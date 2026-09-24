"""Render DOCX placeholders without allowing paths outside approved folders."""

from __future__ import annotations

import os
from pathlib import Path
import tempfile
from typing import Iterable, Mapping

from docx import Document

from app.template_fields import DOCUMENT_PLACEHOLDER_RE


class DocumentTemplateError(ValueError):
    """The template or its placeholder set is invalid."""


class DocumentPublishError(RuntimeError):
    """The completed document could not be published atomically."""


def _paragraphs(document) -> Iterable:
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _container_paragraphs(cell)
    for section in document.sections:
        for container in (section.header, section.footer):
            yield from _container_paragraphs(container)


def _container_paragraphs(container) -> Iterable:
    yield from container.paragraphs
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _container_paragraphs(cell)


def _replace_in_paragraph(paragraph, values: Mapping[str, str]) -> None:
    runs = list(paragraph.runs)
    original = "".join(run.text for run in runs)
    if not original or not DOCUMENT_PLACEHOLDER_RE.search(original):
        return
    boundaries: list[tuple[int, int]] = []
    offset = 0
    for run in runs:
        boundaries.append((offset, offset + len(run.text)))
        offset += len(run.text)

    def locate(position: int) -> tuple[int, int]:
        for index, (start, end) in enumerate(boundaries):
            if start <= position < end:
                return index, position - start
        raise DocumentTemplateError("Не удалось обработать расположение поля в DOCX.")

    for match in reversed(list(DOCUMENT_PLACEHOLDER_RE.finditer(original))):
        start_run, start_offset = locate(match.start())
        end_run, end_offset_last = locate(match.end() - 1)
        end_offset = end_offset_last + 1
        replacement = values[match.group("legacy") or match.group("bank")]
        if start_run == end_run:
            run = runs[start_run]
            run.text = run.text[:start_offset] + replacement + run.text[end_offset:]
            continue
        start = runs[start_run]
        end = runs[end_run]
        start.text = start.text[:start_offset] + replacement
        for index in range(start_run + 1, end_run):
            runs[index].text = ""
        end.text = end.text[end_offset:]


def _safe_child(directory: Path, file_name: str) -> Path:
    if not isinstance(file_name, str) or not file_name.strip():
        raise DocumentTemplateError("Не указано имя DOCX-шаблона.")
    if Path(file_name).name != file_name or "/" in file_name or "\\" in file_name:
        raise DocumentTemplateError("Шаблон должен находиться внутри общей папки templates.")
    if not file_name.casefold().endswith(".docx"):
        raise DocumentTemplateError("Шаблон должен быть файлом DOCX.")
    return directory / file_name


def inspect_placeholders(document) -> set[str]:
    """Return placeholders, including those split between formatted runs."""

    found: set[str] = set()
    for paragraph in _paragraphs(document):
        text = "".join(run.text for run in paragraph.runs)
        found.update(
            match.group("legacy") or match.group("bank")
            for match in DOCUMENT_PLACEHOLDER_RE.finditer(text)
        )
    return found


def render_docx(
    *, template_directory: Path, template_file_name: str,
    output_directory: Path, output_file_name: str,
    values: Mapping[str, object], required_placeholders: Iterable[str],
) -> Path:
    """Validate, render to a temporary file, then atomically publish a DOCX."""

    template_path = _safe_child(Path(template_directory), template_file_name)
    output_path = _safe_child(Path(output_directory), output_file_name)
    if not template_path.is_file():
        raise DocumentTemplateError("Файл шаблона не найден. Обратитесь к администратору.")
    try:
        document = Document(template_path)
    except (OSError, ValueError) as exc:
        raise DocumentTemplateError("Не удалось открыть DOCX-шаблон.") from exc

    placeholders = inspect_placeholders(document)
    required = set(required_placeholders)
    missing_in_template = sorted(required - placeholders)
    if missing_in_template:
        raise DocumentTemplateError(
            "В шаблоне отсутствуют обязательные поля: " + ", ".join(missing_in_template)
        )
    unknown = sorted(placeholders - set(values))
    if unknown:
        raise DocumentTemplateError(
            "В шаблоне найдены неизвестные поля: " + ", ".join(unknown)
        )
    empty = sorted(name for name in placeholders if values.get(name) is None or str(values[name]).strip() == "")
    if empty:
        raise DocumentTemplateError(
            "Для документа не заполнены обязательные данные: " + ", ".join(empty)
        )

    normalized = {name: str(value) for name, value in values.items()}
    for paragraph in _paragraphs(document):
        _replace_in_paragraph(paragraph, normalized)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=".safe-cells-", suffix=".docx", dir=output_path.parent, delete=False
        ) as temporary:
            temporary_path = Path(temporary.name)
        document.save(temporary_path)
        os.replace(temporary_path, output_path)
    except (OSError, ValueError) as exc:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
        raise DocumentPublishError(
            "Не удалось подготовить новый документ."
        ) from exc
    return output_path
