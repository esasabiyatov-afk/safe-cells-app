"""Historical read-only snapshot and XLSX report for every vault cell."""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from io import BytesIO
import json
from pathlib import Path
import sqlite3
from typing import Any, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.config import Settings
from app.db.connections import (
    DatabaseUnavailableError,
    NETWORK_ERROR_MESSAGE,
    open_readonly,
    validate_database_pair,
)
from app.services.journal_reports import REPORT_MIMETYPE


RUSSIAN_MONTHS = (
    "",
    "января",
    "февраля",
    "марта",
    "апреля",
    "мая",
    "июня",
    "июля",
    "августа",
    "сентября",
    "октября",
    "ноября",
    "декабря",
)
STATE_ACTIONS = frozenset(
    {
        "contract.created",
        "contract.closed",
        "cell.manual_occupied",
        "cell.manual_released",
        "cell.bank_occupied",
        "cell.bank_released",
        "cell.key_restored",
    }
)
LIFECYCLE_ACTIONS = frozenset(
    {
        "admin.cell.created",
        "admin.cells.created",
        "admin.cell.retired",
        "admin.cell.restored",
        "admin.cell.deleted",
    }
)


class CellStateReportValidationError(ValueError):
    """The selected report date is absent or invalid."""


class CellStateReportReadError(RuntimeError):
    """The database pair could not be read safely."""


@dataclass(frozen=True, slots=True)
class _StateEvent:
    occurred_at: str
    audit_order: int
    kind: str
    contract_id: str | None = None
    occupation_label: str | None = None
    close_reason: str | None = None


def validate_report_date(value: object, *, today: date) -> date:
    if not isinstance(value, str) or not value.strip():
        raise CellStateReportValidationError("Укажите дату состояния ячеек.")
    try:
        selected = date.fromisoformat(value.strip())
    except ValueError as exc:
        raise CellStateReportValidationError(
            "Укажите корректную дату состояния ячеек."
        ) from exc
    if selected > today:
        raise CellStateReportValidationError(
            "Нельзя сформировать состояние ячеек на будущую дату."
        )
    return selected


def _readonly_uri(path: Path) -> str:
    return f"{path.absolute().as_uri()}?mode=ro"


def _safe_json(value: object) -> dict[str, Any]:
    if not isinstance(value, str):
        return {}
    try:
        parsed = json.loads(value)
    except (TypeError, ValueError):
        return {}
    return parsed if isinstance(parsed, dict) else {}


def _on_or_before(value: object, selected: date) -> bool:
    return isinstance(value, str) and value[:10] <= selected.isoformat()


def _event_sort_key(event: _StateEvent) -> tuple[str, int]:
    return event.occurred_at, event.audit_order


def _natural_cell_key(number: object) -> tuple[int, int | str, str]:
    text = str(number)
    try:
        return 0, int(text), text
    except ValueError:
        return 1, text.casefold(), text


def _contract_rows(connection: sqlite3.Connection) -> list[sqlite3.Row]:
    return connection.execute(
        """
        SELECT
            contract_id, cell_number, client_full_name, start_date, end_date,
            created_at, NULL AS closed_at, NULL AS close_reason,
            NULL AS closure_operation_id
        FROM working.contracts
        UNION ALL
        SELECT
            contract_id, cell_number, client_full_name, start_date, end_date,
            created_at, closed_at, close_reason,
            operation_id AS closure_operation_id
        FROM contracts_archive
        """
    ).fetchall()


def _read_snapshot_source(
    settings: Settings,
) -> tuple[
    list[sqlite3.Row],
    list[sqlite3.Row],
    list[sqlite3.Row],
    list[sqlite3.Row],
    list[sqlite3.Row],
]:
    paths = validate_database_pair(settings)
    with open_readonly(
        paths.archive, busy_timeout_ms=settings.busy_timeout_ms
    ) as connection:
        connection.execute(
            "ATTACH DATABASE ? AS working", (_readonly_uri(paths.working),)
        )
        connection.execute("BEGIN")
        cells = connection.execute(
            """
            SELECT
                cells.number, cells.height_mm,
                COALESCE(cells.width_mm, defaults.width_mm) AS width_mm,
                COALESCE(cells.depth_mm, defaults.depth_mm) AS depth_mm,
                cells.is_active, cells.retired_at
            FROM working.cells AS cells
            CROSS JOIN working.vault_defaults AS defaults
            WHERE defaults.id=1
            """
        ).fetchall()
        contracts = _contract_rows(connection)
        renewals = connection.execute(
            """
            SELECT
                renewal_id, contract_id, cell_number, old_end_date,
                renewal_date, new_end_date, created_at, operation_id
            FROM renewals
            ORDER BY created_at, rowid
            """
        ).fetchall()
        logs = connection.execute(
            f"""
            SELECT
                rowid AS audit_order, operation_id, occurred_at, action,
                contract_id, cell_number, changes_json
            FROM log
            WHERE action IN (
                {",".join("?" for _ in STATE_ACTIONS | LIFECYCLE_ACTIONS)}
            )
            ORDER BY occurred_at, rowid
            """,
            tuple(STATE_ACTIONS | LIFECYCLE_ACTIONS),
        ).fetchall()
        cancellations = connection.execute(
            """
            SELECT original_operation_id, cancelled_at
            FROM operation_cancellations
            """
        ).fetchall()
        connection.rollback()
    return cells, contracts, renewals, logs, cancellations


def _creation_events(logs: Sequence[sqlite3.Row]) -> dict[str, list[str]]:
    events: dict[str, list[str]] = defaultdict(list)
    for row in logs:
        action = str(row["action"])
        occurred_at = str(row["occurred_at"])
        if action == "admin.cell.created" and row["cell_number"] is not None:
            events[str(row["cell_number"])].append(occurred_at)
        elif action == "admin.cells.created":
            for cell in _safe_json(row["changes_json"]).get("cells", []):
                if isinstance(cell, dict) and cell.get("number") is not None:
                    events[str(cell["number"])].append(occurred_at)
    return events


def _cell_is_visible(
    cell: sqlite3.Row,
    *,
    selected: date,
    logs: Sequence[sqlite3.Row],
    creation_events: Mapping[str, Sequence[str]],
) -> bool:
    number = str(cell["number"])
    created = creation_events.get(number, ())
    visible = not created
    if any(_on_or_before(value, selected) for value in created):
        visible = True

    lifecycle_rows = sorted(
        (
            row
            for row in logs
            if str(row["action"]) in LIFECYCLE_ACTIONS
            and row["cell_number"] is not None
            and str(row["cell_number"]) == number
            and _on_or_before(row["occurred_at"], selected)
        ),
        key=lambda row: (str(row["occurred_at"]), int(row["audit_order"])),
    )
    for row in lifecycle_rows:
        action = str(row["action"])
        if action in {"admin.cell.created", "admin.cell.restored"}:
            visible = True
        elif action in {"admin.cell.retired", "admin.cell.deleted"}:
            visible = False

    if (
        visible
        and not lifecycle_rows
        and not bool(cell["is_active"])
        and _on_or_before(cell["retired_at"], selected)
    ):
        visible = False
    return visible


def _effective_operation(
    operation_id: object,
    *,
    selected: date,
    cancellations: Mapping[str, str],
) -> bool:
    if operation_id is None:
        return True
    cancelled_at = cancellations.get(str(operation_id))
    return cancelled_at is None or not _on_or_before(cancelled_at, selected)


def _state_events(
    *,
    selected: date,
    contracts: Sequence[sqlite3.Row],
    logs: Sequence[sqlite3.Row],
    cancellations: Mapping[str, str],
) -> dict[str, list[_StateEvent]]:
    result: dict[str, list[_StateEvent]] = defaultdict(list)
    contract_by_id = {
        str(row["contract_id"]): row
        for row in contracts
    }
    logged_openings = {
        str(row["contract_id"])
        for row in logs
        if str(row["action"]) == "contract.created"
        and row["contract_id"] is not None
    }
    logged_closures = {
        str(row["contract_id"])
        for row in logs
        if str(row["action"]) == "contract.closed"
        and row["contract_id"] is not None
    }

    for row in logs:
        action = str(row["action"])
        if action not in STATE_ACTIONS or not _on_or_before(row["occurred_at"], selected):
            continue
        if not _effective_operation(
            row["operation_id"],
            selected=selected,
            cancellations=cancellations,
        ):
            continue
        cell_number = row["cell_number"]
        if cell_number is None:
            continue
        number = str(cell_number)
        contract_id = (
            str(row["contract_id"]) if row["contract_id"] is not None else None
        )
        common = {
            "occurred_at": str(row["occurred_at"]),
            "audit_order": int(row["audit_order"]),
        }
        if action == "contract.created":
            if contract_id is None:
                continue
            result[number].append(
                _StateEvent(**common, kind="contract", contract_id=contract_id)
            )
        elif action == "contract.closed":
            contract = contract_by_id.get(contract_id or "")
            close_reason = (
                str(contract["close_reason"])
                if contract is not None and contract["close_reason"] is not None
                else str(_safe_json(row["changes_json"]).get("close_reason") or "")
            )
            result[number].append(
                _StateEvent(
                    **common,
                    kind="lost_key" if close_reason == "Потеря ключа" else "free",
                    contract_id=contract_id,
                    close_reason=close_reason or None,
                )
            )
        elif action in {"cell.manual_occupied", "cell.bank_occupied"}:
            label = _safe_json(row["changes_json"]).get("occupation_label")
            result[number].append(
                _StateEvent(
                    **common,
                    kind="manual",
                    occupation_label=(
                        " ".join(label.split())
                        if isinstance(label, str) and label.strip()
                        else "Занята без договора"
                    ),
                )
            )
        else:
            result[number].append(_StateEvent(**common, kind="free"))

    synthetic_order = -2
    for contract in contracts:
        contract_id = str(contract["contract_id"])
        cell_number = str(contract["cell_number"])
        if contract_id not in logged_openings and _on_or_before(
            contract["created_at"], selected
        ):
            result[cell_number].append(
                _StateEvent(
                    occurred_at=str(contract["created_at"]),
                    audit_order=synthetic_order,
                    kind="contract",
                    contract_id=contract_id,
                )
            )
            synthetic_order -= 1

        close_reason = str(contract["close_reason"] or "")
        if (
            contract_id not in logged_closures
            and contract["closed_at"] is not None
            and not close_reason.startswith("Отмена открытия")
            and _on_or_before(contract["closed_at"], selected)
            and _effective_operation(
                contract["closure_operation_id"],
                selected=selected,
                cancellations=cancellations,
            )
        ):
            result[cell_number].append(
                _StateEvent(
                    occurred_at=str(contract["closed_at"]),
                    audit_order=synthetic_order,
                    kind="lost_key" if close_reason == "Потеря ключа" else "free",
                    contract_id=contract_id,
                    close_reason=close_reason or None,
                )
            )
            synthetic_order -= 1

    for events in result.values():
        events.sort(key=_event_sort_key)
    return result


def _contract_period(
    contract: sqlite3.Row,
    *,
    selected: date,
    renewals: Sequence[sqlite3.Row],
    cancellations: Mapping[str, str],
) -> tuple[str, str | None]:
    contract_renewals = [
        row
        for row in renewals
        if str(row["contract_id"]) == str(contract["contract_id"])
    ]
    valid = [
        row
        for row in contract_renewals
        if _on_or_before(row["created_at"], selected)
        and _effective_operation(
            row["operation_id"],
            selected=selected,
            cancellations=cancellations,
        )
    ]
    valid.sort(key=lambda row: str(row["created_at"]))
    if valid:
        latest = valid[-1]
        return str(latest["new_end_date"]), str(latest["renewal_date"])

    if contract_renewals:
        initial = min(contract_renewals, key=lambda row: str(row["created_at"]))
        return str(initial["old_end_date"]), None
    return str(contract["end_date"]), None


def list_cell_states_on_date(
    settings: Settings,
    *,
    as_of_date: date,
) -> list[dict[str, Any]]:
    """Return the end-of-day state of each physical cell visible on the date."""

    try:
        cells, contracts, renewals, logs, cancellation_rows = _read_snapshot_source(
            settings
        )
        cancellations = {
            str(row["original_operation_id"]): str(row["cancelled_at"])
            for row in cancellation_rows
        }
        contract_by_id = {
            str(row["contract_id"]): row
            for row in contracts
        }
        creation_events = _creation_events(logs)
        events_by_cell = _state_events(
            selected=as_of_date,
            contracts=contracts,
            logs=logs,
            cancellations=cancellations,
        )

        rows: list[dict[str, Any]] = []
        for cell in cells:
            if not _cell_is_visible(
                cell,
                selected=as_of_date,
                logs=logs,
                creation_events=creation_events,
            ):
                continue
            state = _StateEvent("", -1, "free")
            for event in events_by_cell.get(str(cell["number"]), ()):
                state = event

            client_name = ""
            start_date: str | None = None
            last_renewal_date: str | None = None
            end_date: str | None = None
            if state.kind in {"contract", "lost_key"} and state.contract_id is not None:
                contract = contract_by_id.get(state.contract_id)
                if contract is not None:
                    client_name = str(contract["client_full_name"])
                    start_date = str(contract["start_date"])
                    end_date, last_renewal_date = _contract_period(
                        contract,
                        selected=as_of_date,
                        renewals=renewals,
                        cancellations=cancellations,
                    )
                if state.kind == "lost_key":
                    client_name = (
                        f"Ключ утерян — {client_name}"
                        if client_name
                        else "Ключ утерян"
                    )
            elif state.kind == "manual":
                client_name = state.occupation_label or "Занята без договора"

            occupied = state.kind != "free"
            rows.append(
                {
                    "cell_number": str(cell["number"]),
                    "height_mm": int(cell["height_mm"]),
                    "width_mm": int(cell["width_mm"]),
                    "depth_mm": int(cell["depth_mm"]),
                    "occupied": occupied,
                    "occupancy_code": 1 if occupied else 0,
                    "client_full_name": client_name,
                    "start_date": start_date,
                    "last_renewal_date": last_renewal_date,
                    "end_date": end_date,
                    "key_count": 1 if occupied else 2,
                }
            )
        rows.sort(key=lambda row: _natural_cell_key(row["cell_number"]))
        return rows
    except CellStateReportReadError:
        raise
    except (DatabaseUnavailableError, OSError, sqlite3.Error) as exc:
        raise CellStateReportReadError(NETWORK_ERROR_MESSAGE) from exc


def _safe_excel_text(value: object) -> str:
    text = str(value or "")
    if text.startswith(("=", "+", "-", "@")):
        return f"'{text}"
    return text


def _display_date(value: object) -> str:
    if not isinstance(value, str):
        return "—"
    try:
        parsed = date.fromisoformat(value)
    except ValueError:
        return "—"
    return f"{parsed.day:02d} {RUSSIAN_MONTHS[parsed.month]} {parsed.year} г."


def build_cell_state_report(
    rows: Sequence[Mapping[str, object]],
    *,
    as_of_date: date,
) -> bytes:
    """Return a styled XLSX snapshot without macros or external links."""

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = as_of_date.strftime("%d.%m.%y")
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A4"
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.print_title_rows = "1:3"
    worksheet.sheet_properties.outlinePr.summaryBelow = True

    worksheet.merge_cells("A1:D1")
    worksheet["A1"] = "Отдел депозитария"
    worksheet["A1"].font = Font(name="Calibri", size=14, bold=True, color="1F2937")
    worksheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    worksheet.merge_cells("E1:H1")
    worksheet["E1"] = "ЗАО АКБ «Толубай»"
    worksheet["E1"].font = Font(name="Calibri", size=14, bold=True, color="1F2937")
    worksheet["E1"].alignment = Alignment(horizontal="right", vertical="center")

    worksheet.merge_cells("A2:H2")
    worksheet["A2"] = (
        f"Состояние всех ячеек на "
        f"{as_of_date.day:02d} {RUSSIAN_MONTHS[as_of_date.month]} "
        f"{as_of_date.year} года"
    )
    worksheet["A2"].font = Font(name="Calibri", size=12, bold=True, color="334E68")
    worksheet["A2"].alignment = Alignment(horizontal="center", vertical="center")

    headers = (
        "№",
        "Размер, мм\n(В × Ш × Г)",
        "0 — свободна\n1 — занята",
        "Ф.И.О. клиента / пометка",
        "Дата открытия",
        "Дата последнего продления",
        "Срок окончания",
        "Количество ключей",
    )
    header_fill = PatternFill("solid", fgColor="5B7083")
    number_fill = PatternFill("solid", fgColor="C0504D")
    occupied_fill = PatternFill("solid", fgColor="FFF2CC")
    free_fill = PatternFill("solid", fgColor="FFFFFF")
    thin = Side(style="thin", color="AAB7C4")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=3, column=column, value=header)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = border

    for row_number, row in enumerate(rows, start=4):
        values = (
            _safe_excel_text(row.get("cell_number")),
            (
                f"{int(row.get('height_mm', 0))} × "
                f"{int(row.get('width_mm', 0))} × "
                f"{int(row.get('depth_mm', 0))}"
            ),
            int(row.get("occupancy_code", 0)),
            _safe_excel_text(row.get("client_full_name")),
            _display_date(row.get("start_date")),
            _display_date(row.get("last_renewal_date")),
            _display_date(row.get("end_date")),
            int(row.get("key_count", 0)),
        )
        occupied = bool(row.get("occupied"))
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(row=row_number, column=column, value=value)
            cell.fill = number_fill if column == 1 else (
                occupied_fill if occupied else free_fill
            )
            cell.font = Font(
                name="Calibri",
                size=10,
                bold=column in {1, 3},
                color="FFFFFF" if column == 1 else "1F2937",
            )
            cell.alignment = Alignment(
                horizontal="center" if column != 4 else "left",
                vertical="center",
                wrap_text=True,
            )
            cell.border = border

    last_row = max(3, 3 + len(rows))
    worksheet.auto_filter.ref = f"A3:H{last_row}"
    worksheet.print_area = f"A1:H{last_row}"
    for column, width in {
        "A": 8,
        "B": 22,
        "C": 17,
        "D": 34,
        "E": 20,
        "F": 25,
        "G": 20,
        "H": 18,
    }.items():
        worksheet.column_dimensions[column].width = width
    worksheet.row_dimensions[1].height = 25
    worksheet.row_dimensions[2].height = 24
    worksheet.row_dimensions[3].height = 42
    for row_number in range(4, last_row + 1):
        worksheet.row_dimensions[row_number].height = 27

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


__all__ = [
    "CellStateReportReadError",
    "CellStateReportValidationError",
    "REPORT_MIMETYPE",
    "build_cell_state_report",
    "list_cell_states_on_date",
    "validate_report_date",
]
