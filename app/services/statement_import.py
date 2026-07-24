"""Read the approved ABS opening statement without storing its client data."""

from __future__ import annotations

from dataclasses import asdict, dataclass
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
import re
from typing import BinaryIO
import zipfile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table


MAX_STATEMENT_BYTES = 2 * 1024 * 1024
MAX_UNCOMPRESSED_STATEMENT_BYTES = 12 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 500

_FIO_LABEL = "2. Фамилия, Имя, Отчество"
_ID_LABEL = "8.2. Серия и номер документа"
_ISSUE_LABEL = "8.3. Дата выдачи и кем выдан документ"
_DATE_PATTERN = re.compile(r"(?<!\d)(\d{2}[.\-/]\d{2}[.\-/]\d{4})(?!\d)")
_ACCOUNT_PATTERN = re.compile(
    r"Банковский\s+сейф\s*\(Ячейки\)\s*№\s*((?:\d[\s-]*){16})(?!\d)",
    re.IGNORECASE,
)


class StatementImportValidationError(ValueError):
    """The uploaded file is not a supported, complete ABS statement."""


@dataclass(frozen=True)
class StatementData:
    client_full_name: str
    id_card_number: str
    id_card_issuer: str
    id_card_issue_date: str
    account_number: str

    def to_dict(self) -> dict[str, str]:
        return asdict(self)


def _normalized(value: object, label: str, maximum: int) -> str:
    if not isinstance(value, str):
        raise StatementImportValidationError(f"Не удалось прочитать поле «{label}».")
    result = " ".join(value.split())
    if not result:
        raise StatementImportValidationError(f"В заявлении не заполнено поле «{label}».")
    if len(result) > maximum:
        raise StatementImportValidationError(f"Поле «{label}» в заявлении слишком длинное.")
    return result


def _read_limited(stream: BinaryIO) -> bytes:
    try:
        data = stream.read(MAX_STATEMENT_BYTES + 1)
    except (OSError, ValueError) as exc:
        raise StatementImportValidationError(
            "Не удалось прочитать выбранный DOCX-файл."
        ) from exc
    if not data:
        raise StatementImportValidationError("Выбран пустой DOCX-файл.")
    if len(data) > MAX_STATEMENT_BYTES:
        raise StatementImportValidationError("Размер заявления не должен превышать 2 МБ.")
    return data


def _validate_docx_archive(data: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ARCHIVE_ENTRIES:
                raise StatementImportValidationError(
                    "DOCX-файл имеет недопустимую структуру."
                )
            if any(entry.flag_bits & 0x1 for entry in entries):
                raise StatementImportValidationError(
                    "Защищённый паролем DOCX-файл не поддерживается."
                )
            if sum(entry.file_size for entry in entries) > MAX_UNCOMPRESSED_STATEMENT_BYTES:
                raise StatementImportValidationError(
                    "DOCX-файл имеет слишком большой распакованный размер."
                )
            names = {entry.filename for entry in entries}
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise StatementImportValidationError("Выбранный файл не является документом DOCX.")
    except StatementImportValidationError:
        raise
    except (OSError, zipfile.BadZipFile, zipfile.LargeZipFile) as exc:
        raise StatementImportValidationError("Не удалось открыть выбранный DOCX-файл.") from exc


def _cell_text(table, row: int, column: int) -> str:
    try:
        return " ".join(table.cell(row, column).text.split())
    except (IndexError, ValueError) as exc:
        raise StatementImportValidationError(
            "Структура заявления отличается от утверждённой формы АБС."
        ) from exc


def _validate_structure(document) -> Table:
    if len(document.tables) != 1:
        raise StatementImportValidationError(
            "Структура заявления отличается от утверждённой формы АБС."
        )
    table = document.tables[0]
    if len(table.rows) != 30 or any(len(row.cells) != 8 for row in table.rows):
        raise StatementImportValidationError(
            "Структура заявления отличается от утверждённой формы АБС."
        )
    expected_labels = (
        (1, 0, _FIO_LABEL),
        (8, 4, _ID_LABEL),
        (10, 0, _ISSUE_LABEL),
    )
    for row, column, expected in expected_labels:
        if _cell_text(table, row, column) != expected:
            raise StatementImportValidationError(
                "Структура заявления отличается от утверждённой формы АБС."
            )
    return table


def _issue_details(value: str, *, as_of_date: date) -> tuple[str, str]:
    matches = list(_DATE_PATTERN.finditer(value))
    if len(matches) != 1:
        raise StatementImportValidationError(
            "Не удалось однозначно определить дату выдачи документа."
        )
    match = matches[0]
    try:
        normalized_date = match.group(1).replace("/", ".").replace("-", ".")
        issued_on = datetime.strptime(normalized_date, "%d.%m.%Y").date()
    except ValueError as exc:
        raise StatementImportValidationError(
            "В заявлении указана неверная дата выдачи документа."
        ) from exc
    if issued_on > as_of_date:
        raise StatementImportValidationError("Дата выдачи документа не может быть в будущем.")
    issuer = _normalized(value[: match.start()].strip(" ,;:-"), "кем выдан документ", 200)
    tail = value[match.end() :].casefold().strip(" .,:;-")
    if tail not in {"", "г"}:
        raise StatementImportValidationError(
            "Не удалось однозначно отделить дату от органа выдачи документа."
        )
    return issuer, issued_on.isoformat()


def extract_statement_data(
    stream: BinaryIO,
    *,
    file_name: object,
    as_of_date: date,
) -> StatementData:
    """Extract five approved fields and never persist the uploaded document."""

    if not isinstance(file_name, str) or not file_name.strip():
        raise StatementImportValidationError("Выберите заявление в формате DOCX.")
    if Path(file_name).suffix.casefold() != ".docx":
        raise StatementImportValidationError("Заявление должно быть файлом DOCX.")
    data = _read_limited(stream)
    _validate_docx_archive(data)
    try:
        document = Document(BytesIO(data))
    except (
        OSError,
        ValueError,
        KeyError,
        PackageNotFoundError,
        zipfile.BadZipFile,
    ) as exc:
        raise StatementImportValidationError("Не удалось открыть выбранный DOCX-файл.") from exc
    table = _validate_structure(document)

    full_name = _normalized(_cell_text(table, 1, 3), "ФИО клиента", 200)
    id_number = re.sub(
        r"\s+",
        "",
        _normalized(_cell_text(table, 9, 4), "серия и номер документа", 100),
    )
    issue_source = _normalized(
        _cell_text(table, 11, 0), "дата выдачи и кем выдан документ", 240
    )
    issuer, issue_date = _issue_details(issue_source, as_of_date=as_of_date)

    account_source = _cell_text(table, 29, 0)
    account_matches = list(_ACCOUNT_PATTERN.finditer(account_source))
    if len(account_matches) != 1:
        raise StatementImportValidationError(
            "Не удалось однозначно определить номер счёта в пункте 17 заявления."
        )
    account_number = "".join(
        character
        for character in account_matches[0].group(1)
        if character.isdigit()
    )
    if len(account_number) != 16:
        raise StatementImportValidationError(
            "Номер счёта в пункте 17 заявления имеет неверный формат."
        )

    return StatementData(
        client_full_name=full_name,
        id_card_number=id_number,
        id_card_issuer=issuer,
        id_card_issue_date=issue_date,
        account_number=account_number,
    )
