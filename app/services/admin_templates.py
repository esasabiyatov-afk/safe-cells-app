"""Safe administrative upload and activation of shared DOCX templates."""

from __future__ import annotations

from datetime import datetime
import json
import os
from pathlib import Path
import sqlite3
import tempfile
from typing import BinaryIO
from uuid import UUID, uuid4
import zipfile

from docx import Document

from app.config import Settings
from app.db.connections import (
    DatabaseCorruptionError,
    DatabaseUnavailableError,
    NETWORK_ERROR_MESSAGE,
    open_readonly,
    open_write,
    validate_database_pair,
)
from app.documents.renderer import DocumentTemplateError, inspect_placeholders
from app.services.admin_settings import (
    AdminBusyError,
    AdminConflictError,
    AdminNetworkError,
    AdminValidationError,
    AdminWriteError,
    AdminWriteResult,
    AdminWriteUncertainError,
    BUSY_MESSAGE,
    UNCERTAIN_MESSAGE,
    _backup_after_commit,
    _employee,
    _existing_operation,
    _operation_id,
    _timestamp,
)
from app.services.documents import ALLOWED_DOCUMENT_PLACEHOLDERS
from app.template_fields import invalid_document_placeholders


MAX_TEMPLATE_BYTES = 10 * 1024 * 1024
MAX_UNCOMPRESSED_TEMPLATE_BYTES = 50 * 1024 * 1024
DOCUMENT_TYPES = frozenset({"opening", "renewal", "closing", "manual"})


def _text(value: object, label: str, maximum: int) -> str:
    if not isinstance(value, str) or not value.strip():
        raise AdminValidationError(f"Поле «{label}» обязательно.")
    normalized = " ".join(value.split())
    if len(normalized) > maximum:
        raise AdminValidationError(f"Поле «{label}» слишком длинное.")
    return normalized


def _document_type(value: object) -> str:
    normalized = _text(value, "Тип документа", 20).lower()
    if normalized not in DOCUMENT_TYPES:
        raise AdminValidationError("Неизвестный тип комплекта документов.")
    return normalized


def _file_name(value: object) -> str:
    name = _text(value, "Имя файла", 180)
    if Path(name).name != name or "/" in name or "\\" in name:
        raise AdminValidationError("Файл шаблона должен находиться внутри папки templates.")
    if not name.casefold().endswith(".docx"):
        raise AdminValidationError("Шаблон должен быть файлом DOCX.")
    return name


def _template_id(value: object | None) -> str:
    if value in (None, ""):
        return str(uuid4())
    if not isinstance(value, str):
        raise AdminValidationError("Неверный идентификатор шаблона.")
    try:
        return str(UUID(value))
    except ValueError:
        # The six approved seed identifiers are stable readable IDs.
        if 1 <= len(value) <= 100 and all(
            character.isalnum() or character in "-_" for character in value
        ):
            return value
        raise AdminValidationError("Неверный идентификатор шаблона.")


def _validate_placeholder_context(
    placeholders: list[str],
    *,
    document_type: str,
) -> None:
    invalid = sorted(
        invalid_document_placeholders(
            set(placeholders), document_type=document_type
        )
    )
    if invalid:
        raise AdminValidationError(
            "Эти поля доступны только для документа продления: "
            + ", ".join(invalid)
            + "."
        )


def _stage_docx(
    template_directory: Path,
    stream: BinaryIO,
    *,
    document_type: str,
) -> tuple[Path, list[str]]:
    template_directory.mkdir(exist_ok=True)
    temporary_path: Path | None = None
    total = 0
    try:
        with tempfile.NamedTemporaryFile(
            prefix=".admin-upload-",
            suffix=".docx",
            dir=template_directory,
            delete=False,
        ) as temporary:
            temporary_path = Path(temporary.name)
            while True:
                chunk = stream.read(64 * 1024)
                if not chunk:
                    break
                total += len(chunk)
                if total > MAX_TEMPLATE_BYTES:
                    raise AdminValidationError("Размер DOCX-шаблона не должен превышать 10 МБ.")
                temporary.write(chunk)
        if total == 0:
            raise AdminValidationError("Загружен пустой файл шаблона.")
        try:
            with zipfile.ZipFile(temporary_path) as archive:
                uncompressed = sum(item.file_size for item in archive.infolist())
        except (OSError, zipfile.BadZipFile) as exc:
            raise AdminValidationError("Не удалось открыть загруженный DOCX-шаблон.") from exc
        if uncompressed > MAX_UNCOMPRESSED_TEMPLATE_BYTES:
            raise AdminValidationError("Распакованный DOCX-шаблон не должен превышать 50 МБ.")
        try:
            document = Document(temporary_path)
        except (OSError, ValueError) as exc:
            raise AdminValidationError("Не удалось открыть загруженный DOCX-шаблон.") from exc
        placeholders = sorted(inspect_placeholders(document))
        unknown = sorted(set(placeholders) - ALLOWED_DOCUMENT_PLACEHOLDERS)
        if unknown:
            raise AdminValidationError(
                "В шаблоне найдены неизвестные поля: " + ", ".join(unknown)
            )
        _validate_placeholder_context(
            placeholders, document_type=document_type
        )
        return temporary_path, placeholders
    except Exception:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
        raise


def save_document_template(
    settings: Settings,
    *,
    operation_id: object,
    template_id: object | None,
    document_type: object,
    display_name: object,
    file_name: object,
    stream: BinaryIO,
    employee: object,
    occurred_at: datetime,
) -> tuple[AdminWriteResult, str]:
    normalized_operation = _operation_id(operation_id)
    normalized_id = _template_id(template_id)
    normalized_type = _document_type(document_type)
    normalized_display = _text(display_name, "Название шаблона", 120)
    normalized_file = _file_name(file_name)
    employee_name = _employee(employee)
    timestamp = _timestamp(occurred_at)
    try:
        paths = validate_database_pair(settings)
    except DatabaseCorruptionError as exc:
        raise AdminNetworkError(str(exc)) from exc
    except DatabaseUnavailableError as exc:
        raise AdminNetworkError(NETWORK_ERROR_MESSAGE) from exc
    template_directory = paths.directory / "templates"
    try:
        staged_path, placeholders = _stage_docx(
            template_directory, stream, document_type=normalized_type
        )
    except OSError as exc:
        raise AdminNetworkError(NETWORK_ERROR_MESSAGE) from exc

    target = template_directory / normalized_file
    previous_file: Path | None = None
    published = False
    committed = False
    phase = "opening"
    try:
        with open_write(settings, attach_archive=True) as connection:
            phase = "begin"
            connection.execute("BEGIN IMMEDIATE")
            phase = "transaction"
            if _existing_operation(connection, normalized_operation, "admin.template.saved"):
                connection.rollback()
                return AdminWriteResult(True, False, None), normalized_id
            current = connection.execute(
                "SELECT relative_file_name FROM document_templates WHERE template_id = ?",
                (normalized_id,),
            ).fetchone()
            owner = connection.execute(
                "SELECT template_id FROM document_templates WHERE relative_file_name = ?",
                (normalized_file,),
            ).fetchone()
            if owner is not None and str(owner["template_id"]) != normalized_id:
                raise AdminConflictError("Это имя файла уже использует другой шаблон.")
            if target.exists():
                if current is None or str(current["relative_file_name"]) != normalized_file:
                    raise AdminConflictError("Файл с таким именем уже существует в папке templates.")
                previous_file = template_directory / f".admin-previous-{uuid4().hex}.docx"
                os.replace(target, previous_file)
            os.replace(staged_path, target)
            published = True
            connection.execute(
                """INSERT INTO document_templates(
                       template_id, document_type, display_name, relative_file_name,
                       required_placeholders_json, is_active, updated_at, updated_by
                   ) VALUES(?, ?, ?, ?, ?, 1, ?, ?)
                   ON CONFLICT(template_id) DO UPDATE SET
                       document_type=excluded.document_type,
                       display_name=excluded.display_name,
                       relative_file_name=excluded.relative_file_name,
                       required_placeholders_json=excluded.required_placeholders_json,
                       is_active=1, updated_at=excluded.updated_at,
                       updated_by=excluded.updated_by""",
                (
                    normalized_id,
                    normalized_type,
                    normalized_display,
                    normalized_file,
                    json.dumps(placeholders, ensure_ascii=False),
                    timestamp,
                    employee_name,
                ),
            )
            changes = {
                "template_id": normalized_id,
                "document_type": normalized_type,
                "display_name": normalized_display,
                "relative_file_name": normalized_file,
                "is_active": True,
                "placeholders": placeholders,
            }
            connection.execute(
                """INSERT INTO archive.log(
                       log_id, operation_id, occurred_at, employee, action,
                       contract_id, cell_number, changes_json
                   ) VALUES(?, ?, ?, ?, 'admin.template.saved', NULL, NULL, ?)""",
                (
                    str(uuid4()), normalized_operation, timestamp, employee_name,
                    json.dumps(changes, ensure_ascii=False, sort_keys=True),
                ),
            )
            phase = "committing"
            connection.commit()
            committed = True
            phase = "verifying"
            saved = connection.execute(
                "SELECT 1 FROM archive.log WHERE operation_id = ? AND action = 'admin.template.saved'",
                (normalized_operation,),
            ).fetchone()
            if saved is None:
                raise AdminWriteUncertainError(UNCERTAIN_MESSAGE)
            backup_created, warning = _backup_after_commit(
                connection,
                settings,
                operation_id=normalized_operation,
                occurred_at=occurred_at,
            )
            previous_file and previous_file.unlink(missing_ok=True)
            return AdminWriteResult(False, backup_created, warning), normalized_id
    except (AdminValidationError, AdminConflictError, AdminWriteUncertainError):
        raise
    except DatabaseCorruptionError as exc:
        raise AdminNetworkError(str(exc)) from exc
    except DatabaseUnavailableError as exc:
        raise AdminNetworkError(NETWORK_ERROR_MESSAGE) from exc
    except sqlite3.OperationalError as exc:
        if "locked" in str(exc).lower() and phase in {"opening", "begin", "transaction"}:
            raise AdminBusyError(BUSY_MESSAGE) from exc
        if phase in {"committing", "verifying"}:
            raise AdminWriteUncertainError(UNCERTAIN_MESSAGE) from exc
        raise AdminNetworkError(NETWORK_ERROR_MESSAGE) from exc
    except (OSError, sqlite3.Error, DocumentTemplateError) as exc:
        if phase in {"committing", "verifying"}:
            raise AdminWriteUncertainError(UNCERTAIN_MESSAGE) from exc
        raise AdminWriteError("Шаблон не сохранён. Изменения отменены.") from exc
    finally:
        staged_path.unlink(missing_ok=True)
        if not committed and published and phase not in {"committing", "verifying"}:
            target.unlink(missing_ok=True)
            if previous_file is not None and previous_file.exists():
                os.replace(previous_file, target)
        if committed and previous_file is not None:
            previous_file.unlink(missing_ok=True)


def get_document_template_path(settings: Settings, template_id: object) -> Path:
    """Return a validated registered DOCX path for an administrator download."""
    normalized_id = _template_id(template_id)
    try:
        paths = validate_database_pair(settings)
        with open_readonly(
            paths.working, busy_timeout_ms=settings.busy_timeout_ms
        ) as connection:
            row = connection.execute(
                "SELECT relative_file_name FROM document_templates WHERE template_id = ?",
                (normalized_id,),
            ).fetchone()
    except (DatabaseUnavailableError, OSError, sqlite3.Error) as exc:
        raise AdminNetworkError(NETWORK_ERROR_MESSAGE) from exc
    if row is None:
        raise AdminConflictError("Шаблон не найден. Обновите настройки.")
    file_name = _file_name(row["relative_file_name"])
    path = paths.directory / "templates" / file_name
    if not path.is_file():
        raise AdminConflictError("Файл DOCX этого шаблона не найден.")
    return path


def update_document_template(
    settings: Settings,
    *,
    payload: object,
    employee: object,
    occurred_at: datetime,
) -> AdminWriteResult:
    if not isinstance(payload, dict) or set(payload) != {
        "operation_id", "template_id", "document_type", "display_name", "is_active"
    }:
        raise AdminValidationError("Переданы неизвестные или неполные данные шаблона.")
    operation_id = _operation_id(payload["operation_id"])
    template_id = _template_id(payload["template_id"])
    document_type = _document_type(payload["document_type"])
    display_name = _text(payload["display_name"], "Название шаблона", 120)
    if not isinstance(payload["is_active"], bool):
        raise AdminValidationError("Состояние шаблона должно быть включено или выключено.")
    is_active = payload["is_active"]
    employee_name = _employee(employee)
    timestamp = _timestamp(occurred_at)
    phase = "opening"
    try:
        with open_write(settings, attach_archive=True) as connection:
            phase = "begin"
            connection.execute("BEGIN IMMEDIATE")
            phase = "transaction"
            if _existing_operation(connection, operation_id, "admin.template.updated"):
                connection.rollback()
                return AdminWriteResult(True, False, None)
            row = connection.execute(
                "SELECT * FROM document_templates WHERE template_id = ?", (template_id,)
            ).fetchone()
            if row is None:
                raise AdminConflictError("Шаблон не найден. Обновите настройки.")
            try:
                stored_placeholders = json.loads(
                    str(row["required_placeholders_json"])
                )
            except (TypeError, json.JSONDecodeError) as exc:
                raise AdminWriteError(
                    "Список полей шаблона повреждён."
                ) from exc
            if not isinstance(stored_placeholders, list) or not all(
                isinstance(item, str) for item in stored_placeholders
            ):
                raise AdminWriteError("Список полей шаблона повреждён.")
            _validate_placeholder_context(
                stored_placeholders, document_type=document_type
            )
            if is_active and not (
                settings.database_directory / "templates" / str(row["relative_file_name"])
            ).is_file():
                raise AdminConflictError("Нельзя включить шаблон: файл DOCX не найден.")
            connection.execute(
                """UPDATE document_templates SET document_type=?, display_name=?,
                       is_active=?, updated_at=?, updated_by=? WHERE template_id=?""",
                (
                    document_type, display_name, int(is_active), timestamp,
                    employee_name, template_id,
                ),
            )
            changes = {
                "template_id": template_id,
                "document_type": {"old": row["document_type"], "new": document_type},
                "display_name": {"old": row["display_name"], "new": display_name},
                "is_active": {"old": bool(row["is_active"]), "new": is_active},
            }
            connection.execute(
                """INSERT INTO archive.log(
                       log_id, operation_id, occurred_at, employee, action,
                       contract_id, cell_number, changes_json
                   ) VALUES(?, ?, ?, ?, 'admin.template.updated', NULL, NULL, ?)""",
                (
                    str(uuid4()), operation_id, timestamp, employee_name,
                    json.dumps(changes, ensure_ascii=False, sort_keys=True),
                ),
            )
            phase = "committing"
            connection.commit()
            phase = "verifying"
            if connection.execute(
                "SELECT 1 FROM archive.log WHERE operation_id = ?", (operation_id,)
            ).fetchone() is None:
                raise AdminWriteUncertainError(UNCERTAIN_MESSAGE)
            backup_created, warning = _backup_after_commit(
                connection, settings, operation_id=operation_id, occurred_at=occurred_at
            )
            return AdminWriteResult(False, backup_created, warning)
    except (AdminValidationError, AdminConflictError, AdminWriteUncertainError):
        raise
    except DatabaseCorruptionError as exc:
        raise AdminNetworkError(str(exc)) from exc
    except DatabaseUnavailableError as exc:
        raise AdminNetworkError(NETWORK_ERROR_MESSAGE) from exc
    except sqlite3.OperationalError as exc:
        if "locked" in str(exc).lower() and phase in {"opening", "begin", "transaction"}:
            raise AdminBusyError(BUSY_MESSAGE) from exc
        if phase in {"committing", "verifying"}:
            raise AdminWriteUncertainError(UNCERTAIN_MESSAGE) from exc
        raise AdminNetworkError(NETWORK_ERROR_MESSAGE) from exc
    except (OSError, sqlite3.Error) as exc:
        if phase in {"committing", "verifying"}:
            raise AdminWriteUncertainError(UNCERTAIN_MESSAGE) from exc
        raise AdminWriteError("Шаблон не сохранён. Изменения отменены.") from exc
