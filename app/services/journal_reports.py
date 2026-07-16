"""Create a local Excel statement from already filtered journal entries."""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


REPORT_MIMETYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _safe_excel_text(value: object) -> str:
    text = str(value or "")
    if text.startswith(("=", "+", "-", "@")):
        return f"'{text}"
    return text


def _occurred_at(value: object) -> str:
    try:
        parsed = datetime.fromisoformat(str(value))
    except ValueError:
        return _safe_excel_text(value)
    return parsed.strftime("%d.%m.%Y %H:%M")


def _display_filter_date(value: object) -> str:
    try:
        return date.fromisoformat(str(value)).strftime("%d.%m.%Y")
    except ValueError:
        return "—"


def _filter_text(filters: Mapping[str, object]) -> str:
    parts = []
    if filters.get("cell_number"):
        parts.append(f"ячейка № {filters['cell_number']}")
    if filters.get("action_label"):
        parts.append(f"действие: {filters['action_label']}")
    if filters.get("date_from"):
        parts.append(f"с {_display_filter_date(filters['date_from'])}")
    if filters.get("date_to"):
        parts.append(f"по {_display_filter_date(filters['date_to'])}")
    return "; ".join(parts) or "все операции"


def build_journal_report(
    entries: Sequence[Mapping[str, object]],
    *,
    generated_on: date,
    filters: Mapping[str, object],
) -> bytes:
    """Return a formatted XLSX workbook without macros or external links."""

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Журнал ячеек"
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A6"
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.print_title_rows = "1:5"

    worksheet.merge_cells("A1:F1")
    title = worksheet["A1"]
    title.value = "Выписка из журнала сейфовых ячеек"
    title.font = Font(name="Calibri", size=16, bold=True, color="1F2937")
    title.alignment = Alignment(horizontal="center")

    worksheet.merge_cells("A2:F2")
    worksheet["A2"] = f"Сформировано: {generated_on.strftime('%d.%m.%Y')}"
    worksheet["A2"].alignment = Alignment(horizontal="center")

    worksheet.merge_cells("A3:F3")
    worksheet["A3"] = f"Отбор: {_safe_excel_text(_filter_text(filters))}"
    worksheet["A3"].alignment = Alignment(horizontal="center")

    headers = (
        "Дата и время",
        "Ячейка",
        "Действие",
        "Клиент",
        "Сотрудник",
        "Сведения",
    )
    header_fill = PatternFill("solid", fgColor="243B53")
    thin = Side(style="thin", color="CBD5E1")
    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=5, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)

    for row_number, entry in enumerate(entries, start=6):
        values = (
            _occurred_at(entry.get("occurred_at")),
            _safe_excel_text(entry.get("cell_number")),
            _safe_excel_text(
                entry.get("report_action_label") or entry.get("action_label")
            ),
            _safe_excel_text(entry.get("client_full_name")),
            _safe_excel_text(entry.get("employee")),
            _safe_excel_text(entry.get("summary")),
        )
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(row=row_number, column=column, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)

    worksheet.auto_filter.ref = f"A5:F{max(5, 5 + len(entries))}"
    for column, width in {
        "A": 19,
        "B": 12,
        "C": 22,
        "D": 34,
        "E": 28,
        "F": 58,
    }.items():
        worksheet.column_dimensions[column].width = width
    worksheet.row_dimensions[1].height = 26
    worksheet.row_dimensions[5].height = 24

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
