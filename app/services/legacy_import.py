"""Preview and atomically import the old occupied/free-cell Excel register."""

from __future__ import annotations

from collections.abc import Callable
from dataclasses import asdict, dataclass
from datetime import date, datetime, time
from hashlib import sha256
from io import BytesIO
import json
import re
import sqlite3
from typing import Any, BinaryIO
from uuid import UUID, NAMESPACE_URL, uuid4, uuid5
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from app.config import Settings
from app.db.connections import (
    DatabaseCorruptionError,
    DatabaseUnavailableError,
    NETWORK_ERROR_MESSAGE,
    open_readonly,
    open_write,
    validate_database_pair,
)
from app.services.backups import (
    BackupBusyError,
    BackupError,
    BackupNetworkError,
    create_backup_pair,
    has_valid_backup_for_operation,
)
from app.services.legacy_contracts import (
    LEGACY_MISSING_DATE,
    LEGACY_MISSING_TEXT,
    legacy_extra_fields,
)


MAX_WORKBOOK_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ZIP_ENTRIES = 1_000
IMPORT_CONFIRMATION = "ПЕРЕНЕСТИ СТАРЫЕ ДОГОВОРЫ"
IMPORT_ACTOR = "Перенос старого учета"
SUMMARY_ACTION = "admin.legacy_contracts_imported"
BUSY_MESSAGE = "База сейчас занята другим сотрудником. Перенос не выполнен."
UNCERTAIN_MESSAGE = (
    "Не удалось подтвердить результат переноса. Не повторяйте операцию автоматически. "
    "Обновите главный экран и проверьте занятые ячейки."
)


class LegacyImportValidationError(ValueError):
    pass


class LegacyImportConflictError(RuntimeError):
    pass


class LegacyImportBusyError(RuntimeError):
    pass


class LegacyImportNetworkError(RuntimeError):
    pass


class LegacyImportWriteError(RuntimeError):
    pass


class LegacyImportWriteUncertainError(RuntimeError):
    pass


@dataclass(frozen=True, slots=True)
class LegacyImportIssue:
    cell_number: str | None
    message: str

    def to_dict(self) -> dict[str, str | None]:
        return asdict(self)


@dataclass(frozen=True, slots=True)
class LegacyImportRow:
    cell_number: str
    kind: str
    client_full_name: str | None = None
    start_date: date | None = None
    end_date: date | None = None
    occupation_label: str | None = None
    id_card_number: str | None = None
    id_card_issuer: str | None = None
    id_card_issue_date: date | None = None
    account_number: str | None = None
    deposit_amount_minor: int | None = None


@dataclass(frozen=True, slots=True)
class LegacyImportPlan:
    sha256: str
    rows: tuple[LegacyImportRow, ...]
    issues: tuple[LegacyImportIssue, ...]
    database_issues: tuple[LegacyImportIssue, ...] = ()

    @property
    def contracts_count(self) -> int:
        return sum(row.kind == "contract" for row in self.rows)

    @property
    def free_count(self) -> int:
        return sum(row.kind == "free" for row in self.rows)

    @property
    def manual_count(self) -> int:
        return sum(row.kind == "manual" for row in self.rows)

    @property
    def passport_details_count(self) -> int:
        return sum(
            row.kind == "contract"
            and bool(row.id_card_number)
            and bool(row.id_card_issuer)
            and row.id_card_issue_date is not None
            for row in self.rows
        )

    @property
    def identity_complete_count(self) -> int:
        return sum(
            row.kind == "contract"
            and bool(row.id_card_number)
            and bool(row.id_card_issuer)
            and row.id_card_issue_date is not None
            and bool(row.account_number)
            for row in self.rows
        )

    @property
    def deposit_known_count(self) -> int:
        return sum(
            row.kind == "contract" and row.deposit_amount_minor is not None
            for row in self.rows
        )

    @property
    def ready(self) -> bool:
        return not self.issues and not self.database_issues

    def to_dict(self) -> dict[str, Any]:
        all_issues = (*self.issues, *self.database_issues)
        return {
            "sha256": self.sha256,
            "ready": self.ready,
            "cells_count": len(self.rows),
            "contracts_count": self.contracts_count,
            "free_count": self.free_count,
            "manual_count": self.manual_count,
            "passport_details_count": self.passport_details_count,
            "identity_complete_count": self.identity_complete_count,
            "deposit_known_count": self.deposit_known_count,
            "issues": [issue.to_dict() for issue in all_issues],
            "confirmation": IMPORT_CONFIRMATION,
        }


@dataclass(frozen=True, slots=True)
class LegacyImportResult:
    operation_id: str
    contracts_count: int
    free_count: int
    manual_count: int
    repeated: bool
    backup_created: bool
    warning: str | None = None

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


RUSSIAN_MONTHS = {
    "январь": 1, "января": 1,
    "февраль": 2, "февраля": 2,
    "март": 3, "марта": 3,
    "апрель": 4, "апреля": 4,
    "май": 5, "мая": 5,
    "июнь": 6, "июня": 6,
    "июль": 7, "июля": 7,
    "август": 8, "августа": 8,
    "сентябрь": 9, "сентября": 9,
    "октябрь": 10, "октября": 10,
    "ноябрь": 11, "ноября": 11,
    "декабрь": 12, "декабря": 12,
}


def read_upload(stream: BinaryIO) -> bytes:
    content = stream.read(MAX_WORKBOOK_BYTES + 1)
    if len(content) > MAX_WORKBOOK_BYTES:
        raise LegacyImportValidationError("Файл Excel больше допустимых 5 МБ.")
    if not content:
        raise LegacyImportValidationError("Выбран пустой файл Excel.")
    return content


def _check_xlsx_container(content: bytes) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                raise LegacyImportValidationError("В файле Excel слишком много внутренних элементов.")
            if sum(item.file_size for item in entries) > MAX_UNCOMPRESSED_BYTES:
                raise LegacyImportValidationError("Распакованный файл Excel слишком большой.")
            if "[Content_Types].xml" not in archive.namelist():
                raise LegacyImportValidationError("Выбранный файл не является корректным XLSX.")
    except BadZipFile as exc:
        raise LegacyImportValidationError("Выбранный файл не является корректным XLSX.") from exc


def _text(value: object) -> str:
    return " ".join(str(value or "").replace("\xa0", " ").split())


def _cell_number(value: object) -> str | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)) and float(value).is_integer():
        return str(int(value))
    normalized = _text(value)
    return normalized if normalized else None


def _status(value: object) -> int | None:
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)) and float(value).is_integer() and int(value) in {0, 1}:
        return int(value)
    normalized = _text(value)
    return int(normalized) if normalized in {"0", "1"} else None


def _date_value(value: object, *, epoch: datetime) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, bool):
        raise ValueError
    if isinstance(value, (int, float)):
        converted = from_excel(value, epoch=epoch)
        return converted.date() if isinstance(converted, datetime) else converted
    raw = _text(value)
    if not raw or raw in {"-", "—"}:
        return None
    if raw.startswith("="):
        raise ValueError
    normalized = raw.casefold().replace("ё", "е")
    normalized = re.sub(r"\b(?:г|года|год)\.?\b", " ", normalized)
    numeric = re.search(r"(?<!\d)(\d{1,2})[./-](\d{1,2})[./-](\d{4})(?!\d)", normalized)
    if numeric:
        return date(int(numeric.group(3)), int(numeric.group(2)), int(numeric.group(1)))
    russian = re.search(
        r"(?<!\d)(\d{1,2})\s*([а-я]+)\s*(\d{4})(?!\d)", normalized
    )
    if russian and russian.group(2) in RUSSIAN_MONTHS:
        return date(
            int(russian.group(3)),
            RUSSIAN_MONTHS[russian.group(2)],
            int(russian.group(1)),
        )
    cleaned = re.sub(r"[,]+", " ", raw)
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    for pattern in (
        "%A %B %d %Y", "%a %B %d %Y", "%B %d %Y", "%d %B %Y",
        "%A %b %d %Y", "%a %b %d %Y", "%b %d %Y", "%d %b %Y",
    ):
        try:
            return datetime.strptime(cleaned, pattern).date()
        except ValueError:
            continue
    raise ValueError


def _header(value: object) -> str:
    return re.sub(r"[^a-zа-я0-9]+", "", _text(value).casefold().replace("ё", "е"))


def _name_key(value: object) -> str:
    return _header(value)


def _id_card_number(value: object) -> str:
    return re.sub(r"\s+", "", _text(value))


def _account_number(value: object) -> str | None:
    normalized = _text(value)
    if not normalized or _header(normalized) == "безномера":
        return None
    return normalized


def _deposit_amount(value: object) -> int | None:
    if value is None or _text(value) in {"", "-", "—"}:
        return None
    if isinstance(value, bool):
        raise ValueError
    if isinstance(value, (int, float)):
        if not float(value).is_integer():
            raise ValueError
        result = int(value)
    else:
        normalized = _text(value).replace(" ", "")
        if not normalized.isdecimal():
            raise ValueError
        result = int(normalized)
    if not 0 <= result <= 10_000_000:
        raise ValueError
    return result


def _maximum_column(worksheet, minimum: int) -> int:
    if worksheet.max_column is None:
        worksheet.calculate_dimension(force=True)
    return max(worksheet.max_column or 0, minimum)


OPTIONAL_FIRST_SHEET_HEADERS = {
    "паспорт": "id_card_number",
    "серияиномерпаспорта": "id_card_number",
    "серияиномерidкарты": "id_card_number",
    "серияиномердокумента": "id_card_number",
    "орган": "id_card_issuer",
    "кемвыдан": "id_card_issuer",
    "органвыдачи": "id_card_issuer",
    "сроквыдачи": "id_card_issue_date",
    "датавыдачи": "id_card_issue_date",
    "датавыдачипаспорта": "id_card_issue_date",
    "номерсчета": "account_number",
    "счетномер": "account_number",
    "номердоговора": "account_number",
    "суммазалога": "deposit_amount_minor",
    "залог": "deposit_amount_minor",
}


def _optional_first_sheet_columns(header: tuple[object, ...]) -> dict[str, int]:
    columns: dict[str, int] = {}
    for index, value in enumerate(header[6:], start=6):
        field = OPTIONAL_FIRST_SHEET_HEADERS.get(_header(value))
        if field is None:
            continue
        if field in columns:
            raise LegacyImportValidationError(
                "В XLSX повторяется столбец с паспортными данными."
            )
        columns[field] = index
    return columns


def _second_sheet_identity(
    workbook,
) -> tuple[dict[str, tuple[object, object, object, object]], tuple[LegacyImportIssue, ...]]:
    if len(workbook.sheetnames) < 2:
        return {}, ()
    worksheet = workbook[workbook.sheetnames[1]]
    maximum = _maximum_column(worksheet, 1)
    header = next(
        worksheet.iter_rows(min_row=1, max_row=1, max_col=maximum, values_only=True),
        (),
    )
    normalized = [_header(value) for value in header]

    def column(*aliases: str) -> int | None:
        return next(
            (index for index, value in enumerate(normalized) if value in aliases),
            None,
        )

    name_column = column("фиоклиента", "фио", "клиент")
    issue_date_column = column("сроквыдачи", "датавыдачи", "датавыдачипаспорта")
    issuer_column = column("орган", "кемвыдан", "органвыдачи")
    passport_column = column(
        "паспорт",
        "серияиномерпаспорта",
        "серияиномерidкарты",
        "серияиномердокумента",
    )
    if (
        name_column is None
        or issue_date_column is None
        or issuer_column is None
        or passport_column is None
    ):
        return {}, ()
    split_passport = (
        normalized[passport_column] == "паспорт"
        and passport_column + 1 < len(normalized)
        and not normalized[passport_column + 1]
    )
    lookup: dict[str, tuple[object, object, object, object]] = {}
    issues: list[LegacyImportIssue] = []
    for values in worksheet.iter_rows(min_row=2, max_col=maximum, values_only=True):
        name = _text(values[name_column])
        if not name:
            continue
        key = _name_key(name)
        raw_passport = values[passport_column]
        if split_passport:
            raw_passport = (
                f"{_text(raw_passport)}{_text(values[passport_column + 1])}"
            )
        candidate = (
            raw_passport,
            values[issuer_column],
            values[issue_date_column],
            None,
        )
        previous = lookup.get(key)
        if previous is not None:
            comparable_previous = (
                _id_card_number(previous[0]),
                _text(previous[1]).casefold(),
                _text(previous[2]).casefold(),
            )
            comparable_candidate = (
                _id_card_number(candidate[0]),
                _text(candidate[1]).casefold(),
                _text(candidate[2]).casefold(),
            )
            if comparable_previous != comparable_candidate:
                issues.append(
                    LegacyImportIssue(
                        None,
                        "На втором листе для одного клиента указаны разные паспортные данные.",
                    )
                )
                continue
        lookup[key] = candidate
    return lookup, tuple(issues)


def _identity_values(
    values: tuple[object, ...],
    columns: dict[str, int],
    second_sheet: dict[str, tuple[object, object, object, object]],
    *,
    client_full_name: str,
    epoch: datetime,
    cell_number: str,
    issues: list[LegacyImportIssue],
) -> tuple[str | None, str | None, date | None, str | None, int | None]:
    raw = {
        field: values[index] if index < len(values) else None
        for field, index in columns.items()
    }
    fallback = second_sheet.get(_name_key(client_full_name))
    if fallback is not None:
        for field, candidate in zip(
            (
                "id_card_number",
                "id_card_issuer",
                "id_card_issue_date",
                "account_number",
            ),
            fallback,
            strict=True,
        ):
            if not _text(raw.get(field)):
                raw[field] = candidate

    number = _id_card_number(raw.get("id_card_number")) or None
    issuer = _text(raw.get("id_card_issuer")) or None
    account = _account_number(raw.get("account_number"))
    deposit = None
    if _text(raw.get("deposit_amount_minor")):
        try:
            deposit = _deposit_amount(raw["deposit_amount_minor"])
        except (TypeError, ValueError, OverflowError):
            issues.append(
                LegacyImportIssue(cell_number, "Некорректная сумма залога.")
            )
    issue_date = None
    if _text(raw.get("id_card_issue_date")):
        try:
            issue_date = _date_value(raw["id_card_issue_date"], epoch=epoch)
        except (TypeError, ValueError, OverflowError):
            issues.append(
                LegacyImportIssue(cell_number, "Некорректная дата выдачи паспорта.")
            )
    for value, maximum, label in (
        (number, 100, "Серия и номер паспорта"),
        (issuer, 200, "Орган выдачи паспорта"),
        (account, 100, "Номер счёта"),
    ):
        if value is not None and len(value) > maximum:
            issues.append(
                LegacyImportIssue(cell_number, f"Поле «{label}» слишком длинное.")
            )
    return number, issuer, issue_date, account, deposit


def _parse_workbook(content: bytes) -> LegacyImportPlan:
    _check_xlsx_container(content)
    digest = sha256(content).hexdigest()
    try:
        workbook = load_workbook(
            BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
    except Exception as exc:
        raise LegacyImportValidationError("Не удалось прочитать XLSX-файл.") from exc
    issues: list[LegacyImportIssue] = []
    rows: list[LegacyImportRow] = []
    try:
        if not workbook.sheetnames:
            raise LegacyImportValidationError("В XLSX-файле нет листов.")
        worksheet = workbook[workbook.sheetnames[0]]
        maximum = _maximum_column(worksheet, 6)
        header = next(
            worksheet.iter_rows(
                min_row=1, max_row=1, max_col=maximum, values_only=True
            ),
            None,
        )
        expected = ("№", "Размер", "0 - свободен", "Ф.И.О. Клиента", "дата открытия", "Срок окончания")
        if header is None or tuple(_header(value) for value in header[:6]) != tuple(_header(value) for value in expected):
            raise LegacyImportValidationError(
                "Столбцы XLSX не совпадают с утверждённым отчётом по ячейкам."
            )
        optional_columns = _optional_first_sheet_columns(tuple(header))
        second_sheet, second_sheet_issues = _second_sheet_identity(workbook)
        issues.extend(second_sheet_issues)
        seen: set[str] = set()
        for values in worksheet.iter_rows(
            min_row=2, max_col=maximum, values_only=True
        ):
            if not any(value is not None and _text(value) for value in values):
                continue
            number = _cell_number(values[0])
            if number is None:
                issues.append(LegacyImportIssue(None, "В строке не указан номер ячейки."))
                continue
            if number in seen:
                issues.append(LegacyImportIssue(number, "Номер ячейки повторяется в файле."))
                continue
            seen.add(number)
            if len(number) > 50:
                issues.append(LegacyImportIssue(number[:50], "Номер ячейки длиннее 50 символов."))
                continue
            checked_columns = (*range(6), *optional_columns.values())
            if any(
                isinstance(values[index], str)
                and values[index].lstrip().startswith("=")
                for index in checked_columns
                if index < len(values)
            ):
                issues.append(LegacyImportIssue(number, "Формулы в обязательных столбцах запрещены."))
                continue
            status = _status(values[2])
            if status is None:
                issues.append(LegacyImportIssue(number, "Статус должен быть 0 (свободна) или 1 (занята)."))
                continue
            name = _text(values[3])
            if len(name) > 200:
                issues.append(LegacyImportIssue(number, "ФИО клиента длиннее 200 символов."))
            start_invalid = False
            try:
                start = _date_value(values[4], epoch=workbook.epoch)
            except (TypeError, ValueError, OverflowError):
                start = None
                start_invalid = True
                issues.append(LegacyImportIssue(number, "Некорректная дата начала аренды."))
            end_invalid = False
            try:
                end = _date_value(values[5], epoch=workbook.epoch)
            except (TypeError, ValueError, OverflowError):
                end = None
                end_invalid = True
                issues.append(LegacyImportIssue(number, "Некорректная дата окончания аренды."))
            if status == 0:
                if (
                    name
                    or start is not None
                    or end is not None
                    or any(
                        _text(values[index])
                        for index in optional_columns.values()
                        if index < len(values)
                    )
                ):
                    issues.append(LegacyImportIssue(number, "Свободная ячейка содержит данные договора."))
                rows.append(LegacyImportRow(number, "free"))
                continue
            if name and start is None and end is None:
                if len(name) > 80:
                    issues.append(LegacyImportIssue(number, "Служебная пометка длиннее 80 символов."))
                rows.append(LegacyImportRow(number, "manual", occupation_label=name[:80]))
                continue
            if not name:
                issues.append(LegacyImportIssue(number, "Для занятой ячейки не указано ФИО клиента."))
            if start is None and not start_invalid:
                issues.append(LegacyImportIssue(number, "Не указана дата начала аренды."))
            if end is None and not end_invalid:
                issues.append(LegacyImportIssue(number, "Не указана дата окончания аренды."))
            if start is not None and end is not None and end < start:
                issues.append(LegacyImportIssue(number, "Дата окончания раньше даты начала."))
            identity = _identity_values(
                tuple(values),
                optional_columns,
                second_sheet,
                client_full_name=name,
                epoch=workbook.epoch,
                cell_number=number,
                issues=issues,
            )
            rows.append(
                LegacyImportRow(
                    number,
                    "contract",
                    name or None,
                    start,
                    end,
                    None,
                    *identity,
                )
            )
    finally:
        workbook.close()
    return LegacyImportPlan(digest, tuple(rows), tuple(issues))


def _database_issues(
    working: sqlite3.Connection,
    archive: sqlite3.Connection | None,
    plan: LegacyImportPlan,
    *,
    as_of_date: date,
) -> tuple[LegacyImportIssue, ...]:
    issues: list[LegacyImportIssue] = []
    database_cells = {str(row[0]) for row in working.execute("SELECT number FROM cells").fetchall()}
    file_cells = {row.cell_number for row in plan.rows}
    missing = sorted(database_cells - file_cells, key=lambda value: (len(value), value))
    extra = sorted(file_cells - database_cells, key=lambda value: (len(value), value))
    if missing:
        issues.append(LegacyImportIssue(None, f"В файле отсутствуют ячейки: {', '.join(missing)}."))
    if extra:
        issues.append(LegacyImportIssue(None, f"В файле есть неизвестные ячейки: {', '.join(extra)}."))
    if working.execute("SELECT 1 FROM contracts LIMIT 1").fetchone() is not None:
        issues.append(LegacyImportIssue(None, "В рабочей базе уже есть договоры. Повторный перенос запрещён."))
    if working.execute("SELECT 1 FROM cell_blocks LIMIT 1").fetchone() is not None:
        issues.append(LegacyImportIssue(None, "В рабочей базе уже есть служебно занятые ячейки."))
    if archive is not None:
        archive_prefix = "archive." if archive is working else ""
        if archive.execute(f"SELECT 1 FROM {archive_prefix}contracts_archive LIMIT 1").fetchone() is not None:
            issues.append(LegacyImportIssue(None, "В архиве уже есть закрытые договоры. Перенос запрещён."))
        if archive.execute(f"SELECT 1 FROM {archive_prefix}renewals LIMIT 1").fetchone() is not None:
            issues.append(LegacyImportIssue(None, "В архиве уже есть продления. Перенос запрещён."))
    for row in plan.rows:
        if row.kind == "contract" and row.start_date is not None and row.start_date > as_of_date:
            issues.append(LegacyImportIssue(row.cell_number, "Дата начала активного договора находится в будущем."))
        if (
            row.kind == "contract"
            and row.id_card_issue_date is not None
            and row.id_card_issue_date > as_of_date
        ):
            issues.append(
                LegacyImportIssue(
                    row.cell_number,
                    "Дата выдачи паспорта находится в будущем.",
                )
            )
    return tuple(issues)


def preview_legacy_import(
    settings: Settings, *, content: bytes, file_name: object, as_of_date: date
) -> LegacyImportPlan:
    if not isinstance(file_name, str) or not file_name.casefold().endswith(".xlsx"):
        raise LegacyImportValidationError("Выберите файл формата XLSX.")
    plan = _parse_workbook(content)
    try:
        paths = validate_database_pair(settings)
        with open_readonly(paths.working, busy_timeout_ms=settings.busy_timeout_ms) as working:
            with open_readonly(paths.archive, busy_timeout_ms=settings.busy_timeout_ms) as archive:
                database_issues = _database_issues(
                    working, archive, plan, as_of_date=as_of_date
                )
    except (DatabaseUnavailableError, OSError, sqlite3.Error) as exc:
        raise LegacyImportNetworkError(NETWORK_ERROR_MESSAGE) from exc
    return LegacyImportPlan(plan.sha256, plan.rows, plan.issues, database_issues)


def _operation_id(value: object) -> str:
    if not isinstance(value, str):
        raise LegacyImportValidationError("Не удалось подготовить операцию переноса.")
    try:
        return str(UUID(value.strip()))
    except (ValueError, AttributeError) as exc:
        raise LegacyImportValidationError("Не удалось подготовить операцию переноса.") from exc


def _result_from_summary(
    row: sqlite3.Row, settings: Settings, *, repeated: bool
) -> LegacyImportResult:
    try:
        changes = json.loads(str(row["changes_json"]))
        return LegacyImportResult(
            operation_id=str(row["operation_id"]),
            contracts_count=int(changes["contracts_count"]),
            free_count=int(changes["free_count"]),
            manual_count=int(changes["manual_count"]),
            repeated=repeated,
            backup_created=has_valid_backup_for_operation(settings, str(row["operation_id"])),
            warning=None,
        )
    except (TypeError, ValueError, KeyError, json.JSONDecodeError) as exc:
        raise LegacyImportWriteError("Запись предыдущего переноса повреждена.") from exc


def _insert_plan(
    connection: sqlite3.Connection,
    plan: LegacyImportPlan,
    *,
    operation_id: str,
    occurred_at: datetime,
    after_insert: Callable[[LegacyImportRow], None] | None,
) -> None:
    timestamp = occurred_at.isoformat(timespec="seconds")
    for row in plan.rows:
        row_operation = str(uuid5(UUID(operation_id), f"cell:{row.cell_number}:{row.kind}"))
        if row.kind == "contract":
            assert row.client_full_name and row.start_date and row.end_date
            contract_id = str(uuid5(NAMESPACE_URL, f"safe-cells-legacy:{plan.sha256}:{row.cell_number}"))
            created_at = datetime.combine(row.start_date, time.min, tzinfo=occurred_at.tzinfo).isoformat(timespec="seconds")
            rent_days = (row.end_date - row.start_date).days + 1
            connection.execute(
                """
                INSERT INTO main.contracts(
                    contract_id, cell_number, client_full_name, id_card_number,
                    id_card_issuer, id_card_issue_date, account_number,
                    extra_fields_json, start_date, end_date, rent_days,
                    price_per_day_minor, rent_price_minor, deposit_amount_minor,
                    created_at, created_by, updated_at, updated_by
                ) VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0, 0, ?, ?, ?, ?, ?)
                """,
                (
                    contract_id, row.cell_number, row.client_full_name,
                    row.id_card_number or LEGACY_MISSING_TEXT,
                    row.id_card_issuer or LEGACY_MISSING_TEXT,
                    (
                        row.id_card_issue_date.isoformat()
                        if row.id_card_issue_date is not None
                        else LEGACY_MISSING_DATE
                    ),
                    row.account_number or LEGACY_MISSING_TEXT,
                    legacy_extra_fields(
                        identity_complete=bool(
                            row.id_card_number
                            and row.id_card_issuer
                            and row.id_card_issue_date is not None
                            and row.account_number
                        ),
                        deposit_known=row.deposit_amount_minor is not None,
                    ),
                    row.start_date.isoformat(), row.end_date.isoformat(), rent_days,
                    row.deposit_amount_minor or 0,
                    created_at, IMPORT_ACTOR, timestamp, IMPORT_ACTOR,
                ),
            )
            connection.execute(
                """
                INSERT INTO archive.log(
                    log_id, operation_id, occurred_at, employee, action,
                    contract_id, cell_number, changes_json
                ) VALUES(?, ?, ?, ?, 'contract.created', ?, ?, ?)
                """,
                (
                    str(uuid4()), row_operation, created_at, IMPORT_ACTOR,
                    contract_id, row.cell_number,
                    json.dumps(
                        {
                            "source": "legacy_register",
                            "start_date": row.start_date.isoformat(),
                            "end_date": row.end_date.isoformat(),
                            "rent_days": rent_days,
                            "financial_terms_known": False,
                        },
                        sort_keys=True, separators=(",", ":"),
                    ),
                ),
            )
        elif row.kind == "manual":
            assert row.occupation_label
            connection.execute(
                """
                INSERT INTO main.cell_blocks(
                    cell_number, block_kind, source_contract_id,
                    occupation_label, created_at, created_by
                ) VALUES(?, 'manual', NULL, ?, ?, ?)
                """,
                (row.cell_number, row.occupation_label, timestamp, IMPORT_ACTOR),
            )
            connection.execute(
                """
                INSERT INTO archive.log(
                    log_id, operation_id, occurred_at, employee, action,
                    contract_id, cell_number, changes_json
                ) VALUES(?, ?, ?, ?, 'cell.manual_occupied', NULL, ?, ?)
                """,
                (
                    str(uuid4()), row_operation, timestamp, IMPORT_ACTOR,
                    row.cell_number,
                    json.dumps(
                        {
                            "source": "legacy_register",
                            "occupation_label": row.occupation_label,
                        },
                        ensure_ascii=False,
                        sort_keys=True,
                        separators=(",", ":"),
                    ),
                ),
            )
        if after_insert is not None:
            after_insert(row)
    connection.execute(
        """
        INSERT INTO archive.log(
            log_id, operation_id, occurred_at, employee, action,
            contract_id, cell_number, changes_json
        ) VALUES(?, ?, ?, ?, ?, NULL, NULL, ?)
        """,
        (
            str(uuid4()), operation_id, timestamp, IMPORT_ACTOR, SUMMARY_ACTION,
            json.dumps(
                {
                    "source": "legacy_register",
                    "sha256": plan.sha256,
                    "contracts_count": plan.contracts_count,
                    "free_count": plan.free_count,
                    "manual_count": plan.manual_count,
                },
                sort_keys=True, separators=(",", ":"),
            ),
        ),
    )


def import_legacy_contracts(
    settings: Settings,
    *,
    content: bytes,
    file_name: object,
    expected_sha256: object,
    confirmation: object,
    operation_id: object,
    occurred_at: datetime,
    after_insert: Callable[[LegacyImportRow], None] | None = None,
) -> LegacyImportResult:
    if not isinstance(file_name, str) or not file_name.casefold().endswith(".xlsx"):
        raise LegacyImportValidationError("Выберите файл формата XLSX.")
    if confirmation != IMPORT_CONFIRMATION:
        raise LegacyImportValidationError("Требуется явное подтверждение переноса.")
    normalized_operation = _operation_id(operation_id)
    if occurred_at.tzinfo is None or occurred_at.utcoffset() is None:
        raise LegacyImportValidationError("Время переноса должно содержать часовой пояс.")
    plan = _parse_workbook(content)
    if not isinstance(expected_sha256, str) or expected_sha256 != plan.sha256:
        raise LegacyImportConflictError("Файл изменился после проверки. Проверьте его заново.")
    if plan.issues:
        raise LegacyImportValidationError("В файле остались ошибки. Выполните проверку заново.")
    phase = "opening"
    try:
        with open_write(settings, attach_archive=True) as connection:
            previous = connection.execute(
                "SELECT * FROM archive.log WHERE operation_id=? AND action=?",
                (normalized_operation, SUMMARY_ACTION),
            ).fetchone()
            if previous is not None:
                changes = json.loads(str(previous["changes_json"]))
                if changes.get("sha256") != plan.sha256:
                    raise LegacyImportConflictError("Идентификатор операции уже использован для другого файла.")
                return _result_from_summary(previous, settings, repeated=True)
            initial_issues = _database_issues(
                connection, connection, plan, as_of_date=occurred_at.date()
            )
            if initial_issues:
                raise LegacyImportConflictError(initial_issues[0].message)
            phase = "backup-before"
            create_backup_pair(
                connection,
                settings,
                operation_id=str(uuid5(UUID(normalized_operation), "before-import")),
                occurred_at=occurred_at,
            )
            phase = "begin"
            connection.execute("BEGIN IMMEDIATE")
            phase = "transaction"
            transaction_issues = _database_issues(
                connection, connection, plan, as_of_date=occurred_at.date()
            )
            if transaction_issues:
                raise LegacyImportConflictError(transaction_issues[0].message)
            _insert_plan(
                connection,
                plan,
                operation_id=normalized_operation,
                occurred_at=occurred_at,
                after_insert=after_insert,
            )
            phase = "committing"
            connection.commit()
            phase = "verifying"
            summary = connection.execute(
                "SELECT * FROM archive.log WHERE operation_id=? AND action=?",
                (normalized_operation, SUMMARY_ACTION),
            ).fetchone()
            contract_count = int(connection.execute("SELECT COUNT(*) FROM main.contracts").fetchone()[0])
            manual_count = int(connection.execute("SELECT COUNT(*) FROM main.cell_blocks").fetchone()[0])
            if summary is None or contract_count != plan.contracts_count or manual_count != plan.manual_count:
                raise LegacyImportWriteUncertainError(UNCERTAIN_MESSAGE)
            phase = "backup-after"
            backup_created = True
            warning = None
            try:
                create_backup_pair(
                    connection,
                    settings,
                    operation_id=normalized_operation,
                    occurred_at=occurred_at,
                )
            except Exception:
                backup_created = False
                warning = (
                    "Договоры перенесены, но итоговую резервную копию создать не удалось. "
                    "Не повторяйте перенос; сообщите администратору."
                )
            return LegacyImportResult(
                normalized_operation,
                plan.contracts_count,
                plan.free_count,
                plan.manual_count,
                False,
                backup_created,
                warning,
            )
    except (LegacyImportValidationError, LegacyImportConflictError, LegacyImportWriteUncertainError):
        raise
    except DatabaseCorruptionError as exc:
        raise LegacyImportNetworkError(str(exc)) from exc
    except DatabaseUnavailableError as exc:
        raise LegacyImportNetworkError(NETWORK_ERROR_MESSAGE) from exc
    except BackupBusyError as exc:
        raise LegacyImportBusyError(BUSY_MESSAGE) from exc
    except (BackupNetworkError, BackupError) as exc:
        raise LegacyImportNetworkError(str(exc)) from exc
    except sqlite3.IntegrityError as exc:
        raise LegacyImportWriteError("Перенос не выполнен. Все изменения отменены.") from exc
    except (OSError, sqlite3.OperationalError) as exc:
        lowered = str(exc).casefold()
        if phase in {"opening", "backup-before", "begin", "transaction"} and ("locked" in lowered or "busy" in lowered or "занята" in lowered):
            raise LegacyImportBusyError(BUSY_MESSAGE) from exc
        if phase in {"committing", "verifying"}:
            raise LegacyImportWriteUncertainError(UNCERTAIN_MESSAGE) from exc
        raise LegacyImportNetworkError(NETWORK_ERROR_MESSAGE) from exc
    except Exception as exc:
        if phase in {"committing", "verifying"}:
            raise LegacyImportWriteUncertainError(UNCERTAIN_MESSAGE) from exc
        raise LegacyImportWriteError("Перенос не выполнен. Все изменения отменены.") from exc
