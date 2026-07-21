"""Open the packaged EXE against a disposable database with synthetic data."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
from uuid import uuid4


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from app.config import Settings
from app.db.connections import open_write
from app.db.schema import initialize_databases
from app.services.employee import (
    EmployeeRecord,
    EMPLOYEES_CONFIG_KEY,
    encode_employee_config,
)
from docx import Document


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--exe",
        type=Path,
        default=PROJECT_ROOT / "dist" / "safe-cells-portable" / "safe-cells.exe",
    )
    parser.add_argument("--ready-file", type=Path, required=True)
    args = parser.parse_args()
    executable = args.exe.resolve()
    ready_file = args.ready_file.resolve()
    if not executable.is_file():
        parser.error("safe-cells.exe не найден; сначала выполните build_portable.ps1")
    ready_file.unlink(missing_ok=True)

    with tempfile.TemporaryDirectory(
        prefix="safe-cells-stage12-browser-", ignore_cleanup_errors=True
    ) as temporary:
        root = Path(temporary)
        portable = root / "portable"
        shared = root / "shared"
        portable.mkdir()
        shared.mkdir()
        copied_executable = portable / "safe-cells.exe"
        shutil.copy2(executable, copied_executable)

        settings = Settings(database_directory=shared, testing=True)
        initialize_databases(
            settings,
            cells_csv_path=PROJECT_ROOT / "data" / "cell_heights.csv",
        )
        templates = shared / "templates"
        templates.mkdir()
        document = Document()
        document.add_heading("ТЕСТОВЫЙ ДОКУМЕНТ", level=1)
        document.add_paragraph("Ячейка: [Сейф.Номер]")
        document.save(templates / "acceptance-opening.docx")

        employee = EmployeeRecord(str(uuid4()), "Тестовый Сотрудник", True)
        with open_write(settings) as connection:
            connection.execute("BEGIN IMMEDIATE")
            connection.execute(
                "UPDATE config SET value = ? WHERE key = ?",
                (encode_employee_config([employee]), EMPLOYEES_CONFIG_KEY),
            )
            connection.execute(
                "INSERT INTO document_templates VALUES(?, 'opening', ?, ?, ?, 1, ?, ?)",
                (
                    "acceptance-opening",
                    "ТЕСТОВЫЙ ДОКУМЕНТ",
                    "acceptance-opening.docx",
                    json.dumps(["Сейф.Номер"], ensure_ascii=False),
                    "2026-07-01T09:00:00+06:00",
                    "test-setup",
                ),
            )
            connection.commit()
        (portable / "config.json").write_text(
            json.dumps({"database_directory": str(shared)}),
            encoding="utf-8",
        )

        creation_flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
        process = subprocess.Popen(
            [
                str(copied_executable),
                "--no-browser",
                "--ready-file",
                str(ready_file),
            ],
            cwd=portable,
            creationflags=creation_flags,
        )
        try:
            return process.wait()
        finally:
            if process.poll() is None:
                process.terminate()
                process.wait(timeout=5)
            ready_file.unlink(missing_ok=True)


if __name__ == "__main__":
    raise SystemExit(main())
