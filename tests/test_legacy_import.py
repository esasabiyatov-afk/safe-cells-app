from __future__ import annotations

from datetime import date, datetime, timedelta, timezone
from io import BytesIO
import json
from pathlib import Path
import sqlite3
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.utils.datetime import to_excel
from docx import Document
import pytest

from app import create_app
from app.config import Settings
from app.db.connections import open_readonly, open_write
from app.services.backups import list_backup_sets
from app.services.closures import ClosureConflictError, calculate_closure_quote
from app.services.contract_details import get_private_contract_details
from app.services.documents import (
    DocumentConflictError,
    DocumentValidationError,
    generate_event_documents,
    list_active_templates,
)
from app.services.editing import edit_contract
from app.services.legacy_contracts import legacy_status
from app.services.legacy_import import (
    IMPORT_CONFIRMATION,
    LegacyImportBusyError,
    LegacyImportConflictError,
    LegacyImportNetworkError,
    LegacyImportValidationError,
    LegacyImportWriteError,
    import_legacy_contracts,
    preview_legacy_import,
)
from app.services.renewals import RenewalConflictError, calculate_renewal_quote


OCCURRED_AT = datetime(2026, 7, 22, 9, 0, tzinfo=timezone(timedelta(hours=6)))
HEADERS = (
    "№",
    "Размер",
    "0 - свободен",
    "Ф.И.О. Клиента",
    "дата открытия",
    "Срок окончания",
)
IDENTITY_HEADERS = (
    "Серия и номер паспорта",
    "Кем выдан",
    "Дата выдачи паспорта",
    "Номер счёта",
)


def _cell_numbers(settings: Settings) -> list[str]:
    with open_readonly(settings.database_directory / settings.working_database_name) as connection:
        return [str(row[0]) for row in connection.execute(
            "SELECT number FROM cells ORDER BY CAST(number AS INTEGER), number"
        ).fetchall()]


def _workbook(
    settings: Settings,
    occupied: dict[str, tuple[str, object, object]] | None = None,
    manual: dict[str, str] | None = None,
    identity: dict[str, tuple[object, object, object, object]] | None = None,
    second_sheet_identity: list[
        tuple[str, object, object, object, object]
    ] | None = None,
) -> bytes:
    occupied = occupied or {}
    manual = manual or {}
    identity = identity or {}
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Лист2"
    sheet.append((*HEADERS, *IDENTITY_HEADERS) if identity else HEADERS)
    for number in _cell_numbers(settings):
        if number in occupied:
            name, start, end = occupied[number]
            row = (number, "100×220×330", 1, name, start, end)
            sheet.append((*row, *identity.get(number, (None, None, None, None))))
        elif number in manual:
            row = (number, "100×220×330", 1, manual[number], None, None)
            sheet.append((*row, *identity.get(number, (None, None, None, None))))
        else:
            row = (number, "100×220×330", 0, None, None, None)
            sheet.append((*row, *identity.get(number, (None, None, None, None))))
    if second_sheet_identity is not None:
        second = workbook.create_sheet("Паспортные данные")
        second.append(("№", "Ф.И.О. Клиента", "срок выдачи", "орган", "паспорт", None))
        for index, (name, issued_on, issuer, series, number) in enumerate(
            second_sheet_identity, start=1
        ):
            second.append((index, name, issued_on, issuer, series, number))
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _import(settings: Settings, content: bytes, **kwargs):
    return import_legacy_contracts(
        settings,
        content=content,
        file_name="old-register.xlsx",
        expected_sha256=preview_legacy_import(
            settings,
            content=content,
            file_name="old-register.xlsx",
            as_of_date=OCCURRED_AT.date(),
        ).sha256,
        confirmation=IMPORT_CONFIRMATION,
        operation_id=str(uuid4()),
        occurred_at=OCCURRED_AT,
        **kwargs,
    )


def test_preview_parses_mixed_dates_without_exposing_client_names(
    settings: Settings, initialized_databases,
) -> None:
    clients = {
        "1": ("Тестовый Клиент Один", "20 октября2025г.", "31 декабря 2026 г."),
        "2": ("Тестовый Клиент Два", "Monday, June 23, 2025", "Friday, July 31, 2026"),
        "3": ("Тестовый Клиент Три", to_excel(date(2025, 5, 2)), to_excel(date(2026, 8, 2))),
    }
    content = _workbook(settings, clients, {"4": "Служебная пометка"})

    preview = preview_legacy_import(
        settings,
        content=content,
        file_name="report.xlsx",
        as_of_date=OCCURRED_AT.date(),
    )

    public = json.dumps(preview.to_dict(), ensure_ascii=False)
    assert preview.ready is True
    assert (preview.contracts_count, preview.manual_count, preview.free_count) == (3, 1, 122)
    assert all(name not in public for name, _start, _end in clients.values())


def test_preview_blocks_missing_and_impossible_dates_without_writes(
    settings: Settings, initialized_databases,
) -> None:
    content = _workbook(
        settings,
        {
            "1": ("Тестовый Клиент Один", None, "31.12.2026"),
            "2": ("Тестовый Клиент Два", "01.01.2026", "31 сентября 2026"),
        },
    )

    preview = preview_legacy_import(
        settings,
        content=content,
        file_name="report.xlsx",
        as_of_date=OCCURRED_AT.date(),
    )

    assert preview.ready is False
    assert {(issue.cell_number, issue.message) for issue in preview.issues} == {
        ("1", "Не указана дата начала аренды."),
        ("2", "Некорректная дата окончания аренды."),
    }
    with open_readonly(settings.database_directory / settings.working_database_name) as connection:
        assert connection.execute("SELECT COUNT(*) FROM contracts").fetchone()[0] == 0


def test_successful_import_is_atomic_audited_and_backed_up_without_pii_in_metadata(
    settings: Settings, initialized_databases,
) -> None:
    client_name = "Тестовый Клиент Импорта"
    content = _workbook(
        settings,
        {"1": (client_name, "01.02.2025", "31.12.2026")},
        {"2": "Служебная пометка"},
    )

    result = _import(settings, content)

    assert result.contracts_count == 1
    assert result.manual_count == 1
    assert result.free_count == 124
    assert result.backup_created is True
    with open_readonly(settings.database_directory / settings.working_database_name) as connection:
        contract = connection.execute("SELECT * FROM contracts WHERE cell_number='1'").fetchone()
        block = connection.execute("SELECT * FROM cell_blocks WHERE cell_number='2'").fetchone()
    assert contract is not None and block is not None
    assert int(contract["rent_days"]) == (date(2026, 12, 31) - date(2025, 2, 1)).days + 1
    assert int(contract["price_per_day_minor"]) == 0
    assert int(contract["rent_price_minor"]) == 0
    assert int(contract["deposit_amount_minor"]) == 0
    assert legacy_status(contract["extra_fields_json"]) == {
        "legacy_imported": True,
        "legacy_identity_complete": False,
        "legacy_deposit_known": False,
        "legacy_rent_terms_known": False,
    }
    with open_readonly(settings.database_directory / settings.archive_database_name) as connection:
        changes = "\n".join(str(row[0]) for row in connection.execute(
            "SELECT changes_json FROM log"
        ).fetchall())
    assert client_name not in changes
    backup_sets = list_backup_sets(settings)
    assert len([item for item in backup_sets if item.valid]) == 2
    manifests = "\n".join(
        (item.directory / "manifest.json").read_text(encoding="utf-8")
        for item in backup_sets
        if item.valid
    )
    assert client_name not in manifests


def test_import_combines_passport_parts_from_second_sheet_without_exposing_them(
    settings: Settings, initialized_databases,
) -> None:
    client_name = "Тестовый Клиент Дополнительного Листа"
    content = _workbook(
        settings,
        {"1": (client_name, "01.02.2025", "31.12.2026")},
        second_sheet_identity=[
            (client_name, "02.03.2020", "Тестовый орган", "ID ", "221331"),
            (client_name, "02.03.2020", "Тестовый орган", "ID", "221331"),
        ],
    )

    preview = preview_legacy_import(
        settings,
        content=content,
        file_name="report.xlsx",
        as_of_date=OCCURRED_AT.date(),
    )

    assert preview.ready is True
    assert preview.passport_details_count == 1
    assert preview.identity_complete_count == 0
    _import(settings, content)
    with open_readonly(settings.database_directory / settings.working_database_name) as connection:
        contract = connection.execute(
            "SELECT * FROM contracts WHERE cell_number='1'"
        ).fetchone()
    assert contract is not None
    assert contract["id_card_number"] == "ID221331"
    assert contract["id_card_issuer"] == "Тестовый орган"
    assert contract["id_card_issue_date"] == "2020-03-02"
    details = get_private_contract_details(
        settings,
        cell_number="1",
        contract_ref=str(contract["contract_id"]),
    )
    assert details.id_card_number == "ID221331"
    assert details.id_card_issuer == "Тестовый орган"
    assert details.id_card_issue_date == "2020-03-02"
    assert details.account_number == ""
    assert details.legacy_identity_complete is False
    with open_readonly(settings.database_directory / settings.archive_database_name) as connection:
        audit = "\n".join(
            str(row[0])
            for row in connection.execute("SELECT changes_json FROM log").fetchall()
        )
    assert "ID221331" not in audit
    assert "Тестовый орган" not in audit


def test_extended_import_removes_passport_spaces_and_marks_complete_identity(
    settings: Settings, initialized_databases,
) -> None:
    content = _workbook(
        settings,
        {"1": ("Тестовый Клиент", "01.02.2025", "31.12.2026")},
        identity={
            "1": (
                "ID 22 13 31",
                "Тестовый орган",
                "02.03.2020",
                "TEST-ACCOUNT-1",
            )
        },
    )

    preview = preview_legacy_import(
        settings,
        content=content,
        file_name="report.xlsx",
        as_of_date=OCCURRED_AT.date(),
    )

    assert preview.ready is True
    assert preview.identity_complete_count == 1
    _import(settings, content)
    with open_readonly(settings.database_directory / settings.working_database_name) as connection:
        contract = connection.execute(
            "SELECT * FROM contracts WHERE cell_number='1'"
        ).fetchone()
    assert contract["id_card_number"] == "ID221331"
    assert legacy_status(contract["extra_fields_json"])["legacy_identity_complete"] is True


def test_failure_during_import_rolls_back_every_contract_and_block(
    settings: Settings, initialized_databases,
) -> None:
    content = _workbook(
        settings,
        {
            "1": ("Тестовый Клиент Один", "01.01.2025", "01.01.2027"),
            "2": ("Тестовый Клиент Два", "01.01.2025", "01.01.2027"),
        },
    )
    calls = 0

    def fail_after_first(_row) -> None:
        nonlocal calls
        calls += 1
        if calls == 1:
            raise RuntimeError("simulated partial failure")

    with pytest.raises(LegacyImportWriteError, match="Все изменения отменены"):
        _import(settings, content, after_insert=fail_after_first)

    with open_readonly(settings.database_directory / settings.working_database_name) as connection:
        assert connection.execute("SELECT COUNT(*) FROM contracts").fetchone()[0] == 0
        assert connection.execute("SELECT COUNT(*) FROM cell_blocks").fetchone()[0] == 0
    with open_readonly(settings.database_directory / settings.archive_database_name) as connection:
        assert connection.execute("SELECT COUNT(*) FROM log WHERE action LIKE 'contract.%'").fetchone()[0] == 0


def test_digest_confirmation_and_existing_business_data_are_rechecked(
    settings: Settings, initialized_databases, insert_test_contract,
) -> None:
    content = _workbook(
        settings,
        {"1": ("Тестовый Клиент", "01.01.2025", "01.01.2027")},
    )
    digest = preview_legacy_import(
        settings, content=content, file_name="report.xlsx", as_of_date=OCCURRED_AT.date()
    ).sha256
    with pytest.raises(LegacyImportValidationError, match="подтверждение"):
        import_legacy_contracts(
            settings, content=content, file_name="report.xlsx", expected_sha256=digest,
            confirmation="", operation_id=str(uuid4()), occurred_at=OCCURRED_AT,
        )
    with pytest.raises(LegacyImportConflictError, match="Файл изменился"):
        import_legacy_contracts(
            settings, content=content, file_name="report.xlsx", expected_sha256="0" * 64,
            confirmation=IMPORT_CONFIRMATION, operation_id=str(uuid4()), occurred_at=OCCURRED_AT,
        )
    insert_test_contract(cell_number="3", start_date="2026-01-01", end_date="2027-01-01")
    preview = preview_legacy_import(
        settings, content=content, file_name="report.xlsx", as_of_date=OCCURRED_AT.date()
    )
    assert preview.ready is False
    assert any("уже есть договоры" in issue.message for issue in preview.database_issues)
    with pytest.raises(LegacyImportConflictError, match="уже есть договоры"):
        import_legacy_contracts(
            settings, content=content, file_name="report.xlsx", expected_sha256=digest,
            confirmation=IMPORT_CONFIRMATION, operation_id=str(uuid4()), occurred_at=OCCURRED_AT,
        )


def test_locked_and_unavailable_database_are_reported_without_local_replacement(
    settings: Settings, initialized_databases,
) -> None:
    content = _workbook(
        settings,
        {"1": ("Тестовый Клиент", "01.01.2025", "01.01.2027")},
    )
    short_settings = Settings(
        database_directory=settings.database_directory,
        busy_timeout_ms=50,
        testing=True,
    )
    digest = preview_legacy_import(
        short_settings, content=content, file_name="report.xlsx", as_of_date=OCCURRED_AT.date()
    ).sha256
    locker = sqlite3.connect(
        str(settings.database_directory / settings.working_database_name),
        isolation_level=None,
    )
    try:
        locker.execute("BEGIN IMMEDIATE")
        with pytest.raises(LegacyImportBusyError):
            import_legacy_contracts(
                short_settings, content=content, file_name="report.xlsx",
                expected_sha256=digest, confirmation=IMPORT_CONFIRMATION,
                operation_id=str(uuid4()), occurred_at=OCCURRED_AT,
            )
    finally:
        locker.rollback()
        locker.close()
    missing = settings.database_directory / settings.working_database_name
    missing.unlink()
    with pytest.raises(LegacyImportNetworkError):
        preview_legacy_import(
            settings, content=content, file_name="report.xlsx", as_of_date=OCCURRED_AT.date()
        )
    assert not missing.exists()


def test_legacy_contract_requires_real_details_before_renewal_or_closure(
    settings: Settings, initialized_databases, tmp_path: Path,
) -> None:
    content = _workbook(
        settings,
        {"1": ("Тестовый Клиент", "01.01.2025", "01.01.2027")},
    )
    _import(settings, content)
    with open_readonly(settings.database_directory / settings.working_database_name) as connection:
        contract_ref = str(connection.execute(
            "SELECT contract_id FROM contracts WHERE cell_number='1'"
        ).fetchone()[0])
    details = get_private_contract_details(settings, cell_number="1", contract_ref=contract_ref)
    assert details.id_card_number == ""
    assert details.account_number == ""
    assert details.deposit_amount is None
    with pytest.raises(RenewalConflictError, match="паспортные данные"):
        calculate_renewal_quote(
            settings, cell_number="1", contract_ref=contract_ref,
            renewal_date=OCCURRED_AT.date(), renewal_days_value=30,
        )
    with pytest.raises(ClosureConflictError, match="паспортные данные"):
        calculate_closure_quote(
            settings, cell_number="1", contract_ref=contract_ref,
            close_date=OCCURRED_AT.date(), reason_code="standard",
        )
    with pytest.raises(DocumentConflictError, match="исходная сумма и тариф неизвестны"):
        list_active_templates(settings, cell_number="1", contract_ref=contract_ref)

    edit_contract(
        settings,
        payload={
            "operation_id": str(uuid4()),
            "contract_ref": contract_ref,
            "cell_number": "1",
            "client_full_name": "Тестовый Клиент",
            "id_card_number": "TEST-ID-1",
            "id_card_issuer": "Тестовый орган",
            "id_card_issue_date": "2020-01-01",
            "account_number": "TEST-ACCOUNT-1",
            "deposit_amount": 500,
        },
        employee="Тестовый Сотрудник",
        occurred_at=OCCURRED_AT,
    )
    quote = calculate_renewal_quote(
        settings, cell_number="1", contract_ref=contract_ref,
        renewal_date=OCCURRED_AT.date(), renewal_days_value=30,
    )
    closure = calculate_closure_quote(
        settings, cell_number="1", contract_ref=contract_ref,
        close_date=OCCURRED_AT.date(), reason_code="standard",
    )
    assert quote.renewal_days == 30
    assert closure.deposit_amount == 500
    templates = settings.database_directory / "templates"
    templates.mkdir(exist_ok=True)
    template = Document()
    template.add_paragraph("[Сумма]")
    template.save(templates / "legacy-unsafe.docx")
    with open_write(settings, attach_archive=True) as connection:
        connection.execute("BEGIN IMMEDIATE")
        connection.execute(
            "INSERT INTO document_templates VALUES(?, ?, ?, ?, ?, 1, ?, ?)",
            (
                "legacy-unsafe", "renewal", "Тестовый документ",
                "legacy-unsafe.docx", json.dumps(["Сумма"], ensure_ascii=False),
                OCCURRED_AT.isoformat(), "Тестовый Сотрудник",
            ),
        )
        connection.execute(
            """INSERT INTO archive.renewals(
                   renewal_id, contract_id, cell_number, old_end_date,
                   renewal_date, new_start_date, new_end_date, renewal_days,
                   price_per_day_minor, renewal_price_minor, penalty_days,
                   penalty_rate_minor, penalty_amount_minor, created_at,
                   created_by, operation_id
               ) VALUES(?, ?, '1', '2027-01-01', '2026-07-22', '2027-01-02',
                        '2027-01-31', 30, 15, 450, 0, 15, 0, ?, ?, ?)""",
            (
                "legacy-renewal-doc", contract_ref,
                OCCURRED_AT.isoformat(), "Тестовый Сотрудник", str(uuid4()),
            ),
        )
        connection.commit()
    with pytest.raises(DocumentValidationError, match="исходную сумму"):
        generate_event_documents(
            settings,
            event_type="renewal",
            contract_ref=contract_ref,
            event_ref="legacy-renewal-doc",
            output_directory=tmp_path / "documents",
            employee="Тестовый Сотрудник",
        )
    with open_readonly(settings.database_directory / settings.archive_database_name) as connection:
        audit = str(connection.execute(
            "SELECT changes_json FROM log WHERE action='contract.edited'"
        ).fetchone()[0])
    assert "TEST-ID-1" not in audit
    assert "TEST-ACCOUNT-1" not in audit


def test_admin_endpoints_require_access_preview_then_explicitly_confirm(
    settings: Settings, initialized_databases,
) -> None:
    content = _workbook(
        settings,
        {"1": ("Тестовый Клиент", "01.01.2025", "01.01.2027")},
    )
    app = create_app(settings)
    app.config.update(
        EMPLOYEE_PROVIDER=lambda: "Тестовый Администратор",
        TIMESTAMP_PROVIDER=lambda: OCCURRED_AT,
    )
    client = app.test_client()
    assert client.post(
        "/api/admin/legacy-import/preview",
        data={"file": (BytesIO(content), "report.xlsx")},
    ).status_code == 401
    setup = client.post(
        "/api/admin/setup",
        json={
            "operation_id": str(uuid4()),
            "password": "TestAdmin-2026",
            "password_confirmation": "TestAdmin-2026",
        },
    )
    token = setup.get_json()["token"]
    headers = {"X-Safe-Cells-Admin-Token": token}
    preview = client.post(
        "/api/admin/legacy-import/preview",
        headers=headers,
        data={"file": (BytesIO(content), "report.xlsx")},
    )
    assert preview.status_code == 200
    assert preview.get_json()["ready"] is True
    confirmation = client.post(
        "/api/admin/legacy-import/confirm",
        headers=headers,
        data={
            "file": (BytesIO(content), "report.xlsx"),
            "expected_sha256": preview.get_json()["sha256"],
            "confirmation": preview.get_json()["confirmation"],
            "operation_id": str(uuid4()),
        },
    )
    assert confirmation.status_code == 201, confirmation.get_json()
    assert confirmation.get_json()["contracts_count"] == 1
