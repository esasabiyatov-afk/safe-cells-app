from __future__ import annotations

from datetime import date, datetime, timedelta, timezone
from io import BytesIO
import json
from pathlib import Path
from uuid import uuid4

from docx import Document
from docx.shared import Cm
import pytest

from app import create_app
from app.db.connections import open_write
from app.documents import DocumentPublishError, DocumentTemplateError, render_docx
from app.services.closures import close_contract
from app.services.documents import (
    GeneratedDocument,
    build_document_values,
    generate_active_contract_document,
    generate_event_documents,
)
from app.services.document_downloads import (
    DocumentDownloadStore,
    DownloadArtifact,
    publish_generated_documents,
)
from app.documents.values import (
    amount_in_words_ky, amount_in_words_ru,
    format_document_issue_date, format_kyrgyz_date,
    format_quoted_kyrgyz_date, format_quoted_kyrgyz_date_stem,
    format_quoted_russian_date, format_russian_date,
)


def _template(path: Path, *, unknown: bool = False) -> None:
    document = Document()
    paragraph = document.add_paragraph("ТЕСТОВЫЙ ШАБЛОН — не банковский документ\nКлиент: ")
    paragraph.add_run("{{CLIENT_")
    paragraph.add_run("FULL_NAME}}")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Ячейка"
    table.cell(0, 1).text = "{{SAFE_NUMBER}}"
    section = document.sections[0]
    section.header.paragraphs[0].text = "Дата: {{CREATION_DATE}}"
    section.footer.paragraphs[0].text = "{{UNKNOWN_FIELD}}" if unknown else "ТЕСТ"
    document.save(path)


def test_renderer_replaces_split_runs_tables_and_header(tmp_path: Path):
    templates = tmp_path / "templates"
    downloads = tmp_path / "downloads"
    templates.mkdir()
    _template(templates / "test.docx")

    result = render_docx(
        template_directory=templates, template_file_name="test.docx",
        output_directory=downloads, output_file_name="result.docx",
        values={"CLIENT_FULL_NAME": "Вымышленный Клиент", "SAFE_NUMBER": "41", "CREATION_DATE": "2026-07-13"},
        required_placeholders=["CLIENT_FULL_NAME", "SAFE_NUMBER", "CREATION_DATE"],
    )

    rendered = Document(result)
    assert "Вымышленный Клиент" in rendered.paragraphs[0].text
    assert rendered.tables[0].cell(0, 1).text == "41"
    assert "2026-07-13" in rendered.sections[0].header.paragraphs[0].text
    assert not list(downloads.glob(".safe-cells-*.docx"))


def test_renderer_preserves_document_page_margins(tmp_path: Path):
    templates = tmp_path / "templates"
    templates.mkdir()
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.1)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)
    document.add_paragraph("Клиент: [Клиент.ФИО]")
    document.save(templates / "margins.docx")
    saved_section = Document(templates / "margins.docx").sections[0]
    expected_margins = (
        saved_section.top_margin,
        saved_section.bottom_margin,
        saved_section.left_margin,
        saved_section.right_margin,
    )

    result = render_docx(
        template_directory=templates,
        template_file_name="margins.docx",
        output_directory=tmp_path / "downloads",
        output_file_name="result.docx",
        values={"Клиент.ФИО": "Вымышленный Клиент"},
        required_placeholders=["Клиент.ФИО"],
    )

    rendered_section = Document(result).sections[0]
    assert (
        rendered_section.top_margin,
        rendered_section.bottom_margin,
        rendered_section.left_margin,
        rendered_section.right_margin,
    ) == expected_margins


def test_renderer_rejects_missing_required_placeholder(tmp_path: Path):
    templates = tmp_path / "templates"
    templates.mkdir()
    _template(templates / "test.docx")
    with pytest.raises(DocumentTemplateError, match="ACCOUNT_NUMBER"):
        render_docx(
            template_directory=templates, template_file_name="test.docx",
            output_directory=tmp_path / "downloads", output_file_name="result.docx",
            values={"CLIENT_FULL_NAME": "Тест", "SAFE_NUMBER": "1", "CREATION_DATE": "2026-07-13"},
            required_placeholders=["ACCOUNT_NUMBER"],
        )


def test_renderer_rejects_unknown_placeholder_and_path_escape(tmp_path: Path):
    templates = tmp_path / "templates"
    templates.mkdir()
    _template(templates / "unknown.docx", unknown=True)
    with pytest.raises(DocumentTemplateError, match="UNKNOWN_FIELD"):
        render_docx(
            template_directory=templates, template_file_name="unknown.docx",
            output_directory=tmp_path / "downloads", output_file_name="result.docx",
            values={"CLIENT_FULL_NAME": "Тест", "SAFE_NUMBER": "1", "CREATION_DATE": "2026-07-13"},
            required_placeholders=[],
        )


def test_renderer_removes_temporary_file_after_publish_failure(tmp_path: Path, monkeypatch):
    templates = tmp_path / "templates"
    downloads = tmp_path / "downloads"
    templates.mkdir()
    _template(templates / "test.docx")
    monkeypatch.setattr("app.documents.renderer.os.replace", lambda *args: (_ for _ in ()).throw(OSError("test failure")))
    with pytest.raises(DocumentPublishError, match="подготовить новый документ"):
        render_docx(
            template_directory=templates, template_file_name="test.docx",
            output_directory=downloads, output_file_name="result.docx",
            values={"CLIENT_FULL_NAME": "Тест", "SAFE_NUMBER": "1", "CREATION_DATE": "2026-07-13"},
            required_placeholders=["CLIENT_FULL_NAME"],
        )
    assert not list(downloads.glob(".safe-cells-*.docx"))
    assert not (downloads / "result.docx").exists()
    with pytest.raises(DocumentTemplateError, match="внутри общей папки"):
        render_docx(
            template_directory=templates, template_file_name="..\\outside.docx",
            output_directory=tmp_path, output_file_name="result.docx", values={}, required_placeholders=[],
        )


def test_renderer_creates_new_file_without_modifying_source_template(tmp_path: Path):
    templates = tmp_path / "templates"
    templates.mkdir()
    source = templates / "source.docx"
    _template(source)
    source_before = source.read_bytes()

    result = render_docx(
        template_directory=templates,
        template_file_name=source.name,
        output_directory=tmp_path / "staging",
        output_file_name="filled.docx",
        values={
            "CLIENT_FULL_NAME": "Вымышленный Клиент",
            "SAFE_NUMBER": "41",
            "CREATION_DATE": "2026-07-13",
        },
        required_placeholders=["CLIENT_FULL_NAME", "SAFE_NUMBER", "CREATION_DATE"],
    )

    assert result != source
    assert result.is_file()
    assert source.read_bytes() == source_before


def test_document_download_store_is_one_time_and_expires():
    now = [100.0]
    store = DocumentDownloadStore(
        ttl_seconds=10,
        max_documents=2,
        time_provider=lambda: now[0],
    )
    first = store.publish(
        [DownloadArtifact(file_name="first.docx", content=b"PK-first")]
    )[0]

    claimed = store.claim(first.download_id)
    assert claimed == DownloadArtifact(file_name="first.docx", content=b"PK-first")
    assert store.claim(first.download_id) is None

    second = store.publish(
        [DownloadArtifact(file_name="second.docx", content=b"PK-second")]
    )[0]
    now[0] = 111.0
    assert store.claim(second.download_id) is None


def test_document_download_store_does_not_publish_partial_set_when_id_creation_fails(
    monkeypatch,
):
    store = DocumentDownloadStore()
    existing = store.publish(
        [DownloadArtifact(file_name="existing.docx", content=b"PK-existing")]
    )[0]
    calls = 0

    def fail_on_second_id(_length):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise OSError("test random source failure")
        return f"test-download-{calls}"

    monkeypatch.setattr(
        "app.services.document_downloads.token_urlsafe", fail_on_second_id
    )
    with pytest.raises(OSError, match="random source failure"):
        store.publish(
            [
                DownloadArtifact(file_name="one.docx", content=b"PK-one"),
                DownloadArtifact(file_name="two.docx", content=b"PK-two"),
            ]
        )

    assert store.claim("test-download-1") is None
    assert store.claim(existing.download_id) == DownloadArtifact(
        file_name="existing.docx", content=b"PK-existing"
    )


def test_generated_documents_are_published_only_as_complete_set(tmp_path: Path):
    staging = tmp_path / "staging"
    staging.mkdir()
    (staging / "one.docx").write_bytes(b"PK-one")
    store = DocumentDownloadStore()

    with pytest.raises(DocumentPublishError, match="передать"):
        publish_generated_documents(
            store,
            output_directory=staging,
            generated=[GeneratedDocument("one.docx"), GeneratedDocument("missing.docx")],
        )

    published = publish_generated_documents(
        store,
        output_directory=staging,
        generated=[GeneratedDocument("one.docx")],
    )
    assert len(published) == 1
    assert store.claim(published[0].download_id) == DownloadArtifact(
        file_name="one.docx", content=b"PK-one"
    )


def test_active_contract_document_and_private_endpoint(
    settings, initialized_databases, insert_test_contract, tmp_path: Path
):
    insert_test_contract(cell_number="41", start_date="2026-07-01", end_date="2026-07-30", client_name="Вымышленный Клиент")
    templates = settings.database_directory / "templates"
    templates.mkdir()
    _template(templates / "TEST_ONLY.docx")
    with open_write(settings) as connection:
        connection.execute("BEGIN IMMEDIATE")
        connection.execute(
            "INSERT INTO document_templates VALUES(?, ?, ?, ?, ?, 1, ?, ?)",
            ("test-template", "TEST_ONLY", "ТЕСТОВЫЙ-ДОКУМЕНТ", "TEST_ONLY.docx",
             json.dumps(["CLIENT_FULL_NAME", "SAFE_NUMBER", "CREATION_DATE"]),
             "2026-07-13T10:00:00+06:00", "test-user"),
        )
        connection.commit()
    downloads = tmp_path / "downloads"

    result = generate_active_contract_document(
        settings, cell_number="41", contract_ref="contract-test-41",
        template_id="test-template", output_directory=downloads,
        creation_date=date(2026, 7, 13), employee="test-user",
    )
    assert result.file_name.startswith("ТЕСТОВЫЙ-ДОКУМЕНТ_Ячейка-41_Вымышленный-Клиент_")
    assert (downloads / result.file_name).is_file()

    app = create_app(settings)
    endpoint_downloads = tmp_path / "endpoint-downloads"
    app.config.update(
        TODAY_PROVIDER=lambda: date(2026, 7, 13),
        DOWNLOADS_DIRECTORY_PROVIDER=lambda: endpoint_downloads,
        EMPLOYEE_PROVIDER=lambda: "Тестовый Сотрудник",
    )
    token = app.extensions["safe_cells_private_token"]
    templates_response = app.test_client().post(
        "/api/documents/templates", headers={"X-Safe-Cells-Token": token},
        json={"cell_number": "41", "contract_ref": "contract-test-41"},
    )
    assert templates_response.status_code == 200
    assert templates_response.get_json()["templates"] == [
        {"template_id": "test-template", "display_name": "ТЕСТОВЫЙ-ДОКУМЕНТ"}
    ]
    response = app.test_client().post(
        "/api/documents/generate", headers={"X-Safe-Cells-Token": token},
        json={"cell_number": "41", "contract_ref": "contract-test-41", "template_id": "test-template"},
    )
    assert response.status_code == 201
    payload = response.get_json()
    assert payload["message"] == "Документ готов к скачиванию в браузере."
    assert payload["documents"] == [
        {
            "download_id": payload["documents"][0]["download_id"],
            "file_name": payload["file_name"],
        }
    ]
    assert not endpoint_downloads.exists()
    download = app.test_client().post(
        "/api/documents/download",
        headers={"X-Safe-Cells-Token": token},
        json={"download_id": payload["documents"][0]["download_id"]},
    )
    assert download.status_code == 200
    assert download.mimetype == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "attachment" in download.headers["Content-Disposition"]
    downloaded = Document(BytesIO(download.data))
    assert "Вымышленный Клиент" in downloaded.paragraphs[0].text
    repeated_download = app.test_client().post(
        "/api/documents/download",
        headers={"X-Safe-Cells-Token": token},
        json={"download_id": payload["documents"][0]["download_id"]},
    )
    assert repeated_download.status_code == 410
    assert not list(settings.database_directory.glob("*.docx"))


def test_document_endpoint_does_not_disclose_without_token(settings, initialized_databases):
    app = create_app(settings)
    response = app.test_client().post("/api/documents/generate", json={})
    assert response.status_code == 403
    assert app.test_client().post(
        "/api/documents/download", json={"download_id": "test"}
    ).status_code == 403


def test_approved_date_and_deposit_formats():
    value = date(2026, 5, 22)
    assert format_russian_date(value) == "22 мая 2026 г."
    assert format_kyrgyz_date(value) == "22-май 2026-ж."
    assert format_document_issue_date(date(2017, 9, 12)) == "12.09.2017-ж/г."
    assert format_quoted_russian_date(value) == "«22» мая 2026 г."
    assert format_quoted_kyrgyz_date(value) == "«22» май 2026-ж."
    assert format_quoted_kyrgyz_date_stem(value) == "«22» май 2026"
    assert amount_in_words_ru(1500) == "Одна тысяча пятьсот"
    assert amount_in_words_ky(1500) == "Бир миң беш жүз"


def test_renderer_supports_split_bank_square_codes(tmp_path: Path):
    templates = tmp_path / "templates"
    templates.mkdir()
    document = Document()
    paragraph = document.add_paragraph("Клиент: ")
    split_start = paragraph.add_run("[Клиент.")
    split_start.bold = True
    paragraph.add_run("ФИО]")
    suffix = paragraph.add_run("; сумма [Сумма]")
    suffix.italic = True
    document.save(templates / "bank.docx")
    result = render_docx(
        template_directory=templates, template_file_name="bank.docx",
        output_directory=tmp_path / "downloads", output_file_name="result.docx",
        values={"Клиент.ФИО": "Вымышленный Клиент", "Сумма": 450},
        required_placeholders=["Клиент.ФИО", "Сумма"],
    )
    rendered = Document(result).paragraphs[0]
    assert rendered.text == "Клиент: Вымышленный Клиент; сумма 450"
    assert rendered.runs[1].bold is True
    assert rendered.runs[3].italic is True


def test_renewal_codes_keep_original_contract_start_and_specific_period(
    tmp_path: Path,
):
    contract = {
        "client_full_name": "Вымышленный Клиент",
        "id_card_number": "TEST-000000",
        "id_card_issuer": "Тестовый орган",
        "id_card_issue_date": "2017-09-12",
        "account_number": "TEST-ACCOUNT-001",
        "cell_number": "41",
        "height_mm": 50,
        "width_mm": 220,
        "depth_mm": 330,
        "start_date": "2026-07-12",
        "end_date": "2026-09-18",
        "rent_days": 30,
        "rent_price_minor": 450,
        "deposit_amount_minor": 1500,
    }
    renewal = {
        "new_start_date": "2026-08-12",
        "new_end_date": "2026-09-18",
        "renewal_days": 38,
        "renewal_price_minor": 380,
    }
    values = build_document_values(
        contract,
        creation_date=date(2026, 8, 12),
        employee="Тестовый Сотрудник",
        renewal=renewal,
    )
    assert values["Договор.НачалоД"] == "«12» июля 2026 г."
    assert values["Договор.НачалоДК"] == "«12» июль 2026"
    assert values["Продление.Начало"] == "«12» августа 2026 г."
    assert values["Продление.Конец"] == "«18» сентября 2026 г."
    assert values["Продление.НачалоК"] == "«12» август 2026-ж."
    assert values["Продление.КонецК"] == "«18» сентябрь 2026-ж."
    assert values["Продление.Сумма"] == 380
    assert values["Продление.Срок"] == 38

    templates = tmp_path / "templates"
    templates.mkdir()
    document = Document()
    document.add_paragraph(
        "[Договор.НачалоДК]-жылдагы | [Договор.НачалоД] | "
        "[Продление.Начало] | [Продление.Конец] | "
        "[Продление.НачалоК] | [Продление.КонецК] | "
        "[Продление.Сумма] | [Продление.Срок]"
    )
    document.save(templates / "renewal.docx")
    required = [
        "Договор.НачалоДК",
        "Договор.НачалоД",
        "Продление.Начало",
        "Продление.Конец",
        "Продление.НачалоК",
        "Продление.КонецК",
        "Продление.Сумма",
        "Продление.Срок",
    ]
    result = render_docx(
        template_directory=templates,
        template_file_name="renewal.docx",
        output_directory=tmp_path / "downloads",
        output_file_name="result.docx",
        values=values,
        required_placeholders=required,
    )
    text = Document(result).paragraphs[0].text
    assert "«12» июль 2026-жылдагы" in text
    assert "2026--жылдагы" not in text
    assert "[Продление." not in text


def test_opening_renewal_and_closing_document_bundles(
    settings, initialized_databases, insert_test_contract, tmp_path: Path
):
    insert_test_contract(
        cell_number="41",
        start_date="2026-07-01",
        end_date="2026-07-30",
        client_name="Вымышленный Клиент",
        account_number="TEST-ACCOUNT-041",
    )
    templates = settings.database_directory / "templates"
    templates.mkdir()
    definitions = (
        ("open-a", "opening", "ТЕСТ-АКТ", "open-a.docx", "Сейф.Номер"),
        ("open-b", "opening", "ТЕСТ-ДОГОВОР", "open-b.docx", "Договор.Начало"),
        ("open-c", "opening", "ТЕСТ-РАСПОРЯЖЕНИЕ", "open-c.docx", "Счет.Номер"),
        ("renew", "renewal", "ТЕСТ-ПРОДЛЕНИЕ", "renew.docx", "Продление.Начало"),
        ("close", "closing", "ТЕСТ-ЗАКРЫТИЕ", "close.docx", "Залог.Пропись"),
    )
    with open_write(settings, attach_archive=True) as connection:
        connection.execute("BEGIN IMMEDIATE")
        for template_id, event, display, file_name, placeholder in definitions:
            document = Document()
            document.add_paragraph(f"ТЕСТОВЫЙ ДОКУМЕНТ: [{placeholder}]")
            document.save(templates / file_name)
            connection.execute(
                "INSERT INTO document_templates VALUES(?, ?, ?, ?, ?, 1, ?, ?)",
                (
                    template_id,
                    event,
                    display,
                    file_name,
                    json.dumps([placeholder], ensure_ascii=False),
                    "2026-07-30T10:00:00+06:00",
                    "test-user",
                ),
            )
        connection.execute(
            """INSERT INTO archive.renewals(
                   renewal_id, contract_id, cell_number, old_end_date,
                   renewal_date, new_start_date, new_end_date, renewal_days,
                   price_per_day_minor, renewal_price_minor, penalty_days,
                   penalty_rate_minor, penalty_amount_minor, created_at,
                   created_by, operation_id
               ) VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                "renewal-test-docs",
                "contract-test-41",
                "41",
                "2026-07-30",
                "2026-07-30",
                "2026-07-31",
                "2026-08-29",
                30,
                15,
                450,
                0,
                15,
                0,
                "2026-07-30T10:00:00+06:00",
                "test-user",
                "renewal-operation-docs",
            ),
        )
        connection.commit()

    downloads = tmp_path / "downloads"
    opening = generate_event_documents(
        settings,
        event_type="opening",
        contract_ref="contract-test-41",
        event_ref=None,
        output_directory=downloads,
        employee="Тестовый Сотрудник",
    )
    renewal = generate_event_documents(
        settings,
        event_type="renewal",
        contract_ref="contract-test-41",
        event_ref="renewal-test-docs",
        output_directory=downloads,
        employee="Тестовый Сотрудник",
    )
    assert len(opening) == 3
    assert len(renewal) == 1

    operation_id = str(uuid4())
    close_contract(
        settings,
        payload={
            "operation_id": operation_id,
            "cell_number": "41",
            "contract_ref": "contract-test-41",
            "expected_end_date": "2026-07-30",
            "reason_code": "standard",
        },
        employee="test-user",
        close_date=date(2026, 7, 30),
        occurred_at=datetime(
            2026, 7, 30, 11, 0, tzinfo=timezone(timedelta(hours=6))
        ),
    )
    closing = generate_event_documents(
        settings,
        event_type="closing",
        contract_ref="contract-test-41",
        event_ref=operation_id,
        output_directory=downloads,
        employee="Тестовый Сотрудник",
    )
    assert len(closing) == 1
    assert len(list(downloads.glob("*.docx"))) == 5
    for path in downloads.glob("*.docx"):
        assert "[" not in "\n".join(paragraph.text for paragraph in Document(path).paragraphs)


def test_event_bundle_is_not_published_when_one_template_is_invalid(
    settings, initialized_databases, insert_test_contract, tmp_path: Path
):
    insert_test_contract(
        cell_number="42",
        end_date="2026-08-31",
        client_name="Вымышленный Клиент",
    )
    templates = settings.database_directory / "templates"
    templates.mkdir()
    valid = Document()
    valid.add_paragraph("ТЕСТОВЫЙ ДОКУМЕНТ: [Сейф.Номер]")
    valid.save(templates / "valid.docx")
    invalid = Document()
    invalid.add_paragraph("ТЕСТОВЫЙ ДОКУМЕНТ: [НЕИЗВЕСТНОЕ.ПОЛЕ]")
    invalid.save(templates / "invalid.docx")
    with open_write(settings) as connection:
        connection.execute("BEGIN IMMEDIATE")
        for template_id, display_name, file_name in (
            ("valid-opening", "А-ТЕСТ", "valid.docx"),
            ("invalid-opening", "Б-ТЕСТ", "invalid.docx"),
        ):
            connection.execute(
                "INSERT INTO document_templates VALUES(?, 'opening', ?, ?, ?, 1, ?, ?)",
                (
                    template_id,
                    display_name,
                    file_name,
                    json.dumps(["Сейф.Номер"], ensure_ascii=False),
                    "2026-07-30T10:00:00+06:00",
                    "test-user",
                ),
            )
        connection.commit()

    downloads = tmp_path / "downloads"
    with pytest.raises(DocumentTemplateError, match="обязательные поля"):
        generate_event_documents(
            settings,
            event_type="opening",
            contract_ref="contract-test-42",
            event_ref=None,
            output_directory=downloads,
            employee="Тестовый Сотрудник",
        )

    assert list(downloads.glob("*.docx")) == []
