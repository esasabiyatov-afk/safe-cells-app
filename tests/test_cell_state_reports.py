from __future__ import annotations

from datetime import date
from io import BytesIO
import json
from pathlib import Path
from urllib.parse import unquote
from uuid import uuid4

from openpyxl import load_workbook
import pytest

from app import create_app
from app.config import Settings
from app.db.connections import NETWORK_ERROR_MESSAGE, open_write
from app.services.cell_state_reports import (
    CellStateReportReadError,
    CellStateReportValidationError,
    build_cell_state_report,
    list_cell_states_on_date,
    validate_report_date,
)


def _insert_log(
    connection,
    *,
    operation_id: str,
    occurred_at: str,
    action: str,
    cell_number: str | None,
    contract_id: str | None = None,
    changes: dict[str, object] | None = None,
) -> None:
    connection.execute(
        """
        INSERT INTO archive.log(
            log_id, operation_id, occurred_at, employee, action,
            contract_id, cell_number, changes_json
        ) VALUES(?, ?, ?, 'Тестовый сотрудник', ?, ?, ?, ?)
        """,
        (
            str(uuid4()),
            operation_id,
            occurred_at,
            action,
            contract_id,
            cell_number,
            json.dumps(changes or {}, ensure_ascii=False),
        ),
    )


def _row(rows: list[dict[str, object]], number: str) -> dict[str, object]:
    return next(row for row in rows if row["cell_number"] == number)


def test_cell_state_snapshot_reconstructs_contract_renewal_and_special_states(
    settings: Settings,
    initialized_databases,
    insert_test_contract,
) -> None:
    insert_test_contract(
        cell_number="7",
        start_date="2026-07-01",
        end_date="2026-07-30",
        client_name="Тестовый Клиент Семь",
    )
    insert_test_contract(
        cell_number="9",
        start_date="2026-07-01",
        end_date="2026-07-31",
        client_name="Тестовый Клиент Девять",
    )
    insert_test_contract(
        cell_number="10",
        start_date="2026-07-05",
        end_date="2026-07-20",
        client_name="=НЕ ФОРМУЛА",
    )
    with open_write(settings, attach_archive=True) as connection:
        connection.execute("BEGIN IMMEDIATE")
        _insert_log(
            connection,
            operation_id="opening-7",
            occurred_at="2026-07-01T09:00:00+06:00",
            action="contract.created",
            contract_id="contract-test-7",
            cell_number="7",
        )
        connection.execute(
            """
            INSERT INTO archive.renewals(
                renewal_id, contract_id, cell_number, old_end_date,
                renewal_date, new_start_date, new_end_date, renewal_days,
                price_per_day_minor, renewal_price_minor, penalty_days,
                penalty_rate_minor, penalty_amount_minor, created_at,
                created_by, operation_id
            ) VALUES(
                'renewal-7', 'contract-test-7', '7', '2026-07-30',
                '2026-07-10', '2026-07-31', '2026-08-30', 31,
                15, 465, 0, 15, 0, '2026-07-10T10:00:00+06:00',
                'Тестовый сотрудник', 'renewal-operation-7'
            )
            """
        )
        connection.execute(
            "UPDATE main.contracts SET end_date='2026-08-30' WHERE cell_number='7'"
        )
        _insert_log(
            connection,
            operation_id="manual-open-8",
            occurred_at="2026-07-02T09:00:00+06:00",
            action="cell.manual_occupied",
            cell_number="8",
            changes={"occupation_label": "Тестовая касса"},
        )
        _insert_log(
            connection,
            operation_id="manual-close-8",
            occurred_at="2026-07-04T09:00:00+06:00",
            action="cell.manual_released",
            cell_number="8",
            changes={"occupation_label": "Тестовая касса"},
        )
        _insert_log(
            connection,
            operation_id="opening-9",
            occurred_at="2026-07-01T09:10:00+06:00",
            action="contract.created",
            contract_id="contract-test-9",
            cell_number="9",
        )
        _insert_log(
            connection,
            operation_id="close-lost-9",
            occurred_at="2026-07-03T11:00:00+06:00",
            action="contract.closed",
            contract_id="contract-test-9",
            cell_number="9",
            changes={"close_reason": "Потеря ключа"},
        )
        _insert_log(
            connection,
            operation_id="restore-key-9",
            occurred_at="2026-07-06T11:00:00+06:00",
            action="cell.key_restored",
            contract_id="contract-test-9",
            cell_number="9",
        )
        _insert_log(
            connection,
            operation_id="opening-cancelled-10",
            occurred_at="2026-07-05T12:00:00+06:00",
            action="contract.created",
            contract_id="contract-test-10",
            cell_number="10",
        )
        connection.execute(
            """
            INSERT INTO archive.operation_cancellations(
                cancellation_id, cancellation_operation_id,
                original_operation_id, original_action, contract_id,
                cell_number, reason_code, cancelled_at, cancelled_by
            ) VALUES(
                'cancel-10', 'cancel-operation-10', 'opening-cancelled-10',
                'contract.created', 'contract-test-10', '10', 'input_error',
                '2026-07-05T12:30:00+06:00', 'Тестовый сотрудник'
            )
            """
        )
        connection.execute(
            """
            INSERT INTO main.cells(
                number, height_mm, width_mm, depth_mm, is_active
            ) VALUES('200', 150, 250, 350, 1)
            """
        )
        _insert_log(
            connection,
            operation_id="cell-created-200",
            occurred_at="2026-07-05T08:00:00+06:00",
            action="admin.cell.created",
            cell_number="200",
            changes={"height_mm": 150},
        )
        _insert_log(
            connection,
            operation_id="cell-retired-200",
            occurred_at="2026-07-10T08:00:00+06:00",
            action="admin.cell.retired",
            cell_number="200",
            changes={"reason": "Тестовый вывод"},
        )
        connection.commit()

    before_renewal = list_cell_states_on_date(
        settings, as_of_date=date(2026, 7, 3)
    )
    after_renewal = list_cell_states_on_date(
        settings, as_of_date=date(2026, 7, 10)
    )
    before_new_cell = list_cell_states_on_date(
        settings, as_of_date=date(2026, 7, 4)
    )
    new_cell_visible = list_cell_states_on_date(
        settings, as_of_date=date(2026, 7, 5)
    )

    assert len(before_new_cell) == 126
    contract_row = _row(before_renewal, "7")
    assert contract_row["occupied"] is True
    assert contract_row["occupancy_code"] == 1
    assert contract_row["client_full_name"] == "Тестовый Клиент Семь"
    assert contract_row["start_date"] == "2026-07-01"
    assert contract_row["last_renewal_date"] is None
    assert contract_row["end_date"] == "2026-07-30"
    assert contract_row["key_count"] == 1
    assert _row(after_renewal, "7")["last_renewal_date"] == "2026-07-10"
    assert _row(after_renewal, "7")["end_date"] == "2026-08-30"
    assert _row(before_renewal, "8")["client_full_name"] == "Тестовая касса"
    assert _row(new_cell_visible, "8")["occupied"] is False
    assert _row(before_new_cell, "9")["client_full_name"] == (
        "Ключ утерян — Тестовый Клиент Девять"
    )
    assert _row(after_renewal, "9")["occupied"] is False
    assert _row(new_cell_visible, "10")["occupied"] is False
    assert all(row["cell_number"] != "200" for row in before_new_cell)
    assert _row(new_cell_visible, "200")["occupied"] is False
    assert all(row["cell_number"] != "200" for row in after_renewal)


def test_cancelled_renewal_keeps_previous_end_date(
    settings: Settings,
    initialized_databases,
    insert_test_contract,
) -> None:
    insert_test_contract(
        cell_number="11",
        start_date="2026-07-01",
        end_date="2026-08-31",
        client_name="Тестовый Клиент Одиннадцать",
    )
    with open_write(settings, attach_archive=True) as connection:
        connection.execute("BEGIN IMMEDIATE")
        _insert_log(
            connection,
            operation_id="opening-11",
            occurred_at="2026-07-01T09:00:00+06:00",
            action="contract.created",
            contract_id="contract-test-11",
            cell_number="11",
        )
        connection.execute(
            """
            INSERT INTO archive.renewals(
                renewal_id, contract_id, cell_number, old_end_date,
                renewal_date, new_start_date, new_end_date, renewal_days,
                price_per_day_minor, renewal_price_minor, penalty_days,
                penalty_rate_minor, penalty_amount_minor, created_at,
                created_by, operation_id
            ) VALUES(
                'renewal-11', 'contract-test-11', '11', '2026-07-31',
                '2026-07-08', '2026-08-01', '2026-08-31', 31,
                15, 465, 0, 15, 0, '2026-07-08T10:00:00+06:00',
                'Тестовый сотрудник', 'renewal-operation-11'
            )
            """
        )
        connection.execute(
            """
            INSERT INTO archive.operation_cancellations(
                cancellation_id, cancellation_operation_id,
                original_operation_id, original_action, contract_id,
                cell_number, reason_code, cancelled_at, cancelled_by
            ) VALUES(
                'cancel-11', 'cancel-operation-11', 'renewal-operation-11',
                'contract.renewed', 'contract-test-11', '11', 'change_term',
                '2026-07-08T10:30:00+06:00', 'Тестовый сотрудник'
            )
            """
        )
        connection.commit()

    row = _row(
        list_cell_states_on_date(settings, as_of_date=date(2026, 7, 8)),
        "11",
    )

    assert row["occupied"] is True
    assert row["end_date"] == "2026-07-31"
    assert row["last_renewal_date"] is None


def test_cell_state_xlsx_matches_approved_layout_and_neutralizes_formulas() -> None:
    report = build_cell_state_report(
        [
            {
                "cell_number": "=1+1",
                "height_mm": 60,
                "width_mm": 220,
                "depth_mm": 300,
                "occupied": True,
                "occupancy_code": 1,
                "client_full_name": "=HYPERLINK(\"unsafe\")",
                "start_date": "2026-05-01",
                "last_renewal_date": None,
                "end_date": "2026-05-18",
                "key_count": 1,
            },
            {
                "cell_number": "2",
                "height_mm": 100,
                "width_mm": 220,
                "depth_mm": 300,
                "occupied": False,
                "occupancy_code": 0,
                "client_full_name": "",
                "start_date": None,
                "last_renewal_date": None,
                "end_date": None,
                "key_count": 2,
            },
        ],
        as_of_date=date(2026, 5, 18),
    )

    worksheet = load_workbook(BytesIO(report), data_only=False)["18.05.26"]

    assert worksheet["A1"].value == "Отдел депозитария"
    assert worksheet["E1"].value == "ЗАО АКБ «Толубай»"
    assert worksheet["A2"].value == (
        "Состояние всех ячеек на 18 мая 2026 года"
    )
    assert tuple(cell.value for cell in worksheet[3]) == (
        "№",
        "Размер, мм\n(В × Ш × Г)",
        "0 — свободна\n1 — занята",
        "Ф.И.О. клиента / пометка",
        "Дата открытия",
        "Дата последнего продления",
        "Срок окончания",
        "Количество ключей",
    )
    assert worksheet["A4"].value == "'=1+1"
    assert worksheet["D4"].value.startswith("'=")
    assert worksheet["E4"].value == "01 мая 2026 г."
    assert worksheet["F4"].value == "—"
    assert worksheet["G4"].value == "18 мая 2026 г."
    assert worksheet["C4"].value == 1
    assert worksheet["H4"].value == 1
    assert worksheet["C5"].value == 0
    assert worksheet["H5"].value == 2
    assert worksheet["B4"].fill.fgColor.rgb == "00FFF2CC"
    assert worksheet["B5"].fill.fgColor.rgb == "00FFFFFF"


@pytest.mark.parametrize(
    ("value", "message"),
    [
        ("", "Укажите дату"),
        ("18.05.2026", "корректную дату"),
        ("2026-07-30", "будущую дату"),
    ],
)
def test_cell_state_report_date_validation(value: str, message: str) -> None:
    with pytest.raises(CellStateReportValidationError, match=message):
        validate_report_date(value, today=date(2026, 7, 29))


def test_cell_state_report_api_requires_employee_and_downloads_xlsx(
    settings: Settings,
    initialized_databases,
) -> None:
    app = create_app(settings)
    app.config["TODAY_PROVIDER"] = lambda: date(2026, 7, 29)
    client = app.test_client()

    missing_employee = client.post(
        "/api/journal/cell-state-report",
        json={"as_of_date": "2026-07-29"},
    )
    assert missing_employee.status_code == 409

    app.config["EMPLOYEE_PROVIDER"] = lambda: "Тестовый сотрудник"
    response = client.post(
        "/api/journal/cell-state-report",
        json={"as_of_date": "2026-07-29"},
    )

    assert response.status_code == 200
    assert response.mimetype == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "Состояние_всех_ячеек_2026-07-29.xlsx" in unquote(
        response.headers["Content-Disposition"]
    )
    worksheet = load_workbook(BytesIO(response.data), read_only=True).active
    assert worksheet.max_row == 129
    assert client.post(
        "/api/journal/cell-state-report",
        json={"as_of_date": "2026-07-30"},
    ).status_code == 400
    assert client.post(
        "/api/journal/cell-state-report",
        json={"as_of_date": "2026-07-29", "unexpected": True},
    ).status_code == 400
    assert client.get(
        "/api/journal/cell-state-report?as_of_date=2026-07-29"
    ).status_code == 405


def test_cell_state_report_unavailable_path_does_not_create_database(
    tmp_path: Path,
) -> None:
    missing = tmp_path / "offline state report"
    settings = Settings(database_directory=missing, testing=True)

    with pytest.raises(CellStateReportReadError, match=NETWORK_ERROR_MESSAGE):
        list_cell_states_on_date(settings, as_of_date=date(2026, 7, 29))

    assert not missing.exists()
