from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from app import create_app
from app.config import Settings
from app.db.schema import initialize_databases
from app.services.statement_import import (
    MAX_STATEMENT_BYTES,
    StatementImportValidationError,
    extract_statement_data,
)


FAKE_NAME = "Вымышленный Тестовый Клиент"
FAKE_ID = "TEST 00 000001"
FAKE_ISSUER = "ТЕСТОВЫЙ ОРГАН 999-999"
FAKE_ACCOUNT = "9999000011112222"


def _statement_docx(
    *,
    issue_text: str = f"{FAKE_ISSUER} 12.09.2017 г.",
    account_text: str = (
        "17. Банковские продукты: Банковский сейф (Ячейки) № " + FAKE_ACCOUNT
    ),
    fio_label: str = "2. Фамилия, Имя, Отчество",
) -> BytesIO:
    document = Document()
    table = document.add_table(rows=30, cols=8)
    table.cell(1, 0).text = fio_label
    table.cell(1, 3).text = FAKE_NAME
    table.cell(8, 4).text = "8.2. Серия и номер документа"
    table.cell(9, 4).text = FAKE_ID
    table.cell(10, 0).text = "8.3. Дата выдачи и кем выдан документ"
    table.cell(11, 0).text = issue_text
    table.cell(29, 0).text = account_text
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream


def _ready_app(settings: Settings, cells_csv_path: Path):
    initialize_databases(settings, cells_csv_path=cells_csv_path)
    app = create_app(settings)
    app.config["TODAY_PROVIDER"] = lambda: date(2026, 7, 20)
    app.config["EMPLOYEE_PROVIDER"] = lambda: "test-user"
    return app


def test_extracts_only_the_five_approved_fields() -> None:
    result = extract_statement_data(
        _statement_docx(), file_name="statement.docx", as_of_date=date(2026, 7, 20)
    )

    assert result.to_dict() == {
        "client_full_name": FAKE_NAME,
        "id_card_number": FAKE_ID,
        "id_card_issuer": FAKE_ISSUER,
        "id_card_issue_date": "2017-09-12",
        "account_number": FAKE_ACCOUNT,
    }


def test_account_is_read_only_after_the_point_17_marker() -> None:
    result = extract_statement_data(
        _statement_docx(
            account_text=(
                "Другой номер 1111222233334444. "
                "17. Банковский сейф (Ячейки) № " + FAKE_ACCOUNT
            )
        ),
        file_name="statement.docx",
        as_of_date=date(2026, 7, 20),
    )

    assert result.account_number == FAKE_ACCOUNT


@pytest.mark.parametrize(
    ("stream", "file_name"),
    [
        (BytesIO(b"not a zip"), "statement.docx"),
        (_statement_docx(), "statement.pdf"),
        (BytesIO(b"x" * (MAX_STATEMENT_BYTES + 1)), "statement.docx"),
    ],
)
def test_rejects_unsafe_or_unsupported_uploads(stream: BytesIO, file_name: str) -> None:
    with pytest.raises(StatementImportValidationError):
        extract_statement_data(
            stream, file_name=file_name, as_of_date=date(2026, 7, 20)
        )


def test_rejects_changed_abs_structure_without_exposing_client_data() -> None:
    with pytest.raises(StatementImportValidationError) as captured:
        extract_statement_data(
            _statement_docx(fio_label="Изменённая форма " + FAKE_NAME),
            file_name="private-client-name.docx",
            as_of_date=date(2026, 7, 20),
        )

    message = str(captured.value)
    assert FAKE_NAME not in message
    assert "private-client-name.docx" not in message


@pytest.mark.parametrize(
    "issue_text",
    [
        f"{FAKE_ISSUER} 31.02.2017 г.",
        f"{FAKE_ISSUER} 21.07.2026 г.",
        f"{FAKE_ISSUER} 12.09.2017 и 13.09.2017",
    ],
)
def test_rejects_invalid_or_ambiguous_issue_date(issue_text: str) -> None:
    with pytest.raises(StatementImportValidationError):
        extract_statement_data(
            _statement_docx(issue_text=issue_text),
            file_name="statement.docx",
            as_of_date=date(2026, 7, 20),
        )


def test_rejects_missing_or_ambiguous_point_17_account() -> None:
    for account_text in (
        "17. Другой банковский продукт № " + FAKE_ACCOUNT,
        (
            "17. Банковский сейф (Ячейки) № "
            + FAKE_ACCOUNT
            + " Банковский сейф (Ячейки) № 8888000011112222"
        ),
    ):
        with pytest.raises(StatementImportValidationError):
            extract_statement_data(
                _statement_docx(account_text=account_text),
                file_name="statement.docx",
                as_of_date=date(2026, 7, 20),
            )


def test_private_route_returns_fields_without_saving_uploaded_file(
    settings: Settings, cells_csv_path: Path
) -> None:
    app = _ready_app(settings, cells_csv_path)
    token = app.extensions["safe_cells_private_token"]

    response = app.test_client().post(
        "/api/contract-statements/extract",
        headers={"X-Safe-Cells-Token": token},
        data={"file": (_statement_docx(), "statement.docx")},
    )

    assert response.status_code == 200
    assert response.get_json()["account_number"] == FAKE_ACCOUNT
    assert not list(settings.database_directory.rglob("*.docx"))


def test_private_route_requires_local_page_token(
    settings: Settings, cells_csv_path: Path
) -> None:
    app = _ready_app(settings, cells_csv_path)

    response = app.test_client().post(
        "/api/contract-statements/extract",
        data={"file": (_statement_docx(), "statement.docx")},
    )

    assert response.status_code == 403
    assert FAKE_NAME not in response.get_data(as_text=True)


def test_route_error_does_not_expose_file_name_or_client_data(
    settings: Settings, cells_csv_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    app = _ready_app(settings, cells_csv_path)
    token = app.extensions["safe_cells_private_token"]

    response = app.test_client().post(
        "/api/contract-statements/extract",
        headers={"X-Safe-Cells-Token": token},
        data={
            "file": (
                _statement_docx(fio_label="Изменённая форма " + FAKE_NAME),
                "private-client-name.docx",
            )
        },
    )

    combined = response.get_data(as_text=True) + caplog.text
    assert response.status_code == 400
    assert FAKE_NAME not in combined
    assert "private-client-name.docx" not in combined
