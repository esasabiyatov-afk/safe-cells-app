from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
import json
from pathlib import Path
import sqlite3
from uuid import uuid4

import pytest
from docx import Document

from app import create_app
from app.config import Settings
from app.db.connections import DatabasePaths, open_readonly, open_write
from app.services.backups import list_backup_sets
from app.services.contracts import (
    BUSY_MESSAGE,
    ContractBusyError,
    ContractConflictError,
    ContractNetworkError,
    ContractValidationError,
    ContractWriteError,
    ContractWriteUncertainError,
    create_contract,
)


OCCURRED_AT = datetime(2026, 7, 12, 9, 30, tzinfo=timezone(timedelta(hours=6)))


def contract_payload(**overrides) -> dict:
    payload = {
        "operation_id": str(uuid4()),
        "cell_number": "1",
        "client_full_name": "Тестовый Клиент",
        "client_phone": "+996 (555) 123-456",
        "id_card_number": "TEST-ID-001",
        "id_card_issuer": "Тестовый орган выдачи",
        "id_card_issue_date": "2017-09-12",
        "account_number": "TEST-ACCOUNT-001",
        "start_date": "2026-07-12",
        "end_date": "2026-08-10",
        "rent_days": 30,
    }
    payload.update(overrides)
    return payload


def _counts(settings: Settings) -> tuple[int, int]:
    paths = DatabasePaths.from_settings(settings)
    with open_readonly(paths.working) as connection:
        contracts = connection.execute("SELECT COUNT(*) FROM contracts").fetchone()[0]
    with open_readonly(paths.archive) as connection:
        logs = connection.execute("SELECT COUNT(*) FROM log").fetchone()[0]
    return int(contracts), int(logs)


def test_document_failure_after_api_save_returns_warning_without_undoing_contract(
    settings: Settings, initialized_databases
) -> None:
    with open_write(settings) as connection:
        connection.execute("BEGIN IMMEDIATE")
        connection.execute(
            "INSERT INTO document_templates VALUES(?, ?, ?, ?, ?, 1, ?, ?)",
            (
                "missing-opening-template",
                "opening",
                "ТЕСТОВЫЙ ДОКУМЕНТ",
                "missing.docx",
                json.dumps(["Сейф.Номер"], ensure_ascii=False),
                OCCURRED_AT.isoformat(),
                "test-user",
            ),
        )
        connection.commit()
    app = create_app(settings)
    app.config.update(
        TODAY_PROVIDER=lambda: date(2026, 7, 12),
        TIMESTAMP_PROVIDER=lambda: OCCURRED_AT,
        DOWNLOADS_DIRECTORY_PROVIDER=lambda: settings.database_directory / "downloads",
        EMPLOYEE_PROVIDER=lambda: "Тестовый Сотрудник",
    )

    response = app.test_client().post("/api/contracts", json=contract_payload())

    assert response.status_code == 201
    body = response.get_json()
    assert body["documents"] == []
    assert "Файл шаблона не найден" in body["document_warning"]
    assert _counts(settings) == (1, 1)


def test_contract_api_prepares_complete_browser_download_bundle_without_silent_files(
    settings: Settings, initialized_databases, tmp_path: Path
) -> None:
    templates = settings.database_directory / "templates"
    templates.mkdir()
    with open_write(settings) as connection:
        connection.execute("BEGIN IMMEDIATE")
        for index, (template_id, display_name) in enumerate(
            (
                ("opening-browser-a", "ТЕСТ-АКТ"),
                ("opening-browser-b", "ТЕСТ-ДОГОВОР"),
                ("opening-browser-c", "ТЕСТ-РАСПОРЯЖЕНИЕ"),
                ("opening-browser-d", "ТЕСТ-БИРКА"),
            ),
            start=1,
        ):
            file_name = f"browser-{index}.docx"
            document = Document()
            document.add_paragraph("ТЕСТОВЫЙ ДОКУМЕНТ: [Сейф.Номер]")
            document.save(templates / file_name)
            connection.execute(
                "INSERT INTO document_templates VALUES(?, 'opening', ?, ?, ?, 1, ?, ?)",
                (
                    template_id,
                    display_name,
                    file_name,
                    json.dumps(["Сейф.Номер"], ensure_ascii=False),
                    OCCURRED_AT.isoformat(),
                    "test-user",
                ),
            )
        connection.commit()

    silent_downloads = tmp_path / "silent-downloads"
    app = create_app(settings)
    app.config.update(
        TODAY_PROVIDER=lambda: date(2026, 7, 12),
        TIMESTAMP_PROVIDER=lambda: OCCURRED_AT,
        DOWNLOADS_DIRECTORY_PROVIDER=lambda: silent_downloads,
        EMPLOYEE_PROVIDER=lambda: "Тестовый Сотрудник",
    )
    token = app.extensions["safe_cells_private_token"]
    client = app.test_client()

    response = client.post("/api/contracts", json=contract_payload(cell_number="4"))

    assert response.status_code == 201
    body = response.get_json()
    assert body["document_warning"] is None
    assert len(body["documents"]) == 4
    assert all(
        set(document_info) == {"download_id", "file_name"}
        for document_info in body["documents"]
    )
    assert all(
        "_Ячейка-04_" in document_info["file_name"]
        for document_info in body["documents"]
    )
    assert not silent_downloads.exists()
    for document_info in body["documents"]:
        downloaded = client.post(
            "/api/documents/download",
            headers={"X-Safe-Cells-Token": token},
            json={"download_id": document_info["download_id"]},
        )
        assert downloaded.status_code == 200
        assert Document(BytesIO(downloaded.data)).paragraphs[0].text.endswith("04")


def test_create_contract_saves_active_row_audit_and_verified_backups(
    settings: Settings, initialized_databases
) -> None:
    payload = contract_payload()
    result = create_contract(
        settings,
        payload=payload,
        employee="test-user",
        occurred_at=OCCURRED_AT,
    )

    assert result.cell_number == "1"
    assert result.rent_days == 30
    assert result.price_per_day == 15
    assert result.rent_price == 450
    assert result.deposit_amount == 1500
    assert result.repeated is False
    assert result.backup_created is True
    assert result.warning is None
    assert _counts(settings) == (1, 1)

    paths = DatabasePaths.from_settings(settings)
    with open_readonly(paths.working) as connection:
        row = connection.execute("SELECT * FROM contracts").fetchone()
        assert row["cell_number"] == "1"
        assert row["client_full_name"] == "Тестовый Клиент"
        assert row["client_phone"] == "+996 (555) 123-456"
        assert row["rent_price_minor"] == 450
        assert row["deposit_amount_minor"] == 1500
        assert row["created_by"] == "test-user"
    with open_readonly(paths.archive) as connection:
        audit = connection.execute("SELECT * FROM log").fetchone()
        changes = json.loads(audit["changes_json"])
        assert audit["operation_id"] == payload["operation_id"]
        assert audit["action"] == "contract.created"
        assert audit["employee"] == "test-user"
        assert changes["rent_price"] == 450
        assert "client_full_name" not in changes
        assert "id_card_number" not in changes
        assert "account_number" not in changes

    backup_sets = list_backup_sets(settings)
    assert len(backup_sets) == 1
    backup_set = backup_sets[0]
    working_backups = [backup_set.directory / backup_set.working.name]
    archive_backups = [backup_set.directory / backup_set.archive.name]
    for path in working_backups + archive_backups:
        connection = sqlite3.connect(path)
        try:
            assert connection.execute("PRAGMA quick_check").fetchone()[0] == "ok"
        finally:
            connection.close()
    working_backup = sqlite3.connect(working_backups[0])
    archive_backup = sqlite3.connect(archive_backups[0])
    try:
        assert working_backup.execute("SELECT COUNT(*) FROM contracts").fetchone()[0] == 1
        backup_audit = archive_backup.execute(
            "SELECT operation_id FROM log"
        ).fetchone()
        assert backup_audit[0] == payload["operation_id"]
    finally:
        working_backup.close()
        archive_backup.close()


def test_repeated_operation_returns_existing_contract_without_duplicate(
    settings: Settings, initialized_databases
) -> None:
    payload = contract_payload()
    first = create_contract(
        settings, payload=payload, employee="test-user", occurred_at=OCCURRED_AT
    )
    second = create_contract(
        settings, payload=payload, employee="test-user", occurred_at=OCCURRED_AT
    )

    assert second.contract_id == first.contract_id
    assert second.repeated is True
    assert _counts(settings) == (1, 1)
    assert len(list_backup_sets(settings)) == 1

    with pytest.raises(ContractConflictError, match="уже использован"):
        create_contract(
            settings,
            payload=contract_payload(
                operation_id=payload["operation_id"], cell_number="2"
            ),
            employee="test-user",
            occurred_at=OCCURRED_AT,
        )


def test_second_operation_cannot_occupy_same_cell(
    settings: Settings, initialized_databases
) -> None:
    create_contract(
        settings,
        payload=contract_payload(),
        employee="first-user",
        occurred_at=OCCURRED_AT,
    )
    with pytest.raises(ContractConflictError, match="уже занята"):
        create_contract(
            settings,
            payload=contract_payload(operation_id=str(uuid4())),
            employee="second-user",
            occurred_at=OCCURRED_AT,
        )
    assert _counts(settings) == (1, 1)


def test_concurrent_requests_leave_one_active_contract(
    settings: Settings, initialized_databases
) -> None:
    payloads = [
        contract_payload(operation_id=str(uuid4()), account_number=f"TEST-{index}")
        for index in range(2)
    ]

    def run(payload: dict):
        try:
            return create_contract(
                settings,
                payload=payload,
                employee="parallel-user",
                occurred_at=OCCURRED_AT,
            )
        except ContractConflictError as exc:
            return exc

    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(run, payloads))

    assert sum(not isinstance(result, Exception) for result in results) == 1
    assert sum(isinstance(result, ContractConflictError) for result in results) == 1
    assert _counts(settings) == (1, 1)


@pytest.mark.parametrize(
    ("field", "value", "message"),
    [
        ("client_full_name", "", "ФИО клиента"),
        ("client_phone", "", "Номер телефона"),
        ("id_card_number", "", "ID-карты"),
        ("id_card_issuer", "", "Орган выдачи"),
        ("id_card_issue_date", "", "дату выдачи ID-карты"),
        ("account_number", "", "Номер счёта"),
        ("operation_id", "not-a-uuid", "подготовить операцию"),
        ("rent_days", 0, "не меньше 1"),
    ],
)
def test_required_and_invalid_fields_are_rejected_before_write(
    settings: Settings,
    initialized_databases,
    field: str,
    value: object,
    message: str,
) -> None:
    with pytest.raises(ContractValidationError, match=message):
        create_contract(
            settings,
            payload=contract_payload(**{field: value}),
            employee="test-user",
            occurred_at=OCCURRED_AT,
        )
    assert _counts(settings) == (0, 0)


def test_future_id_card_issue_date_is_rejected_before_write(
    settings: Settings, initialized_databases
) -> None:
    with pytest.raises(ContractValidationError, match="не может быть в будущем"):
        create_contract(
            settings,
            payload=contract_payload(id_card_issue_date="2026-07-13"),
            employee="test-user",
            occurred_at=OCCURRED_AT,
        )
    assert _counts(settings) == (0, 0)


def test_backend_rejects_tampered_days_and_does_not_write(
    settings: Settings, initialized_databases
) -> None:
    with pytest.raises(ContractValidationError, match="не соответствует"):
        create_contract(
            settings,
            payload=contract_payload(rent_days=1),
            employee="test-user",
            occurred_at=OCCURRED_AT,
        )
    assert _counts(settings) == (0, 0)


def test_backend_rejects_contract_starting_after_today(
    settings: Settings, initialized_databases
) -> None:
    with pytest.raises(ContractValidationError, match="сегодняшней"):
        create_contract(
            settings,
            payload=contract_payload(
                start_date="2026-07-13", end_date="2026-07-13", rent_days=1
            ),
            employee="test-user",
            occurred_at=OCCURRED_AT,
            as_of_date=date(2026, 7, 12),
        )
    assert _counts(settings) == (0, 0)


def test_backend_recalculates_current_tariff_inside_write_transaction(
    settings: Settings, initialized_databases
) -> None:
    with open_write(settings) as connection:
        connection.execute("BEGIN IMMEDIATE")
        connection.execute(
            """
            UPDATE tariffs SET price_per_day_minor = 19
            WHERE height_mm = 50 AND period_from_days = 1
            """
        )
        connection.commit()

    result = create_contract(
        settings,
        payload=contract_payload(),
        employee="test-user",
        occurred_at=OCCURRED_AT,
    )
    assert result.price_per_day == 19
    assert result.rent_price == 570


def test_uncertain_commit_is_not_retried_automatically_and_repeat_is_safe(
    settings: Settings,
    initialized_databases,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    real_open_write = open_write

    class CommitThenDisconnect:
        def __init__(self, connection: sqlite3.Connection) -> None:
            self.connection = connection

        def __getattr__(self, name: str):
            return getattr(self.connection, name)

        def commit(self) -> None:
            self.connection.commit()
            raise sqlite3.OperationalError("simulated disconnect after commit")

    @contextmanager
    def uncertain_open_write(settings_arg: Settings, *, attach_archive: bool = False):
        with real_open_write(settings_arg, attach_archive=attach_archive) as connection:
            yield CommitThenDisconnect(connection)

    payload = contract_payload()
    monkeypatch.setattr("app.services.contracts.open_write", uncertain_open_write)
    with pytest.raises(ContractWriteUncertainError, match="Не повторяйте"):
        create_contract(
            settings,
            payload=payload,
            employee="test-user",
            occurred_at=OCCURRED_AT,
        )
    assert _counts(settings) == (1, 1)

    monkeypatch.setattr("app.services.contracts.open_write", real_open_write)
    repeated = create_contract(
        settings,
        payload=payload,
        employee="test-user",
        occurred_at=OCCURRED_AT,
    )
    assert repeated.repeated is True
    assert _counts(settings) == (1, 1)


def test_partial_failure_rolls_back_contract_and_audit(
    settings: Settings, initialized_databases
) -> None:
    def fail_after_insert() -> None:
        raise RuntimeError("simulated partial failure")

    with pytest.raises(ContractWriteError, match="Изменения отменены"):
        create_contract(
            settings,
            payload=contract_payload(),
            employee="test-user",
            occurred_at=OCCURRED_AT,
            after_contract_insert=fail_after_insert,
        )
    assert _counts(settings) == (0, 0)


def test_locked_database_returns_clear_error_without_partial_write(
    settings: Settings, initialized_databases
) -> None:
    short_settings = Settings(
        database_directory=settings.database_directory,
        busy_timeout_ms=40,
        testing=True,
    )
    with open_write(settings, attach_archive=True) as blocker:
        blocker.execute("BEGIN IMMEDIATE")
        with pytest.raises(ContractBusyError, match="Другая операция") as error:
            create_contract(
                short_settings,
                payload=contract_payload(),
                employee="test-user",
                occurred_at=OCCURRED_AT,
            )
        assert str(error.value) == BUSY_MESSAGE
        blocker.rollback()
    assert _counts(settings) == (0, 0)


def test_missing_network_path_does_not_create_local_database(tmp_path: Path) -> None:
    missing = tmp_path / "offline network"
    settings = Settings(database_directory=missing, testing=True)
    with pytest.raises(ContractNetworkError, match="сетевого диска"):
        create_contract(
            settings,
            payload=contract_payload(),
            employee="test-user",
            occurred_at=OCCURRED_AT,
        )
    assert not missing.exists()


def test_personal_markers_are_not_written_to_technical_logs(
    settings: Settings, initialized_databases, caplog: pytest.LogCaptureFixture
) -> None:
    payload = contract_payload(
        client_full_name="PRIVATE-NAME-MARKER",
        id_card_number="PRIVATE-ID-MARKER",
        account_number="PRIVATE-ACCOUNT-MARKER",
    )
    create_contract(
        settings,
        payload=payload,
        employee="test-user",
        occurred_at=OCCURRED_AT,
    )
    assert "PRIVATE-NAME-MARKER" not in caplog.text
    assert "PRIVATE-ID-MARKER" not in caplog.text
    assert "PRIVATE-ACCOUNT-MARKER" not in caplog.text


def test_backup_failure_is_reported_after_saved_contract(
    settings: Settings, initialized_databases, monkeypatch: pytest.MonkeyPatch
) -> None:
    def fail_backup(*args, **kwargs):
        raise OSError("simulated backup failure")

    monkeypatch.setattr("app.services.contracts.create_backup_pair", fail_backup)
    payload = contract_payload()
    result = create_contract(
        settings,
        payload=payload,
        employee="test-user",
        occurred_at=OCCURRED_AT,
    )
    assert result.backup_created is False
    assert "администратору" in result.warning
    assert _counts(settings) == (1, 1)

    repeated = create_contract(
        settings,
        payload=payload,
        employee="test-user",
        occurred_at=OCCURRED_AT,
    )
    assert repeated.repeated is True
    assert repeated.backup_created is False
    assert "не найден" in repeated.warning


def test_api_creates_contract_without_returning_personal_data(
    settings: Settings, initialized_databases
) -> None:
    app = create_app(settings)
    app.config["EMPLOYEE_PROVIDER"] = lambda: "api-user"
    app.config["TIMESTAMP_PROVIDER"] = lambda: OCCURRED_AT
    response = app.test_client().post("/api/contracts", json=contract_payload())
    body = response.get_data(as_text=True)
    payload = response.get_json()

    assert response.status_code == 201
    assert payload["cell_number"] == "1"
    assert payload["rent_price"] == 450
    assert payload["repeated"] is False
    assert "Тестовый Клиент" not in body
    assert "TEST-ID-001" not in body
    assert "TEST-ACCOUNT-001" not in body
    assert response.headers["Cache-Control"] == "no-store"


def test_api_repeat_is_idempotent_and_conflict_is_clear(
    settings: Settings, initialized_databases
) -> None:
    app = create_app(settings)
    app.config["EMPLOYEE_PROVIDER"] = lambda: "api-user"
    app.config["TIMESTAMP_PROVIDER"] = lambda: OCCURRED_AT
    client = app.test_client()
    payload = contract_payload()

    assert client.post("/api/contracts", json=payload).status_code == 201
    assert client.post("/api/contracts", json=payload).status_code == 200
    conflict = client.post(
        "/api/contracts", json=contract_payload(operation_id=str(uuid4()))
    )
    assert conflict.status_code == 409
    assert "занята" in conflict.get_json()["message"]
    assert _counts(settings) == (1, 1)
