from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
import json
from pathlib import Path
from typing import Callable
from uuid import uuid4

from openpyxl import load_workbook
import pytest

from app import create_app
from app.config import Settings
from app.db.connections import NETWORK_ERROR_MESSAGE, open_write
from app.services import journal
from app.services.closures import close_contract
from app.services.journal import (
    JournalReadError,
    JournalValidationError,
    list_journal_entries,
    list_journal_report_entries,
)
from app.services.journal_reports import build_journal_report


def _insert_log(
    settings: Settings,
    *,
    occurred_at: str,
    action: str,
    cell_number: str | None,
    changes: dict[str, object],
    contract_id: str | None = None,
    employee: str = "Тестовый сотрудник",
) -> None:
    operation_id = str(uuid4())
    with open_write(settings, attach_archive=True) as connection:
        connection.execute("BEGIN IMMEDIATE")
        connection.execute(
            """
            INSERT INTO archive.log(
                log_id, operation_id, occurred_at, employee, action,
                contract_id, cell_number, changes_json
            ) VALUES(?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                str(uuid4()),
                operation_id,
                occurred_at,
                employee,
                action,
                contract_id,
                cell_number,
                json.dumps(changes, ensure_ascii=False),
            ),
        )
        connection.commit()


def _insert_sample_entries(
    settings: Settings,
    insert_test_contract: Callable[..., None],
) -> None:
    insert_test_contract(
        cell_number="7",
        start_date="2026-07-01",
        end_date="2026-07-30",
        client_name="Тестовый Клиент Семь",
    )
    insert_test_contract(
        cell_number="9",
        end_date="2026-07-30",
        client_name="Тестовый Клиент Девять",
    )
    _insert_log(
        settings,
        occurred_at="2026-07-01T09:00:00+06:00",
        action="contract.created",
        cell_number="7",
        contract_id="contract-test-7",
        changes={
            "start_date": "2026-07-01",
            "end_date": "2026-07-30",
            "rent_days": 30,
            "id_card_number": "SECRET-ID-CREATED",
        },
    )
    _insert_log(
        settings,
        occurred_at="2026-07-10T10:30:00+06:00",
        action="contract.renewed",
        cell_number="7",
        contract_id="contract-test-7",
        changes={
            "old_end_date": "2026-07-30",
            "new_start_date": "2026-07-31",
            "new_end_date": "2026-08-30",
            "renewal_days": 31,
            "penalty_days": 2,
            "account_number": "SECRET-ACCOUNT",
        },
        employee="Второй тестовый сотрудник",
    )
    _insert_log(
        settings,
        occurred_at="2026-07-11T11:00:00+06:00",
        action="contract.edited",
        cell_number="9",
        contract_id="contract-test-9",
        changes={
            "client_full_name": {
                "old": "СЕКРЕТНОЕ СТАРОЕ ФИО",
                "new": "СЕКРЕТНОЕ НОВОЕ ФИО",
            },
            "account_number": {"old": "SECRET-OLD", "new": "SECRET-NEW"},
        },
    )
    _insert_log(
        settings,
        occurred_at="2026-07-12T12:00:00+06:00",
        action="contract.closed",
        cell_number="7",
        contract_id="contract-test-7",
        changes={
            "close_date": "2026-07-12",
            "close_reason": "Досрочное расторжение",
            "penalty_days": 0,
            "id_card_number": "SECRET-ID-CLOSED",
        },
    )
    _insert_log(
        settings,
        occurred_at="2026-07-13T12:00:00+06:00",
        action="admin.settings.updated",
        cell_number=None,
        changes={"deposit_amount": {"old": 1500, "new": 1600}},
    )


def test_journal_includes_client_but_excludes_edits_and_secret_fields(
    settings: Settings, initialized_databases, insert_test_contract
) -> None:
    _insert_sample_entries(settings, insert_test_contract)

    payload = list_journal_entries(settings)
    serialized = json.dumps(payload, ensure_ascii=False)

    assert [entry["action"] for entry in payload["entries"]] == [
        "contract.closed",
        "contract.renewed",
        "contract.created",
    ]
    assert payload["pagination"]["total"] == 3
    assert all(
        entry["client_full_name"] == "Тестовый Клиент Семь"
        for entry in payload["entries"]
    )
    assert payload["entries"][-1]["action_label"] == "Открытие"
    assert payload["entries"][1]["action_label"] == "Продление"
    assert payload["entries"][1]["is_overdue"] is True
    assert payload["entries"][1]["report_action_label"] == "Продление / Просрочка"
    assert "contract.edited" not in serialized
    assert "admin.settings.updated" not in serialized
    assert "СЕКРЕТНОЕ" not in serialized
    assert "SECRET-" not in serialized
    assert "Период продления: 31.07.2026 — 30.08.2026; Срок: 31 дн." in serialized
    assert "Период аренды: 01.07.2026 — 30.07.2026; Срок: 30 дн." in serialized
    assert "Дата закрытия:" not in serialized
    assert "Причина:" not in serialized


def test_journal_resolves_client_from_archive_after_real_closure(
    settings: Settings, initialized_databases, insert_test_contract
) -> None:
    insert_test_contract(
        cell_number="12",
        start_date="2026-07-01",
        end_date="2026-07-12",
        client_name="Архивный Тестовый Клиент",
    )
    close_contract(
        settings,
        payload={
            "operation_id": str(uuid4()),
            "cell_number": "12",
            "contract_ref": "contract-test-12",
            "expected_end_date": "2026-07-12",
            "reason_code": "standard",
        },
        employee="Тестовый сотрудник",
        close_date=date(2026, 7, 12),
        occurred_at=datetime.fromisoformat("2026-07-12T12:00:00+06:00"),
    )

    payload = list_journal_entries(
        settings, cell_number="12", action="contract.closed"
    )

    assert payload["pagination"]["total"] == 1
    assert payload["entries"][0]["client_full_name"] == "Архивный Тестовый Клиент"


def test_journal_filters_and_paginates_with_parameterized_values(
    settings: Settings, initialized_databases, insert_test_contract
) -> None:
    _insert_sample_entries(settings, insert_test_contract)

    first = list_journal_entries(settings, cell_number="7", page=1, page_size=2)
    second = list_journal_entries(settings, cell_number="7", page=2, page_size=2)
    filtered = list_journal_entries(
        settings,
        action="contract.renewed",
        date_from="2026-07-10",
        date_to="2026-07-10",
    )
    overdue = list_journal_entries(settings, action="overdue")
    by_client = list_journal_entries(settings, client_name="кЛиЕнТ сЕмЬ")
    by_employee = list_journal_entries(
        settings, employee="Второй тестовый сотрудник"
    )
    injection = list_journal_entries(settings, cell_number="7' OR 1=1 --")
    client_injection = list_journal_entries(
        settings, client_name="x%' OR 1=1 --"
    )

    assert first["pagination"] == {
        "page": 1,
        "page_size": 2,
        "page_count": 2,
        "total": 3,
        "has_previous": False,
        "has_next": True,
    }
    assert len(first["entries"]) == 2
    assert len(second["entries"]) == 1
    assert [entry["action"] for entry in filtered["entries"]] == [
        "contract.renewed"
    ]
    assert [entry["action"] for entry in overdue["entries"]] == [
        "contract.renewed"
    ]
    assert len(by_client["entries"]) == 3
    assert [entry["action"] for entry in by_employee["entries"]] == [
        "contract.renewed"
    ]
    assert "Второй тестовый сотрудник" in first["filters"]["employees"]
    assert injection["entries"] == []
    assert client_injection["entries"] == []


def test_overdue_badge_is_only_added_to_renewal_or_closure(
    settings: Settings, initialized_databases
) -> None:
    _insert_log(
        settings,
        occurred_at="2026-07-14T09:00:00+06:00",
        action="cell.manual_occupied",
        cell_number="3",
        changes={"block_kind": "manual", "penalty_days": 99},
    )
    _insert_log(
        settings,
        occurred_at="2026-07-14T10:00:00+06:00",
        action="contract.closed",
        cell_number="4",
        changes={
            "close_reason": "Закрытие после окончания срока",
            "penalty_days": 3,
        },
    )

    payload = list_journal_entries(settings)
    manual = next(
        entry for entry in payload["entries"]
        if entry["action"] == "cell.manual_occupied"
    )
    closed = next(entry for entry in payload["entries"] if entry["action"] == "contract.closed")
    overdue = list_journal_entries(settings, action="overdue")

    assert manual["is_overdue"] is False
    assert manual["report_action_label"] == "Занятие без договора"
    assert closed["is_overdue"] is True
    assert closed["report_action_label"] == "Закрытие / Просрочка"
    assert [entry["action"] for entry in overdue["entries"]] == [
        "contract.closed"
    ]


@pytest.mark.parametrize(
    ("arguments", "message"),
    [
        ({"action": "contract.edited"}, "допустимое действие"),
        ({"action": "admin.settings.updated"}, "допустимое действие"),
        ({"date_from": "2026-07-02", "date_to": "2026-07-01"}, "раньше"),
        ({"date_from": "01.07.2026"}, "корректную дату"),
        ({"page": 0}, "от 1"),
        ({"page_size": 101}, "от 1 до 100"),
    ],
)
def test_journal_rejects_invalid_filters(
    settings: Settings, initialized_databases, arguments: dict[str, object], message: str
) -> None:
    with pytest.raises(JournalValidationError, match=message):
        list_journal_entries(settings, **arguments)


def test_journal_unavailable_path_does_not_create_local_database(tmp_path: Path) -> None:
    missing = tmp_path / "offline journal data"
    settings = Settings(database_directory=missing, testing=True)

    with pytest.raises(JournalReadError, match=NETWORK_ERROR_MESSAGE):
        list_journal_entries(settings)

    assert not missing.exists()


def test_journal_api_requires_employee_and_returns_client_safely(
    settings: Settings, initialized_databases, insert_test_contract
) -> None:
    _insert_sample_entries(settings, insert_test_contract)
    app = create_app(settings)
    client = app.test_client()

    missing_employee = client.get("/api/journal")
    assert missing_employee.status_code == 409
    assert missing_employee.get_json()["selection_required"] is True

    app.config["EMPLOYEE_PROVIDER"] = lambda: "Тестовый сотрудник"
    response = client.get(
        "/api/journal?cell_number=7&action=contract.closed&date_from=2026-07-01&page=1"
    )
    serialized = response.get_data(as_text=True)

    assert response.status_code == 200
    assert response.get_json()["pagination"]["total"] == 1
    assert response.get_json()["entries"][0]["action_label"] == "Закрытие"
    assert response.get_json()["entries"][0]["client_full_name"] == "Тестовый Клиент Семь"
    assert "SECRET-ID" not in serialized
    assert client.get("/api/journal?action=unknown").status_code == 400


def test_journal_report_downloads_real_filtered_xlsx_without_edits_or_secrets(
    settings: Settings, initialized_databases, insert_test_contract
) -> None:
    _insert_sample_entries(settings, insert_test_contract)
    app = create_app(settings)
    app.config["EMPLOYEE_PROVIDER"] = lambda: "Тестовый сотрудник"
    app.config["TODAY_PROVIDER"] = lambda: date(2026, 7, 16)
    client = app.test_client()

    response = client.get(
        "/api/journal/report?cell_number=7&date_from=2026-07-01&date_to=2026-07-12"
    )

    assert response.status_code == 200
    assert response.mimetype == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in response.headers["Content-Disposition"]
    workbook = load_workbook(BytesIO(response.data), read_only=True, data_only=False)
    worksheet = workbook["Журнал ячеек"]
    assert worksheet["A1"].value == "Выписка из журнала сейфовых ячеек"
    assert worksheet["A5"].value == "Дата и время"
    assert worksheet["D5"].value == "Клиент"
    rows = list(worksheet.iter_rows(min_row=6, values_only=True))
    serialized = json.dumps(rows, ensure_ascii=False)
    assert len(rows) == 3
    assert serialized.count("Тестовый Клиент Семь") == 3
    assert "Открытие" in serialized
    assert "Продление / Просрочка" in serialized
    assert "Закрытие" in serialized
    assert "Период аренды: 01.07.2026 — 30.07.2026; Срок: 30 дн." in serialized
    assert "Причина:" not in serialized
    assert "contract.edited" not in serialized
    assert "СЕКРЕТНОЕ" not in serialized
    assert "SECRET-" not in serialized


def test_journal_report_requires_employee_and_applies_row_limit(
    settings: Settings, initialized_databases, insert_test_contract, monkeypatch
) -> None:
    _insert_sample_entries(settings, insert_test_contract)
    app = create_app(settings)
    client = app.test_client()

    assert client.get("/api/journal/report").status_code == 409
    monkeypatch.setattr(journal, "MAX_REPORT_ROWS", 2)
    with pytest.raises(JournalValidationError, match="слишком много"):
        list_journal_report_entries(settings)


def test_journal_report_applies_client_and_employee_filters(
    settings: Settings, initialized_databases, insert_test_contract
) -> None:
    _insert_sample_entries(settings, insert_test_contract)
    app = create_app(settings)
    app.config["EMPLOYEE_PROVIDER"] = lambda: "Тестовый сотрудник"
    app.config["TODAY_PROVIDER"] = lambda: date(2026, 7, 16)

    response = app.test_client().get(
        "/api/journal/report",
        query_string={
            "client_name": "клиент семь",
            "employee": "Второй тестовый сотрудник",
        },
    )

    assert response.status_code == 200
    worksheet = load_workbook(BytesIO(response.data), read_only=True).active
    rows = list(worksheet.iter_rows(min_row=6, values_only=True))
    assert len(rows) == 1
    assert rows[0][2] == "Продление / Просрочка"
    assert rows[0][3] == "Тестовый Клиент Семь"
    assert rows[0][4] == "Второй тестовый сотрудник"
    assert "клиент: клиент семь" in worksheet["A3"].value
    assert "сотрудник: Второй тестовый сотрудник" in worksheet["A3"].value


def test_journal_report_neutralizes_excel_formula_values() -> None:
    report = build_journal_report(
        [
            {
                "occurred_at": "2026-07-01T09:00:00+06:00",
                "cell_number": "=1+1",
                "action_label": "Открытие",
                "client_full_name": "=HYPERLINK(\"unsafe\")",
                "employee": "+Тест",
                "summary": "@Тест",
            }
        ],
        generated_on=date(2026, 7, 16),
        filters={},
    )

    worksheet = load_workbook(BytesIO(report), data_only=False).active
    assert worksheet["B6"].value == "'=1+1"
    assert worksheet["D6"].value.startswith("'=")
    assert worksheet["E6"].value == "'+Тест"
    assert worksheet["F6"].value == "'@Тест"


def test_journal_interface_uses_text_content_and_local_assets(
    settings: Settings, initialized_databases
) -> None:
    app = create_app(settings)
    client = app.test_client()

    main_html = client.get("/").get_data(as_text=True)
    page_response = client.get("/journal")
    html = page_response.get_data(as_text=True)
    script = client.get("/static/js/journal.js").get_data(as_text=True)
    styles = client.get("/static/css/main.css").get_data(as_text=True)

    assert 'id="journalOpen" href="/journal" target="_blank" rel="noopener"' in main_html
    assert 'src="/static/js/journal.js"' not in main_html
    assert 'data-journal-url="/api/journal"' in html
    assert 'data-journal-report-url="/api/journal/report"' in html
    assert 'src="/static/js/journal.js"' in html
    assert 'id="journalPage"' in html
    assert 'id="journalReport"' in html
    assert 'id="journalSelectionRequired"' in html
    assert 'href="/"' in html
    assert "Общий журнал ячеек" in html
    assert "Скачать отчёт Excel" in html
    assert "ID-карту и номер счёта" in html
    assert '<option value="contract.edited">' not in html
    assert '<option value="overdue">Просрочка</option>' in html
    assert 'id="journalClient"' in html
    assert 'id="journalEmployee"' in html
    assert '<option value="cell.key_restored">Ключ восстановлен</option>' in html
    assert html.count("data-journal-date") == 2
    assert page_response.headers["Cache-Control"] == "no-store"
    assert "innerHTML" not in script
    assert "textContent" in script
    assert "URLSearchParams" in script
    assert "URL.createObjectURL" in script
    assert "journal-action-group" in script
    assert "journal-action-overdue" in script
    assert "client_name: elements.client.value.trim()" in script
    assert "employee: elements.employee.value" in script
    assert "event.preventDefault()" in script
    assert "showModal" not in script
    assert "loadJournal();" in script
    assert ".journal-dialog.journal-page-panel" in styles
    assert "width: min(1460px, calc(100% - 48px));" in styles
