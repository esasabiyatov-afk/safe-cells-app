from __future__ import annotations

from datetime import datetime, timedelta, timezone
from io import BytesIO
import json
from pathlib import Path
from uuid import uuid4

from docx import Document
import pytest

from app import create_app
from app.config import Settings
from app.db.connections import open_readonly, open_write
from app.services.admin_auth import (
    AdminAccessManager,
    AdminAuthenticationError,
    hash_password,
    verify_password,
)
from app.services.admin_settings import AdminValidationError, update_admin_settings
from app.services.admin_templates import save_document_template
from app.services.backups import list_backup_sets


PASSWORD = "TestAdmin-2026"
NEW_PASSWORD = "NewTestAdmin-2027"
OCCURRED_AT = datetime(2026, 7, 14, 10, 0, tzinfo=timezone(timedelta(hours=6)))


def _ready_app(settings):
    app = create_app(settings)
    app.config.update(
        EMPLOYEE_PROVIDER=lambda: "test-admin",
        TIMESTAMP_PROVIDER=lambda: OCCURRED_AT,
    )
    return app


def _setup(client) -> str:
    response = client.post(
        "/api/admin/setup",
        json={
            "operation_id": str(uuid4()),
            "password": PASSWORD,
            "password_confirmation": PASSWORD,
        },
    )
    assert response.status_code == 201, response.get_json()
    return response.get_json()["token"]


def _headers(token: str) -> dict[str, str]:
    return {"X-Safe-Cells-Admin-Token": token}


def _cell(number: str, height_mm: int) -> dict[str, int | str]:
    return {
        "number": number,
        "height_mm": height_mm,
        "width_mm": 220,
        "depth_mm": 330,
    }


def _snapshot(client, token: str) -> dict:
    response = client.get("/api/admin/settings", headers=_headers(token))
    assert response.status_code == 200
    return response.get_json()


def _docx_bytes(placeholder: str = "Сейф.Номер") -> BytesIO:
    buffer = BytesIO()
    document = Document()
    document.add_paragraph(f"ТЕСТОВЫЙ ШАБЛОН [{placeholder}]")
    document.save(buffer)
    buffer.seek(0)
    return buffer


def _static_docx_bytes() -> BytesIO:
    buffer = BytesIO()
    document = Document()
    document.add_paragraph("ТЕСТОВЫЙ ДОКУМЕНТ БЕЗ ПОЛЕЙ")
    document.save(buffer)
    buffer.seek(0)
    return buffer


def test_password_hash_uses_random_salt_and_never_contains_plaintext():
    first = hash_password(PASSWORD)
    second = hash_password(PASSWORD)
    assert first != second
    assert PASSWORD not in first
    assert verify_password(PASSWORD, first)
    assert not verify_password("WrongPassword-1", first)
    assert verify_password("7", hash_password("7"))


def test_initial_settings_use_acknowledgement_without_forced_password_setup(
    settings, initialized_databases
):
    # A database created before this option existed must stay usable.
    with open_write(settings) as working:
        working.execute(
            "DELETE FROM config WHERE key IN ('admin_access_mode', 'penalty_rate_mode', 'penalty_manual_rates_json')"
        )
        working.commit()
    app = _ready_app(settings)
    client = app.test_client()
    assert client.get("/api/admin/status").get_json() == {
        "configured": False,
        "access_mode": "acknowledgement",
        "recovery_mode": False,
    }
    assert client.get("/api/admin/settings").status_code == 401
    acknowledged = client.post(
        "/api/admin/acknowledge",
        json={"accepted": True},
    )
    assert acknowledged.status_code == 200
    assert _snapshot(client, acknowledged.get_json()["token"])[
        "password_configured"
    ] is False

    token = _setup(client)
    assert client.get("/api/admin/status").get_json() == {
        "configured": True,
        "access_mode": "password",
        "recovery_mode": False,
    }
    initial_snapshot = _snapshot(client, token)
    assert initial_snapshot["config"]["deposit_amount_minor"] == 1500
    assert initial_snapshot["penalty"]["mode"] == "linked"
    assert len(initial_snapshot["penalty"]["manual_rates"]) == 6

    with open_readonly(settings.database_directory / "vault_cells.sqlite3") as working:
        stored = working.execute(
            "SELECT password_hash FROM admin_credentials WHERE id = 1"
        ).fetchone()["password_hash"]
    with open_readonly(settings.database_directory / "vault_archive.sqlite3") as archive:
        log = archive.execute(
            "SELECT changes_json FROM log WHERE action = 'admin.password.created'"
        ).fetchone()["changes_json"]
    assert stored != PASSWORD and verify_password(PASSWORD, stored)
    assert PASSWORD not in log
    assert "password_hash" not in log


def test_simple_password_can_be_created_from_settings_without_current_password(
    settings, initialized_databases
):
    app = _ready_app(settings)
    client = app.test_client()
    acknowledged = client.post(
        "/api/admin/acknowledge",
        json={"accepted": True},
    )
    token = acknowledged.get_json()["token"]

    created = client.put(
        "/api/admin/password",
        headers=_headers(token),
        json={
            "operation_id": str(uuid4()),
            "current_password": "",
            "new_password": "7",
            "password_confirmation": "7",
        },
    )

    assert created.status_code == 200
    assert client.get("/api/admin/status").get_json() == {
        "configured": True,
        "access_mode": "password",
        "recovery_mode": False,
    }
    assert client.post("/api/admin/login", json={"password": "7"}).status_code == 200


def test_shared_password_has_no_attempt_lock(settings, initialized_databases):
    app = _ready_app(settings)
    client = app.test_client()
    _setup(client)
    for _ in range(7):
        response = client.post("/api/admin/login", json={"password": "WrongPassword-1"})
        assert response.status_code == 401
    assert client.post("/api/admin/login", json={"password": PASSWORD}).status_code == 200


def test_access_manager_session_lasts_until_revoked():
    manager = AdminAccessManager()
    encoded = hash_password(PASSWORD)
    for _ in range(7):
        with pytest.raises(AdminAuthenticationError, match="Неверный"):
            manager.authenticate("WrongPassword-1", encoded)
    session = manager.authenticate(PASSWORD, encoded)
    manager.require(session.token)
    manager.revoke(session.token)
    with pytest.raises(AdminAuthenticationError, match="завершён"):
        manager.require(session.token)


def test_update_tariffs_and_config_is_atomic_audited_and_backed_up(
    settings, initialized_databases
):
    app = _ready_app(settings)
    client = app.test_client()
    token = _setup(client)
    snapshot = _snapshot(client, token)
    snapshot["config"]["expiring_soon_days"] = 10
    snapshot["config"]["deposit_amount_minor"] = 2000
    snapshot["tariffs"][0]["price_per_day_minor"] = 16
    snapshot["penalty"]["mode"] = "manual"
    snapshot["penalty"]["manual_rates"][0]["price_per_day_minor"] = 23
    operation_id = str(uuid4())
    response = client.put(
        "/api/admin/settings",
        headers=_headers(token),
        json={
            "operation_id": operation_id,
            "config": snapshot["config"],
            "tariffs": snapshot["tariffs"],
            "penalty": snapshot["penalty"],
        },
    )
    assert response.status_code == 200, response.get_json()
    assert response.get_json()["backup_created"] is True
    saved = _snapshot(client, token)
    assert saved["config"] == {
        "deposit_amount_minor": 2000,
        "abs_session_minutes": 60,
        "expiring_soon_days": 10,
        "display_date_words": True,
        "show_ui_hints": True,
    }
    assert saved["tariffs"][0]["price_per_day_minor"] == 16
    assert len(saved["tariffs"]) == 24
    assert saved["penalty"]["mode"] == "manual"
    assert saved["penalty"]["manual_rates"][0] == {
        "height_mm": 50,
        "width_mm": 220,
        "depth_mm": 330,
        "price_per_day_minor": 23,
    }
    with open_readonly(settings.database_directory / "vault_archive.sqlite3") as archive:
        row = archive.execute(
            "SELECT employee, changes_json FROM log WHERE operation_id = ?",
            (operation_id,),
        ).fetchone()
    assert row["employee"] == "test-admin"
    assert "password" not in row["changes_json"].lower()
    assert len(list_backup_sets(settings)) >= 2


def test_admin_can_change_whatsapp_templates_with_required_placeholders(
    settings, initialized_databases
):
    app = _ready_app(settings)
    client = app.test_client()
    token = _setup(client)
    before = _snapshot(client, token)
    assert "[Клиент.Обращение]" in before["reminder_templates"]["expiring"]
    assert "[Договор.Конец]" in before["reminder_templates"]["overdue"]
    templates = {
        "expiring": (
            "Добрый день, [Обращение]!\n"
            "Ячейка №[Номер ячейки] действует до [Дата окончания]."
        ),
        "overdue": (
            "Здравствуйте, [Обращение]!\n"
            "Ячейка №[Номер ячейки] просрочена с [Дата окончания]."
        ),
    }
    operation_id = str(uuid4())

    response = client.put(
        "/api/admin/reminder-templates",
        headers=_headers(token),
        json={"operation_id": operation_id, "templates": templates},
    )

    assert response.status_code == 200, response.get_json()
    assert response.get_json()["backup_created"] is True
    assert _snapshot(client, token)["reminder_templates"] == templates
    with open_readonly(settings.database_directory / "vault_cells.sqlite3") as working:
        stored = dict(
            working.execute(
                """
                SELECT key, value FROM config
                WHERE key LIKE 'whatsapp_reminder_%_template'
                """
            )
        )
    assert stored == {
        "whatsapp_reminder_expiring_template": templates["expiring"],
        "whatsapp_reminder_overdue_template": templates["overdue"],
    }
    with open_readonly(settings.database_directory / "vault_archive.sqlite3") as archive:
        audit = archive.execute(
            "SELECT action FROM log WHERE operation_id = ?", (operation_id,)
        ).fetchone()
    assert audit["action"] == "admin.reminder_templates.updated"


@pytest.mark.parametrize(
    "replacement, expected",
    [
        (
            "[Обращение] [Номер ячейки] [Дата окончания] [Лишнее]",
            "неизвестные подстановки",
        ),
        ("{{ID_CARD_NUMBER}}", "неизвестные подстановки"),
        (" ", "не может быть пустым"),
    ],
)
def test_invalid_whatsapp_template_is_rejected_without_change(
    settings, initialized_databases, replacement, expected
):
    app = _ready_app(settings)
    client = app.test_client()
    token = _setup(client)
    before = _snapshot(client, token)["reminder_templates"]

    response = client.put(
        "/api/admin/reminder-templates",
        headers=_headers(token),
        json={
            "operation_id": str(uuid4()),
            "templates": {
                "expiring": replacement,
                "overdue": before["overdue"],
            },
        },
    )

    assert response.status_code == 400
    assert expected in response.get_json()["message"]
    assert _snapshot(client, token)["reminder_templates"] == before


def test_whatsapp_template_can_remove_or_add_known_placeholders(
    settings, initialized_databases
):
    app = _ready_app(settings)
    client = app.test_client()
    token = _setup(client)
    before = _snapshot(client, token)
    expiring = "Уважаемый клиент! Ячейка [Номер ячейки], [Размер ячейки]."

    response = client.put(
        "/api/admin/reminder-templates",
        headers=_headers(token),
        json={
            "operation_id": str(uuid4()),
            "templates": {
                "expiring": expiring,
                "overdue": before["reminder_templates"]["overdue"],
            },
        },
    )

    assert response.status_code == 200
    snapshot = _snapshot(client, token)
    assert snapshot["reminder_templates"]["expiring"] == expiring
    assert any(
        item["code"] == "[Клиент.ФИО]"
        and item["usage"] == "WhatsApp и DOCX"
        for item in snapshot["template_fields"]
    )


def test_admin_adds_new_cell_idempotently_with_audit_and_backup(
    settings, initialized_databases
):
    app = _ready_app(settings)
    client = app.test_client()
    token = _setup(client)
    before = _snapshot(client, token)
    assert before["cells"]["count"] == 126
    assert before["cells"]["allowed_sizes"] == [
        {"height_mm": height, "width_mm": 220, "depth_mm": 330}
        for height in (50, 75, 100, 125, 175, 300)
    ]
    operation_id = str(uuid4())
    payload = {
        "operation_id": operation_id,
        "cells": [
            _cell("127", 50),
            _cell("128", 75),
        ],
        "new_tariffs": [],
        "new_penalty_rates": [],
    }

    first = client.post(
        "/api/admin/cells", headers=_headers(token), json=payload
    )
    repeated = client.post(
        "/api/admin/cells", headers=_headers(token), json=payload
    )

    assert first.status_code == 201, first.get_json()
    assert first.get_json()["cells"] == [
        _cell("127", 50),
        _cell("128", 75),
    ]
    assert first.get_json()["backup_created"] is True
    assert repeated.status_code == 200
    assert repeated.get_json()["repeated"] is True
    assert _snapshot(client, token)["cells"]["count"] == 128
    with open_readonly(settings.database_directory / "vault_cells.sqlite3") as working:
        cell = working.execute(
            "SELECT number, height_mm, width_mm, depth_mm FROM cells "
            "WHERE number IN ('127', '128') ORDER BY number"
        ).fetchall()
    assert [tuple(row) for row in cell] == [
        ("127", 50, 220, 330),
        ("128", 75, 220, 330),
    ]
    with open_readonly(settings.database_directory / "vault_archive.sqlite3") as archive:
        audit = archive.execute(
            "SELECT action, cell_number, changes_json FROM log WHERE operation_id=?",
            (operation_id,),
        ).fetchone()
    assert audit["action"] == "admin.cells.created"
    assert audit["cell_number"] is None
    assert json.loads(audit["changes_json"]) == {
        "cells": [
            _cell("127", 50),
            _cell("128", 75),
        ],
        "new_tariffs": [],
    }


def test_admin_can_change_shared_periods_and_add_a_completely_new_size(
    settings, initialized_databases
):
    app = _ready_app(settings)
    app.config["TODAY_PROVIDER"] = lambda: OCCURRED_AT.date()
    client = app.test_client()
    token = _setup(client)
    snapshot = _snapshot(client, token)

    periods = ((1, 45), (46, 120), (121, None))
    rates_by_size = {
        (
            row["height_mm"],
            row["width_mm"],
            row["depth_mm"],
        ): row["price_per_day_minor"]
        for row in snapshot["tariffs"]
        if row["period_from_days"] == 1
    }
    replacement_tariffs = [
        {
            "height_mm": size[0],
            "width_mm": size[1],
            "depth_mm": size[2],
            "period_from_days": start,
            "period_to_days": end,
            "price_per_day_minor": rates_by_size[size] + index,
        }
        for size in sorted(rates_by_size)
        for index, (start, end) in enumerate(periods)
    ]
    changed = client.put(
        "/api/admin/settings",
        headers=_headers(token),
        json={
            "operation_id": str(uuid4()),
            "config": snapshot["config"],
            "tariffs": replacement_tariffs,
            "penalty": snapshot["penalty"],
        },
    )
    assert changed.status_code == 200, changed.get_json()

    new_size = {"height_mm": 60, "width_mm": 250, "depth_mm": 400}
    new_tariffs = [
        {
            **new_size,
            "period_from_days": start,
            "period_to_days": end,
            "price_per_day_minor": rate,
        }
        for (start, end), rate in zip(periods, (20, 18, 15), strict=True)
    ]
    added = client.post(
        "/api/admin/cells",
        headers=_headers(token),
        json={
            "operation_id": str(uuid4()),
            "cells": [{"number": "127", **new_size}],
            "new_tariffs": new_tariffs,
            "new_penalty_rates": [],
        },
    )
    assert added.status_code == 201, added.get_json()

    with open_readonly(
        settings.database_directory / settings.working_database_name
    ) as working:
        saved_cell = working.execute(
            """
            SELECT height_mm, width_mm, depth_mm
            FROM cells WHERE number='127'
            """
        ).fetchone()
        saved_tariffs = working.execute(
            """
            SELECT period_from_days, period_to_days, price_per_day_minor
            FROM tariffs
            WHERE height_mm=60 AND width_mm=250 AND depth_mm=400
            ORDER BY period_from_days
            """
        ).fetchall()
    assert tuple(saved_cell) == (60, 250, 400)
    assert [tuple(row) for row in saved_tariffs] == [
        (1, 45, 20),
        (46, 120, 18),
        (121, None, 15),
    ]

    quote = client.post(
        "/api/rental/calculate",
        json={
            "cell_number": "127",
            "start_date": "2026-07-14",
            "end_date": "2026-08-28",
            "rent_days": 46,
        },
    )
    assert quote.status_code == 200, quote.get_json()
    assert quote.get_json()["price_per_day"] == 18


def test_admin_rejects_duplicate_cell_and_height_without_tariff(
    settings, initialized_databases
):
    app = _ready_app(settings)
    client = app.test_client()
    token = _setup(client)

    duplicate = client.post(
        "/api/admin/cells",
        headers=_headers(token),
        json={
            "operation_id": str(uuid4()),
            "cells": [_cell("1", 50)],
            "new_tariffs": [],
            "new_penalty_rates": [],
        },
    )
    unknown_height = client.post(
        "/api/admin/cells",
        headers=_headers(token),
        json={
            "operation_id": str(uuid4()),
            "cells": [_cell("127", 999)],
            "new_tariffs": [],
            "new_penalty_rates": [],
        },
    )

    assert duplicate.status_code == 409
    assert "уже есть" in duplicate.get_json()["message"]
    assert unknown_height.status_code == 400
    assert "тарифы" in unknown_height.get_json()["message"].lower()
    assert _snapshot(client, token)["cells"]["count"] == 126


def test_locked_database_does_not_partially_add_cell(
    settings, initialized_databases
):
    short = Settings(
        database_directory=settings.database_directory,
        busy_timeout_ms=40,
        testing=True,
    )
    app = _ready_app(short)
    client = app.test_client()
    token = _setup(client)

    with open_write(short, attach_archive=True) as blocker:
        blocker.execute("BEGIN IMMEDIATE")
        response = client.post(
            "/api/admin/cells",
            headers=_headers(token),
            json={
                    "operation_id": str(uuid4()),
                    "cells": [_cell("127", 50)],
                    "new_tariffs": [],
                    "new_penalty_rates": [],
            },
        )

    assert response.status_code == 423
    assert _snapshot(client, token)["cells"]["count"] == 126


def test_batch_add_is_all_or_nothing(settings, initialized_databases):
    app = _ready_app(settings)
    client = app.test_client()
    token = _setup(client)

    response = client.post(
        "/api/admin/cells",
        headers=_headers(token),
        json={
            "operation_id": str(uuid4()),
            "cells": [
                _cell("127", 50),
                _cell("1", 50),
            ],
            "new_tariffs": [],
            "new_penalty_rates": [],
        },
    )

    assert response.status_code == 409
    with open_readonly(settings.database_directory / "vault_cells.sqlite3") as working:
        assert working.execute(
            "SELECT 1 FROM cells WHERE number='127'"
        ).fetchone() is None


def test_admin_can_retire_restore_and_delete_never_used_cell(
    settings, initialized_databases
):
    app = _ready_app(settings)
    client = app.test_client()
    token = _setup(client)
    created = client.post(
        "/api/admin/cells",
        headers=_headers(token),
        json={
            "operation_id": str(uuid4()),
            "cells": [_cell("127", 50)],
            "new_tariffs": [],
            "new_penalty_rates": [],
        },
    )
    assert created.status_code == 201

    retired = client.put(
        "/api/admin/cells/lifecycle",
        headers=_headers(token),
        json={
            "operation_id": str(uuid4()),
            "number": "127",
            "action": "retire",
            "reason": "Сейф демонтирован",
        },
    )
    assert retired.status_code == 200, retired.get_json()
    snapshot = _snapshot(client, token)
    assert snapshot["cells"]["count"] == 126
    assert snapshot["cells"]["retired_count"] == 1
    assert all(
        cell["number"] != "127"
        for cell in client.get("/api/cells").get_json()["cells"]
    )

    restored = client.put(
        "/api/admin/cells/lifecycle",
        headers=_headers(token),
        json={
            "operation_id": str(uuid4()),
            "number": "127",
            "action": "restore",
            "reason": "",
        },
    )
    assert restored.status_code == 200
    assert _snapshot(client, token)["cells"]["count"] == 127

    deleted = client.put(
        "/api/admin/cells/lifecycle",
        headers=_headers(token),
        json={
            "operation_id": str(uuid4()),
            "number": "127",
            "action": "delete",
            "reason": "",
        },
    )
    assert deleted.status_code == 200
    with open_readonly(settings.database_directory / "vault_cells.sqlite3") as working:
        assert working.execute(
            "SELECT 1 FROM cells WHERE number='127'"
        ).fetchone() is None


def test_occupied_cell_cannot_be_retired_or_deleted(
    settings, initialized_databases, insert_test_contract
):
    insert_test_contract(cell_number="1", end_date="2026-08-01")
    app = _ready_app(settings)
    client = app.test_client()
    token = _setup(client)

    for action in ("retire", "delete"):
        response = client.put(
            "/api/admin/cells/lifecycle",
            headers=_headers(token),
            json={
                "operation_id": str(uuid4()),
                "number": "1",
                "action": action,
                "reason": "Проверка" if action == "retire" else "",
            },
        )
        assert response.status_code == 409
        assert "Занятую" in response.get_json()["message"]


def test_repeated_settings_operation_does_not_duplicate_audit(
    settings, initialized_databases
):
    app = _ready_app(settings)
    client = app.test_client()
    token = _setup(client)
    snapshot = _snapshot(client, token)
    operation_id = str(uuid4())
    payload = {
        "operation_id": operation_id,
        "config": snapshot["config"],
        "tariffs": snapshot["tariffs"],
        "penalty": snapshot["penalty"],
    }
    first = client.put(
        "/api/admin/settings", headers=_headers(token), json=payload
    )
    second = client.put(
        "/api/admin/settings", headers=_headers(token), json=payload
    )
    assert first.status_code == second.status_code == 200
    assert first.get_json()["repeated"] is False
    assert second.get_json()["repeated"] is True
    with open_readonly(settings.database_directory / "vault_archive.sqlite3") as archive:
        count = archive.execute(
            "SELECT COUNT(*) FROM log WHERE operation_id = ?", (operation_id,)
        ).fetchone()[0]
    assert count == 1


@pytest.mark.parametrize(
    "mutate, expected",
    [
        (lambda payload: payload["config"].update({"currency_code": "USD"}), "Разрешено"),
        (lambda payload: payload["config"].update({"deposit_amount_minor": -1}), "от 0"),
        (lambda payload: payload["tariffs"].pop(), "Последний тариф"),
        (
            lambda payload: payload["tariffs"].__setitem__(
                1, {**payload["tariffs"][1], "period_from_days": 30}
            ),
            "пропуск или пересечение",
        ),
        (lambda payload: payload["penalty"].update({"mode": "other"}), "способ"),
        (lambda payload: payload["penalty"]["manual_rates"].pop(), "всех существующих"),
        (
            lambda payload: payload["penalty"]["manual_rates"][0].update(
                {"price_per_day_minor": -1}
            ),
            "от 0",
        ),
    ],
)
def test_invalid_admin_settings_are_rejected_without_partial_change(
    settings, initialized_databases, mutate, expected
):
    app = _ready_app(settings)
    client = app.test_client()
    token = _setup(client)
    before = _snapshot(client, token)
    payload = {
        "operation_id": str(uuid4()),
        "config": dict(before["config"]),
        "tariffs": [dict(row) for row in before["tariffs"]],
        "penalty": {
            "mode": before["penalty"]["mode"],
            "manual_rates": [dict(row) for row in before["penalty"]["manual_rates"]],
        },
    }
    mutate(payload)
    response = client.put("/api/admin/settings", headers=_headers(token), json=payload)
    assert response.status_code == 400
    assert expected in response.get_json()["message"]
    assert _snapshot(client, token)["config"] == before["config"]
    assert _snapshot(client, token)["tariffs"] == before["tariffs"]
    assert _snapshot(client, token)["penalty"] == before["penalty"]


def test_database_lock_does_not_partially_save_admin_settings(
    settings, initialized_databases
):
    short = Settings(
        database_directory=settings.database_directory,
        busy_timeout_ms=40,
        testing=True,
    )
    app = _ready_app(short)
    client = app.test_client()
    token = _setup(client)
    before = _snapshot(client, token)
    payload = {
        "operation_id": str(uuid4()),
        "config": {**before["config"], "deposit_amount_minor": 2100},
        "tariffs": before["tariffs"],
        "penalty": before["penalty"],
    }
    with open_write(short, attach_archive=True) as blocker:
        blocker.execute("BEGIN IMMEDIATE")
        response = client.put(
            "/api/admin/settings", headers=_headers(token), json=payload
        )
        assert response.status_code == 423
    assert _snapshot(client, token)["config"] == before["config"]
    assert _snapshot(client, token)["penalty"] == before["penalty"]


def test_unexpected_mid_transaction_failure_rolls_back(
    settings, initialized_databases, monkeypatch
):
    app = _ready_app(settings)
    client = app.test_client()
    token = _setup(client)
    before = _snapshot(client, token)
    payload = {
        "operation_id": str(uuid4()),
        "config": {**before["config"], "deposit_amount_minor": 2200},
        "tariffs": before["tariffs"],
        "penalty": before["penalty"],
    }
    with monkeypatch.context() as patcher:
        patcher.setattr(
            "app.services.admin_settings.json.dumps",
            lambda *args, **kwargs: (_ for _ in ()).throw(TypeError("test failure")),
        )
        with pytest.raises(TypeError, match="test failure"):
            update_admin_settings(
                settings,
                payload=payload,
                employee="test-admin",
                occurred_at=OCCURRED_AT,
            )
    assert _snapshot(client, token)["config"] == before["config"]
    assert _snapshot(client, token)["penalty"] == before["penalty"]


def test_admin_can_upload_and_disable_valid_docx_template(
    settings, initialized_databases
):
    app = _ready_app(settings)
    client = app.test_client()
    token = _setup(client)
    upload = client.post(
        "/api/admin/templates",
        headers=_headers(token),
        data={
            "operation_id": str(uuid4()),
            "document_type": "manual",
            "display_name": "ТЕСТОВЫЙ ШАБЛОН",
            "file": (_docx_bytes(), "TEST_ONLY_ADMIN.docx"),
        },
        content_type="multipart/form-data",
    )
    assert upload.status_code == 201, upload.get_json()
    template_id = upload.get_json()["template_id"]
    assert (settings.database_directory / "templates" / "TEST_ONLY_ADMIN.docx").is_file()
    row = next(item for item in _snapshot(client, token)["templates"] if item["template_id"] == template_id)
    assert row["is_active"] is True

    response = client.put(
        "/api/admin/templates",
        headers=_headers(token),
        json={
            "operation_id": str(uuid4()),
            "template_id": template_id,
            "document_type": "manual",
            "display_name": "ТЕСТОВЫЙ ШАБЛОН",
            "is_active": False,
        },
    )
    assert response.status_code == 200
    row = next(item for item in _snapshot(client, token)["templates"] if item["template_id"] == template_id)
    assert row["is_active"] is False


def test_template_upload_rejects_unknown_field_and_path_escape(
    settings, initialized_databases
):
    with pytest.raises(AdminValidationError, match="неизвестные поля"):
        save_document_template(
            settings,
            operation_id=str(uuid4()),
            template_id=None,
            document_type="manual",
            display_name="ТЕСТ",
            file_name="test.docx",
            stream=_docx_bytes("НЕИЗВЕСТНОЕ.ПОЛЕ"),
            employee="test-admin",
            occurred_at=OCCURRED_AT,
        )
    with pytest.raises(AdminValidationError, match="внутри папки"):
        save_document_template(
            settings,
            operation_id=str(uuid4()),
            template_id=None,
            document_type="manual",
            display_name="ТЕСТ",
            file_name="../outside.docx",
            stream=_docx_bytes(),
            employee="test-admin",
            occurred_at=OCCURRED_AT,
        )
    assert not (settings.database_directory.parent / "outside.docx").exists()


def test_template_upload_rejects_renewal_field_for_other_document_type(
    settings, initialized_databases
):
    with pytest.raises(AdminValidationError, match="только для документа продления"):
        save_document_template(
            settings,
            operation_id=str(uuid4()),
            template_id=None,
            document_type="closing",
            display_name="ТЕСТ",
            file_name="wrong-context.docx",
            stream=_docx_bytes("Продление.Конец"),
            employee="test-admin",
            occurred_at=OCCURRED_AT,
        )


def test_static_docx_without_placeholders_can_be_uploaded_and_downloaded(
    settings, initialized_databases
):
    app = _ready_app(settings)
    client = app.test_client()
    token = _setup(client)
    upload = client.post(
        "/api/admin/templates",
        headers=_headers(token),
        data={
            "operation_id": str(uuid4()),
            "document_type": "manual",
            "display_name": "ГОТОВЫЙ ТЕСТОВЫЙ ДОКУМЕНТ",
            "file": (_static_docx_bytes(), "TEST_STATIC.docx"),
        },
        content_type="multipart/form-data",
    )
    assert upload.status_code == 201, upload.get_json()
    template_id = upload.get_json()["template_id"]
    template = next(
        item
        for item in _snapshot(client, token)["templates"]
        if item["template_id"] == template_id
    )
    assert template["required_placeholders"] == []
    assert template["file_present"] is True
    with client.get(
        f"/api/admin/templates/{template_id}/file", headers=_headers(token)
    ) as downloaded:
        assert downloaded.status_code == 200
        assert downloaded.data.startswith(b"PK")
    assert client.get(f"/api/admin/templates/{template_id}/file").status_code == 401


def test_access_mode_can_switch_between_password_and_acknowledgement(
    settings, initialized_databases
):
    app = _ready_app(settings)
    client = app.test_client()
    password_token = _setup(client)
    operation_id = str(uuid4())
    switched = client.put(
        "/api/admin/access",
        headers=_headers(password_token),
        json={"operation_id": operation_id, "access_mode": "acknowledgement"},
    )
    assert switched.status_code == 200, switched.get_json()
    assert client.get("/api/admin/status").get_json()["access_mode"] == "acknowledgement"
    assert client.post("/api/admin/login", json={"password": PASSWORD}).status_code == 409
    assert client.post("/api/admin/acknowledge", json={"accepted": False}).status_code == 400
    acknowledged = client.post("/api/admin/acknowledge", json={"accepted": True})
    assert acknowledged.status_code == 200
    acknowledgement_token = acknowledged.get_json()["token"]
    assert _snapshot(client, acknowledgement_token)["access_mode"] == "acknowledgement"

    restored = client.put(
        "/api/admin/access",
        headers=_headers(acknowledgement_token),
        json={"operation_id": str(uuid4()), "access_mode": "password"},
    )
    assert restored.status_code == 200
    assert client.post("/api/admin/acknowledge", json={"accepted": True}).status_code == 409
    assert client.post("/api/admin/login", json={"password": PASSWORD}).status_code == 200

    with open_readonly(settings.database_directory / "vault_archive.sqlite3") as archive:
        changes = archive.execute(
            "SELECT changes_json FROM log WHERE operation_id = ?", (operation_id,)
        ).fetchone()["changes_json"]
    assert PASSWORD not in changes
    assert "acknowledgement" in changes


def test_password_change_requires_current_password_and_invalidates_old_one(
    settings, initialized_databases
):
    app = _ready_app(settings)
    client = app.test_client()
    token = _setup(client)
    wrong = client.put(
        "/api/admin/password",
        headers=_headers(token),
        json={
            "operation_id": str(uuid4()),
            "current_password": "WrongPassword-1",
            "new_password": NEW_PASSWORD,
            "password_confirmation": NEW_PASSWORD,
        },
    )
    assert wrong.status_code == 400
    changed = client.put(
        "/api/admin/password",
        headers=_headers(token),
        json={
            "operation_id": str(uuid4()),
            "current_password": PASSWORD,
            "new_password": NEW_PASSWORD,
            "password_confirmation": NEW_PASSWORD,
        },
    )
    assert changed.status_code == 200
    assert client.get("/api/admin/settings", headers=_headers(token)).status_code == 401
    assert client.post("/api/admin/login", json={"password": PASSWORD}).status_code == 401
    assert client.post("/api/admin/login", json={"password": NEW_PASSWORD}).status_code == 200
